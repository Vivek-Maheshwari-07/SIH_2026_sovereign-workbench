"""
Deliverables (ticket A7): Word approval note, Excel tag list / findings, and a
PowerPoint summary, saved under workspace/artifacts/ and registered as
shared.contracts.Artifact so /api/artifacts/{artifact_id} can serve them.

Rules
- ref_no and date are ALWAYS set here, never taken from the model.
- Money amounts in cost_implication are kept only if they appear in the
  source document text; otherwise "To be filled by the originator."
- Every text value is cleaned of XML-illegal characters and length-capped
  before it is written, so Word/Excel never show a "repair" prompt.
"""
from __future__ import annotations

import copy
import json
import os
import re
import secrets
import threading
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Optional, Sequence, Union

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook, load_workbook  # noqa: F401  (load_workbook used by callers/tests)
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.util import Inches, Pt

from backend.settings import settings
from backend.tools.documents import PID_TILE_NAMES
from backend.tools.files import FileSafetyError, safe_path, sanitize_filename
from shared.contracts import (
    API_PREFIX,
    ERROR_CODES,
    ApprovalNote,
    Artifact,
    ArtifactKind,
    Finding,
    PidTagList,
    Severity,
)

_REPO_ROOT = Path(__file__).resolve().parent.parent.parent

# ---- named constants (no .env key exists for these)
TEMPLATE_NAME = "approval_note.docx"
ARTIFACTS_DIR = "artifacts"                      # under WB_WORKSPACE_DIR
ARTIFACT_INDEX_FILE = "artifacts/index.json"     # artifact_id -> file, survives restarts
REF_COUNTER_FILE = "ref_counter.json"            # {"2026": 3}, survives restarts
REF_PREFIX = "INSP"
DATE_FORMAT = "%d-%m-%Y"
DRAFT_FOOTER = "DRAFT - AI-generated with Sovereign AI Workbench (offline). Requires human review before approval."
NOT_APPLICABLE = "Not applicable"
COST_PLACEHOLDER = "To be filled by the originator."
SOP_AUTO_NOTE = "References retrieved automatically by knowledge-base search; please verify."
MAX_CELL_CHARS = 4000            # table / Excel cells / bullets
MAX_PARAGRAPH_CHARS = 10000      # Word body paragraphs
PREVIEW_CHARS = 500              # Artifact.preview limit in the contract
MAX_BULLETS_PER_SLIDE = 6
SEVERITY_FILLS = {               # Word/Excel cell colour per contract Severity
    Severity.CRITICAL: "FF9999",
    Severity.HIGH: "FFC7CE",
    Severity.MEDIUM: "FFE0B2",
    Severity.LOW: "C6EFCE",
}
MEDIA_TYPES = {
    ArtifactKind.DOCX: "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ArtifactKind.XLSX: "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ArtifactKind.PPTX: "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}
_HEADER_FILL_XLSX = "1F3864"

_ARTIFACT_ID_RE = re.compile(r"^a_[0-9a-f]{12}$")
# XML 1.0 forbids C0 controls except tab/newline/carriage return, plus lone surrogates and U+FFFE/U+FFFF.
_ILLEGAL_XML_RE = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\x7f\ud800-\udfff￾￿]")
_NUMBER_RE = re.compile(r"\d[\d,]*(?:\.\d+)?")
_PLACEHOLDER_RE = re.compile(r"\{\{(?:ref_no|date|subject|background|recommendation|cost_implication|"
                             r"prepared_by|sop_note|finding\.\w+|sop\.\w+)\}\}")
_SOP_PAGE_RE = re.compile(r"^(?P<doc>.*?)[\s,;:\-]*(?:p\.|pp\.|pg\.?|page)\s*(?P<page>\d+)\s*\.?$", re.IGNORECASE)

_ref_lock = threading.Lock()
_index_lock = threading.Lock()


class OfficeError(Exception):
    """A deliverable could not be produced. `code` is a key from ERROR_CODES."""

    def __init__(self, code: str, message: Optional[str] = None) -> None:
        if code not in ERROR_CODES:
            raise ValueError(f"unknown error code: {code}")
        super().__init__(message or ERROR_CODES[code])
        self.code = code


# ------------------------------------------------------------------ text cleaning
def clean_text(value: Any, limit: int = MAX_CELL_CHARS) -> str:
    """Drop XML-illegal characters (keeps tab/newline, emoji, Hindi, ₹) and cap the length."""
    text = "" if value is None else str(value)
    text = _ILLEGAL_XML_RE.sub("", text).replace("\r\n", "\n").replace("\r", "\n")
    if len(text) > limit:
        text = text[: max(0, limit - 1)].rstrip() + "…"
    return text


def _or_na(value: Any, limit: int = MAX_PARAGRAPH_CHARS) -> str:
    text = clean_text(value, limit).strip()
    return text or NOT_APPLICABLE


def _preview(text: str) -> str:
    return clean_text(" ".join(text.split()), PREVIEW_CHARS)


# ------------------------------------------------------------------ small JSON state files
def _read_json(path: Path, default: Any) -> Any:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (FileNotFoundError, json.JSONDecodeError):
        return default


def _write_json_atomic(path: Path, data: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(f".{path.name}.{os.getpid()}.{threading.get_ident()}.tmp")
    tmp.write_text(json.dumps(data, indent=2), encoding="utf-8")
    os.replace(tmp, path)


# ------------------------------------------------------------------ ref numbers
def next_ref_no(now: Optional[datetime] = None) -> str:
    """INSP-<year>-<NNN>, counted per year in workspace/ref_counter.json; never repeats."""
    year = str((now or datetime.now()).year)
    with _ref_lock:
        path = safe_path(REF_COUNTER_FILE)
        counters = _read_json(path, {})
        number = int(counters.get(year, 0)) + 1
        counters[year] = number
        _write_json_atomic(path, counters)
    return f"{REF_PREFIX}-{year}-{number:03d}"


def today_str(now: Optional[datetime] = None) -> str:
    return (now or datetime.now()).strftime(DATE_FORMAT)


# ------------------------------------------------------------------ artifact registry
def _new_artifact_id() -> str:
    return "a_" + secrets.token_hex(6)


def is_valid_artifact_id(artifact_id: str) -> bool:
    return bool(_ARTIFACT_ID_RE.match(artifact_id or ""))


def register_artifact(path: Path, kind: ArtifactKind, *, task_id: str = "", preview: str = "",
                      artifact_id: Optional[str] = None) -> Artifact:
    artifact_id = artifact_id or _new_artifact_id()
    workspace = safe_path(".")
    artifact = Artifact(
        artifact_id=artifact_id,
        task_id=task_id,
        filename=path.name,
        kind=kind,
        size_bytes=path.stat().st_size,
        created_at=datetime.now(timezone.utc),
        download_url=f"{API_PREFIX}/artifacts/{artifact_id}",
        preview=_preview(preview) or None,
    )
    with _index_lock:
        index_path = safe_path(ARTIFACT_INDEX_FILE)
        index = _read_json(index_path, {})
        index[artifact_id] = {
            "path": path.relative_to(workspace).as_posix(),
            "artifact": artifact.model_dump(mode="json"),
        }
        _write_json_atomic(index_path, index)
    return artifact


def artifact_path(artifact_id: str) -> Optional[tuple[Path, Artifact]]:
    """(file path, Artifact) for a registered id whose file still exists, else None."""
    if not is_valid_artifact_id(artifact_id):
        return None
    with _index_lock:
        entry = _read_json(safe_path(ARTIFACT_INDEX_FILE), {}).get(artifact_id)
    if not entry:
        return None
    try:
        path = safe_path(entry["path"])
    except (FileSafetyError, KeyError):
        return None
    if not path.is_file():
        return None
    return path, Artifact.model_validate(entry["artifact"])


def media_type(kind: ArtifactKind) -> str:
    return MEDIA_TYPES.get(kind, "application/octet-stream")


def _artifact_file(filename: str) -> Path:
    path = safe_path(f"{ARTIFACTS_DIR}/{sanitize_filename(filename)}")
    path.parent.mkdir(parents=True, exist_ok=True)
    return path


def _guard(fn):
    """Turn any unexpected failure into OfficeError so callers get an error code, never a crash."""
    def wrapper(*args, **kwargs):
        try:
            return fn(*args, **kwargs)
        except OfficeError:
            raise
        except FileSafetyError as exc:
            raise OfficeError(exc.code, str(exc)) from exc
        except (TypeError, ValueError) as exc:
            raise OfficeError("BAD_REQUEST", f"{fn.__name__}: bad input: {exc}") from exc
        except Exception as exc:
            raise OfficeError("INTERNAL", f"{fn.__name__} failed: {exc!r}") from exc
    wrapper.__name__ = fn.__name__
    wrapper.__doc__ = fn.__doc__
    return wrapper


@_guard
def save_text_artifact(filename: str, text: str, kind: ArtifactKind, *, task_id: str = "",
                       preview: str = "") -> Artifact:
    """Save a plain-text deliverable (e.g. generated solution.py) and register it like the Office files."""
    path = _artifact_file(filename)
    path.write_text(_ILLEGAL_XML_RE.sub("", text), encoding="utf-8")
    return register_artifact(path, kind, task_id=task_id, preview=preview or text)


# ------------------------------------------------------------------ Word helpers
def _template_path() -> Path:
    root = settings.WB_TEMPLATES_DIR
    return (root if root.is_absolute() else _REPO_ROOT / root) / TEMPLATE_NAME


def _set_paragraph_text(paragraph, text: str) -> None:
    """Replace a paragraph's text, keeping the first run's formatting."""
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run._r.getparent().remove(run._r)


def _set_cell_text(cell, text: str) -> None:
    _set_paragraph_text(cell.paragraphs[0], text)
    for extra in cell.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)


def _shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    for old in tc_pr.findall(qn("w:shd")):
        tc_pr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _all_paragraphs(doc) -> Iterable:
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def _fill_placeholders(doc, values: dict[str, str]) -> None:
    for paragraph in _all_paragraphs(doc):
        text = paragraph.text
        if "{{" not in text:
            continue
        for key, value in values.items():
            text = text.replace("{{" + key + "}}", value)
        if text != paragraph.text:
            _set_paragraph_text(paragraph, text)


def _table_with(doc, placeholder: str):
    for table in doc.tables:
        if len(table.rows) > 1 and placeholder in table.rows[1].cells[0].text:
            return table
    raise OfficeError("INTERNAL", f"template {TEMPLATE_NAME} is missing the {placeholder} table")


def _fill_rows(table, rows: list[list[str]], fills: Optional[list[Optional[str]]] = None, fill_col: int = -1) -> None:
    """Copy the placeholder row once per data row, fill it, then drop the placeholder row."""
    template_tr = table.rows[1]._tr
    for i, values in enumerate(rows):
        new_tr = copy.deepcopy(template_tr)
        template_tr.addprevious(new_tr)
        row = table.rows[1 + i]
        for cell, value in zip(row.cells, values):
            _set_cell_text(cell, value)
        if fills and fills[i]:
            _shade_cell(row.cells[fill_col], fills[i])
    template_tr.getparent().remove(template_tr)


def _set_sop_note(doc, text: str) -> None:
    """Fill the {{sop_note}} line under the SOP table, or remove the line when there is nothing to say."""
    for paragraph in doc.paragraphs:
        if "{{sop_note}}" in paragraph.text:
            if text:
                _set_paragraph_text(paragraph, text)
            else:
                paragraph._p.getparent().remove(paragraph._p)
            return


def split_sop_reference(reference: str) -> tuple[str, str]:
    """'SOP-INSP-012, p.4' -> ('SOP-INSP-012', '4'); no page found -> (text, '-')."""
    text = clean_text(reference, MAX_CELL_CHARS).strip()
    match = _SOP_PAGE_RE.match(text)
    if match and match.group("doc").strip():
        return match.group("doc").strip(), match.group("page")
    return text, "-"


def _normalize_number(number: str) -> str:
    """'4,50,000' / '450,000' / '4,50,000.00' -> '450000'. Words like lakh/crore are NOT converted."""
    plain = number.replace(",", "")
    if re.fullmatch(r"\d+\.0+", plain):
        plain = plain.split(".", 1)[0]
    return plain


def cost_text(cost: Optional[str], source_text: Optional[str]) -> str:
    """
    Keep cost_implication only if every number in it appears in the source
    document text; the model must never invent money amounts. Text with no
    numbers (e.g. "Within the approved maintenance budget") is kept.
    """
    text = clean_text(cost, MAX_PARAGRAPH_CHARS).strip()
    if not text:
        return COST_PLACEHOLDER
    numbers = [_normalize_number(n) for n in _NUMBER_RE.findall(text)]
    if not numbers:
        return text
    source_numbers = {_normalize_number(n) for n in _NUMBER_RE.findall(source_text or "")}
    if all(n in source_numbers for n in numbers):
        return text
    return COST_PLACEHOLDER


def _word_preview(note: ApprovalNote) -> str:
    parts = [f"Subject: {note.subject}.", "Findings:"]
    for i, finding in enumerate(note.findings, start=1):
        parts.append(f"({i}) {finding.severity.value.upper()} {finding.item}: {finding.observation};")
    return " ".join(parts)


# ------------------------------------------------------------------ make_word
@_guard
def make_word(note: ApprovalNote, *, task_id: str = "", source_text: Optional[str] = None,
              sop_auto: bool = False) -> Artifact:
    """
    Fill templates/approval_note.docx from `note`. ref_no and date are always
    generated here. `source_text` (the inspection report text) is used only
    to verify money amounts in cost_implication.
    """
    if not isinstance(note, ApprovalNote):
        raise OfficeError("BAD_REQUEST", "make_word expects an ApprovalNote")
    template = _template_path()
    if not template.is_file():
        raise OfficeError("INTERNAL", f"Word template missing: {template} (run scripts/make_template.py)")

    note = note.model_copy(update={"ref_no": next_ref_no(), "date": today_str()})
    doc = Document(str(template))

    _fill_placeholders(doc, {
        "ref_no": note.ref_no,
        "date": note.date,
        "subject": _or_na(note.subject, MAX_CELL_CHARS),
        "background": _or_na(note.background),
        "recommendation": _or_na(note.recommendation),
        "cost_implication": cost_text(note.cost_implication, source_text),
        "prepared_by": _or_na(note.prepared_by, MAX_CELL_CHARS),
    })

    findings = _table_with(doc, "{{finding.sno}}")
    if note.findings:
        rows = [[str(i), _or_na(f.item, MAX_CELL_CHARS), _or_na(f.observation, MAX_CELL_CHARS),
                 f.severity.value.capitalize(), str(f.source_page) if f.source_page else "-"]
                for i, f in enumerate(note.findings, start=1)]
        _fill_rows(findings, rows, [SEVERITY_FILLS.get(f.severity) for f in note.findings], fill_col=3)
    else:
        _fill_rows(findings, [["-", NOT_APPLICABLE, "", "", ""]])

    sops = _table_with(doc, "{{sop.sno}}")
    references = [r for r in note.sop_references if clean_text(r).strip()]
    if references:
        _fill_rows(sops, [[str(i), *split_sop_reference(r)] for i, r in enumerate(references, start=1)])
    else:
        _fill_rows(sops, [["-", NOT_APPLICABLE, "-"]])
    _set_sop_note(doc, SOP_AUTO_NOTE if sop_auto and references else "")

    leftover = [p.text for p in _all_paragraphs(doc) if _PLACEHOLDER_RE.search(p.text)]
    if leftover:
        raise OfficeError("INTERNAL", f"unfilled template placeholders: {leftover}")

    path = _artifact_file(f"{note.ref_no}_approval_note.docx")
    doc.save(str(path))
    return register_artifact(path, ArtifactKind.DOCX, task_id=task_id, preview=_word_preview(note))


# ------------------------------------------------------------------ Excel helpers
def _style_sheet(ws, header: Sequence[str], rows: list[list[Any]]) -> None:
    ws.append([clean_text(h) for h in header])
    for row in rows:
        ws.append([clean_text(v) if isinstance(v, str) else v for v in row])
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor=_HEADER_FILL_XLSX)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{max(1, len(rows) + 1)}"
    for col_idx, title in enumerate(header, start=1):
        longest = max([len(str(title))] + [len(str(r[col_idx - 1])) for r in rows if r[col_idx - 1] is not None])
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(10, longest + 2), 60)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def _tile_name(tile: Optional[int]) -> str:
    if tile is None:
        return ""
    if 1 <= tile <= len(PID_TILE_NAMES):
        return PID_TILE_NAMES[tile - 1]
    return str(tile)


def dedupe_tags(tag_list: PidTagList) -> list[dict[str, str]]:
    """One row per tag (trimmed, case-insensitive); tiles of duplicates are merged in order."""
    rows: dict[str, dict[str, Any]] = {}
    for tag in tag_list.tags:
        text = clean_text(tag.tag, MAX_CELL_CHARS).strip()
        key = text.upper()
        if not key:
            continue
        tile = _tile_name(tag.tile)
        if key not in rows:
            rows[key] = {
                "tag": text,
                "equipment_type": clean_text(tag.equipment_type).strip(),
                "description": clean_text(tag.description).strip(),
                "tiles": [tile] if tile else [],
            }
        elif tile and tile not in rows[key]["tiles"]:
            rows[key]["tiles"].append(tile)
    return [{**r, "tile": ", ".join(r["tiles"])} for r in rows.values()]


def count_by_type(rows: list[dict[str, str]]) -> list[tuple[str, int]]:
    counts: dict[str, list[Any]] = {}
    for row in rows:
        name = row["equipment_type"] or "Unknown"
        entry = counts.setdefault(name.lower(), [name, 0])
        entry[1] += 1
    return sorted(((name, n) for name, n in counts.values()), key=lambda x: (-x[1], x[0].lower()))


def _tags_workbook(tag_list: PidTagList) -> tuple[Workbook, str]:
    rows = dedupe_tags(tag_list)
    wb = Workbook()
    tags = wb.active
    tags.title = "Tags"
    _style_sheet(tags, ["Tag", "Equipment type", "Description", "Tile"],
                 [[r["tag"], r["equipment_type"], r["description"], r["tile"]] for r in rows])
    counts = count_by_type(rows)
    summary = wb.create_sheet("Summary")
    _style_sheet(summary, ["Equipment type", "Count"], [[name, n] for name, n in counts])
    summary.append(["Total", len(rows)])
    summary.cell(row=summary.max_row, column=1).font = Font(bold=True)
    summary.cell(row=summary.max_row, column=2).font = Font(bold=True)
    if tag_list.drawing_title or tag_list.notes:
        summary.append([])
        summary.append(["Drawing title", clean_text(tag_list.drawing_title) or NOT_APPLICABLE])
        summary.append(["Notes", clean_text(tag_list.notes) or NOT_APPLICABLE])
    preview = f"{len(rows)} unique tags: " + ", ".join(f"{n} {name}" for name, n in counts)
    return wb, preview


def _findings_workbook(findings: list[Finding]) -> tuple[Workbook, str]:
    wb = Workbook()
    ws = wb.active
    ws.title = "Findings"
    rows = [[i, clean_text(f.item), clean_text(f.observation), f.severity.value.capitalize(),
             f.source_page if f.source_page is not None else ""]
            for i, f in enumerate(findings, start=1)]
    _style_sheet(ws, ["S.No", "Item", "Observation", "Severity", "Source page"], rows)
    for row_idx, finding in enumerate(findings, start=2):
        fill = SEVERITY_FILLS.get(finding.severity)
        if fill:
            ws.cell(row=row_idx, column=4).fill = PatternFill("solid", fgColor=fill)
    high = sum(f.severity in (Severity.HIGH, Severity.CRITICAL) for f in findings)
    preview = f"{len(findings)} findings ({high} high/critical): " + "; ".join(
        f"{f.severity.value.upper()} {f.item}" for f in findings)
    return wb, preview


# ------------------------------------------------------------------ make_excel
@_guard
def make_excel(data: Union[PidTagList, list[Finding]], *, task_id: str = "") -> Artifact:
    """PidTagList -> Tags + Summary sheets (duplicates merged); list[Finding] -> Findings sheet."""
    short = secrets.token_hex(4)
    if isinstance(data, PidTagList):
        wb, preview = _tags_workbook(data)
        filename = f"pid_tags_{short}.xlsx"
    elif isinstance(data, list) and all(isinstance(f, Finding) for f in data):
        wb, preview = _findings_workbook(data)
        filename = f"findings_{short}.xlsx"
    else:
        raise OfficeError("BAD_REQUEST", "make_excel expects a PidTagList or a list of Finding")
    path = _artifact_file(filename)
    wb.save(str(path))
    return register_artifact(path, ArtifactKind.XLSX, task_id=task_id, preview=preview)


# ------------------------------------------------------------------ make_ppt
def split_bullets(bullets: Sequence[str], per_slide: int = MAX_BULLETS_PER_SLIDE) -> list[list[str]]:
    items = [clean_text(b, MAX_CELL_CHARS).strip() for b in bullets]
    items = [b for b in items if b] or [NOT_APPLICABLE]
    return [items[i:i + per_slide] for i in range(0, len(items), per_slide)]


def _add_footer(prs, slide) -> None:
    box = slide.shapes.add_textbox(Inches(0.3), prs.slide_height - Inches(0.45),
                                   prs.slide_width - Inches(0.6), Inches(0.35))
    frame = box.text_frame
    frame.word_wrap = True
    run = frame.paragraphs[0].add_run()
    run.text = DRAFT_FOOTER
    run.font.size = Pt(9)
    run.font.italic = True


@_guard
def make_ppt(title: str, sections: Sequence[tuple[str, Sequence[str]]], *, task_id: str = "") -> Artifact:
    """Title slide + one slide per section, at most 6 bullets per slide ("(cont.)" slides after that)."""
    clean_title = _or_na(title, 200)
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = clean_title
    slide.placeholders[1].text = f"Draft generated {today_str()}"
    _add_footer(prs, slide)

    headings: list[str] = []
    for heading, bullets in sections:
        heading = _or_na(heading, 200)
        headings.append(heading)
        for part, chunk in enumerate(split_bullets(bullets)):
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = heading if part == 0 else f"{heading} (cont.)"
            body = slide.placeholders[1].text_frame
            body.text = chunk[0]
            for bullet in chunk[1:]:
                body.add_paragraph().text = bullet
            _add_footer(prs, slide)

    slug = sanitize_filename(clean_title.replace(" ", "_"))[:40] or "presentation"
    path = _artifact_file(f"{slug}_{secrets.token_hex(4)}.pptx")
    prs.save(str(path))
    preview = f"{clean_title} - {len(prs.slides)} slides: " + "; ".join(headings)
    return register_artifact(path, ArtifactKind.PPTX, task_id=task_id, preview=preview)
