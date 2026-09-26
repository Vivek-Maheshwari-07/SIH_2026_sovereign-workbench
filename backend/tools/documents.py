"""
Document extraction for Track A. Per page, in order:
  1. PDF page has a real text layer (>= MIN_TEXT_LAYER_CHARS)  -> use it directly.
  2. No text layer                                             -> render at WB_OCR_DPI, deskew
                                                                   (see _deskew), run Tesseract.
  3. OCR came out poor (see _is_ocr_poor)                      -> send the page image to the vision model.
  4. kind="pid"                                                -> skip 1-3 entirely; cut the image
                                                                   into 4 overlapping tiles and send
                                                                   each to the vision model.
A standalone image (png/jpg) is treated like a single page with no text
layer, so it goes straight into step 2.

Non-image, non-PDF text-bearing types skip the OCR/vision pipeline entirely
(there's nothing to render or recognize) and are read directly:
  - txt/md/py/csv -> read as plain text, method "text_file".
  - docx          -> paragraphs + table cells via python-docx, method "docx".
  - xlsx          -> sheet names + cell values via openpyxl, method "xlsx".

Results are cached under WB_CACHE_DIR, keyed by a sha256 of the file bytes
plus every setting that changes the output (dpi, max px, vision model id,
kind, deskew on/off), so re-extracting the same file with the same settings is instant and
touches neither Tesseract nor Ollama.

PageResult/ExtractResult are internal dataclasses, not shared.contracts
types (AGENTS.md rule 5 only governs API request/response models — these
never cross the HTTP boundary; callers turn them into contract types, e.g.
Finding.source_page, themselves).
"""
from __future__ import annotations

import hashlib
import io
import json
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Literal, Optional, Union

import docx
import openpyxl
import pymupdf
import pytesseract
from PIL import Image, ImageOps

from backend.llm_client import LLMError, chat
from backend.registry import registry
from backend.settings import settings
from shared.contracts import ERROR_CODES, TaskType

_REPO_ROOT = Path(__file__).resolve().parent.parent.parent

# ---------------------------------------------------------------- constants
# No .env key covers these — they're extraction-quality knobs, not deploy
# config, so they live here as named constants.

# A PDF page counts as having a usable text layer once it has at least this
# many non-whitespace characters. A tiny stray header/footer text object can
# exist even on an otherwise-scanned page, so a small floor avoids treating
# that as "has text" and skipping OCR on a page that actually needs it.
MIN_TEXT_LAYER_CHARS = 20

# Tesseract's image_to_data gives a 0-100 confidence per recognized word
# (-1 for non-text regions, which we exclude before averaging). OCR counts
# as "poor" -> fall back to the vision model when EITHER:
#   - the mean confidence of the words it did recognize is below this, or
#   - it recognized fewer than this many words at all (a near-blank read
#     has no confidence figures to average, which would otherwise look
#     like a lucky 0.0-mean-of-nothing rather than the failure it is).
OCR_MIN_MEAN_CONFIDENCE = 60.0
OCR_MIN_WORD_COUNT = 3

# P&ID tiling: a 2x2 grid with ~15% overlap between neighbouring tiles on
# each axis, so a tag straddling the midline isn't split across two tiles.
PID_TILE_OVERLAP = 0.15
PID_TILE_NAMES = ("top-left", "top-right", "bottom-left", "bottom-right")

# Deskew before OCR. Tesseract silently drops table cells on scans rotated by
# only 0.6-1.5 degrees while still reporting ~90 mean confidence (so the
# vision fallback never kicks in); straightening the page first fixes that.
# The angle is found by a projection-profile search: rotate a small binarized
# copy through candidate angles and keep the one whose row-ink profile is the
# sharpest (text lines line up with pixel rows).
DESKEW_ENABLED = True
DESKEW_MAX_ANGLE = 5.0        # degrees searched either side of straight
DESKEW_COARSE_STEP = 0.5
DESKEW_FINE_STEP = 0.1        # searched within +/- one coarse step of the coarse best
DESKEW_MIN_ANGLE = 0.1        # smaller detected angles are left alone
DESKEW_DETECT_MAX_PX = 1200   # longest side of the copy the angle is detected on
DESKEW_INK_THRESHOLD = 128    # grey level below which a pixel counts as ink

# Bump if the cached JSON shape or the extraction pipeline changes in a way
# that makes old cached results wrong. 2 = deskew before OCR (A8c).
_CACHE_VERSION = 2


class DocumentExtractionError(Exception):
    """Raised for a whole-document extraction failure. `code` is an ERROR_CODES key."""

    def __init__(self, code: str, message: Optional[str] = None) -> None:
        if code not in ERROR_CODES:
            raise ValueError(f"unknown error code: {code}")
        super().__init__(message or ERROR_CODES[code])
        self.code = code


ExtractMethod = Literal["text_layer", "ocr", "vision", "vision_tiles", "text_file", "docx", "xlsx"]

# Extensions read directly as plain text (no OCR/vision needed).
_TEXT_FILE_SUFFIXES = {".txt", ".md", ".py", ".csv"}


@dataclass
class PageResult:
    page: int  # 1-indexed
    text: str
    method: ExtractMethod
    ocr_confidence: Optional[float] = None
    tile: Optional[str] = None  # set only for kind="pid" results
    deskew_angle: Optional[float] = None  # detected skew in degrees (CCW +); set when the page went through OCR
    error: Optional[str] = None  # set if this page/tile failed; text is best-effort


@dataclass
class ExtractResult:
    pages: list[PageResult] = field(default_factory=list)
    from_cache: bool = False

    def joined_text(self) -> str:
        """Joins all pages into one string, with '--- Page N ---' markers (Scenario A needs source_page)."""
        return "\n\n".join(f"--- Page {p.page} ---\n{p.text}" for p in self.pages)


def _resolve(path: Path) -> Path:
    return path if path.is_absolute() else _REPO_ROOT / path


# ---------------------------------------------------------------- vision model / resize
def _vision_model_name() -> str:
    return registry.model_for_task(TaskType.VISION).ollama_name


def _resize_for_vision(image: Image.Image) -> Image.Image:
    """Resizes so the longest side is at most WB_VISION_MAX_PX, keeping aspect ratio."""
    max_px = settings.WB_VISION_MAX_PX
    width, height = image.size
    longest = max(width, height)
    if longest <= max_px:
        return image
    scale = max_px / longest
    new_size = (max(1, round(width * scale)), max(1, round(height * scale)))
    return image.resize(new_size, Image.LANCZOS)


def _image_to_png_bytes(image: Image.Image) -> bytes:
    buf = io.BytesIO()
    image.save(buf, format="PNG")
    return buf.getvalue()


def _vision_extract_page(image: Image.Image, *, purpose: str) -> str:
    model = _vision_model_name()
    png_bytes = _image_to_png_bytes(_resize_for_vision(image))
    result = chat(
        model,
        [
            {
                "role": "user",
                "content": (
                    "Read every word of text visible in this image and transcribe it as "
                    "plain text. Output only the transcription, no commentary."
                ),
            }
        ],
        images=[png_bytes],
        purpose=purpose,
    )
    return result.text.strip()


# ---------------------------------------------------------------- OCR
def _configure_tesseract() -> None:
    pytesseract.pytesseract.tesseract_cmd = str(settings.TESSERACT_CMD)


def _ocr_page(image: Image.Image) -> tuple[str, float, int]:
    """Returns (text, mean_word_confidence, word_count)."""
    _configure_tesseract()
    data = pytesseract.image_to_data(image, output_type=pytesseract.Output.DICT)

    words: list[str] = []
    confidences: list[float] = []
    for raw_text, raw_conf in zip(data["text"], data["conf"]):
        text = raw_text.strip()
        if not text:
            continue
        words.append(text)
        try:
            conf_value = float(raw_conf)
        except (TypeError, ValueError):
            conf_value = -1.0
        if conf_value >= 0:
            confidences.append(conf_value)

    mean_confidence = sum(confidences) / len(confidences) if confidences else 0.0
    return " ".join(words), mean_confidence, len(words)


def _ink_mask(image: Image.Image) -> Image.Image:
    """Small greyscale copy with ink = 255 and paper = 0 (so rotation fill adds no ink)."""
    grey = image.convert("L")
    longest = max(grey.size)
    if longest > DESKEW_DETECT_MAX_PX:
        scale = DESKEW_DETECT_MAX_PX / longest
        grey = grey.resize((max(1, round(grey.width * scale)), max(1, round(grey.height * scale))), Image.BILINEAR)
    return grey.point(lambda v: 255 if v < DESKEW_INK_THRESHOLD else 0)


def _profile_score(mask: Image.Image, angle: float) -> float:
    """Sharpness of the row-ink profile after rotating by -angle: sum of squared row sums."""
    rotated = mask.rotate(-angle, resample=Image.NEAREST, expand=True, fillcolor=0)
    row_means = rotated.convert("F").resize((1, rotated.height), Image.BOX)
    width = rotated.width
    return sum((row_means.getpixel((0, y)) * width) ** 2 for y in range(rotated.height))


def _best_angle(mask: Image.Image, candidates: list[float]) -> float:
    # Candidates nearest 0 first, and a strict ">" below, so ties (e.g. a blank page) resolve to straight.
    best_angle, best_score = 0.0, float("-inf")
    for angle in sorted(candidates, key=abs):
        score = _profile_score(mask, angle)
        if score > best_score:
            best_angle, best_score = angle, score
    return best_angle


def _frange(start: float, stop: float, step: float) -> list[float]:
    count = round((stop - start) / step)
    return [round(start + i * step, 4) for i in range(count + 1)]


def detect_skew(image: Image.Image) -> float:
    """Skew of the page in degrees, counter-clockwise positive (PIL's rotate convention)."""
    mask = _ink_mask(image)
    coarse = _best_angle(mask, _frange(-DESKEW_MAX_ANGLE, DESKEW_MAX_ANGLE, DESKEW_COARSE_STEP))
    fine = _frange(coarse - DESKEW_COARSE_STEP, coarse + DESKEW_COARSE_STEP, DESKEW_FINE_STEP)
    return _best_angle(mask, fine)


def _deskew(image: Image.Image) -> tuple[Image.Image, float]:
    """Returns (straightened image, detected angle). Tiny angles are detected but not rotated."""
    angle = detect_skew(image)
    if abs(angle) < DESKEW_MIN_ANGLE:
        return image, angle
    straight = image.convert("RGB").rotate(-angle, resample=Image.BICUBIC, expand=True, fillcolor="white")
    return straight, angle


def _is_ocr_poor(mean_confidence: float, word_count: int) -> bool:
    return word_count < OCR_MIN_WORD_COUNT or mean_confidence < OCR_MIN_MEAN_CONFIDENCE


def _ocr_then_vision(image: Image.Image, page_number: int) -> PageResult:
    angle: Optional[float] = None
    if DESKEW_ENABLED:
        image, angle = _deskew(image)
    result = _ocr_then_vision_straight(image, page_number)
    result.deskew_angle = angle
    return result


def _ocr_then_vision_straight(image: Image.Image, page_number: int) -> PageResult:
    try:
        text, mean_conf, word_count = _ocr_page(image)
        ocr_failed = False
    except Exception as exc:
        text, mean_conf, word_count = "", 0.0, 0
        ocr_failed = True
        ocr_error = str(exc)

    if not ocr_failed and not _is_ocr_poor(mean_conf, word_count):
        return PageResult(page=page_number, text=text, method="ocr", ocr_confidence=mean_conf)

    try:
        vision_text = _vision_extract_page(image, purpose="document_extract_vision")
        return PageResult(
            page=page_number,
            text=vision_text,
            method="vision",
            ocr_confidence=None if ocr_failed else mean_conf,
        )
    except LLMError as exc:
        # Keep whatever OCR text we already have (even if poor) instead of losing the page.
        note = f"OCR failed ({ocr_error})" if ocr_failed else "OCR was poor"
        return PageResult(
            page=page_number,
            text=text,
            method="ocr",
            ocr_confidence=None if ocr_failed else mean_conf,
            error=f"{note}; vision fallback also failed: {exc}",
        )


# ---------------------------------------------------------------- P&ID tiling
def _make_pid_tiles(image: Image.Image) -> list[tuple[str, Image.Image]]:
    """2x2 grid, each tile overlapping its neighbours by PID_TILE_OVERLAP on each axis."""
    width, height = image.size
    half_w, half_h = width / 2, height / 2
    overlap_w, overlap_h = half_w * PID_TILE_OVERLAP, half_h * PID_TILE_OVERLAP

    boxes = {
        "top-left": (0, 0, half_w + overlap_w, half_h + overlap_h),
        "top-right": (half_w - overlap_w, 0, width, half_h + overlap_h),
        "bottom-left": (0, half_h - overlap_h, half_w + overlap_w, height),
        "bottom-right": (half_w - overlap_w, half_h - overlap_h, width, height),
    }
    return [(name, image.crop(tuple(round(v) for v in boxes[name]))) for name in PID_TILE_NAMES]


_PID_TILE_PROMPT = (
    "This is one tile of a larger P&ID (piping and instrumentation diagram). "
    "List every equipment or instrument tag you can read in this tile (for "
    "example P-101, V-201, FT-301), one per line. If you can't read any tags "
    "in this tile, say so plainly."
)


def _vision_extract_pid_tiles(image: Image.Image, *, purpose: str) -> list[PageResult]:
    model = _vision_model_name()
    results: list[PageResult] = []
    for tile_name, tile_image in _make_pid_tiles(image):
        png_bytes = _image_to_png_bytes(_resize_for_vision(tile_image))
        try:
            chat_result = chat(
                model,
                [{"role": "user", "content": _PID_TILE_PROMPT}],
                images=[png_bytes],
                purpose=purpose,
            )
            results.append(PageResult(page=1, text=chat_result.text.strip(), method="vision_tiles", tile=tile_name))
        except LLMError as exc:  # a broken tile must not stop the others
            results.append(PageResult(page=1, text="", method="vision_tiles", tile=tile_name, error=str(exc)))
    return results


# ---------------------------------------------------------------- caching
def _cache_dir() -> Path:
    cache_dir = _resolve(settings.WB_CACHE_DIR)
    cache_dir.mkdir(parents=True, exist_ok=True)
    return cache_dir


def _cache_key(file_bytes: bytes, *, kind: str) -> str:
    hasher = hashlib.sha256()
    hasher.update(file_bytes)
    parts = (str(_CACHE_VERSION), kind, str(settings.WB_OCR_DPI), str(settings.WB_VISION_MAX_PX),
             _vision_model_name(), f"deskew={DESKEW_ENABLED}")
    for part in parts:
        hasher.update(b"\x00")
        hasher.update(part.encode("utf-8"))
    return hasher.hexdigest()


def _load_cache(key: str) -> Optional[ExtractResult]:
    path = _cache_dir() / f"{key}.json"
    if not path.exists():
        return None
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
        pages = [PageResult(**p) for p in data["pages"]]
    except Exception:
        return None  # corrupted cache file -> ignore it and redo
    return ExtractResult(pages=pages, from_cache=True)


def _save_cache(key: str, result: ExtractResult) -> None:
    path = _cache_dir() / f"{key}.json"
    try:
        path.write_text(json.dumps({"pages": [asdict(p) for p in result.pages]}), encoding="utf-8")
    except Exception:
        pass  # caching is best-effort; never fail extraction because of it


# ---------------------------------------------------------------- per-format extraction
def _extract_pdf_page(doc: "pymupdf.Document", index: int, page_number: int) -> PageResult:
    page = doc[index]
    text_layer = page.get_text().strip()
    if len(text_layer) >= MIN_TEXT_LAYER_CHARS:
        return PageResult(page=page_number, text=text_layer, method="text_layer")

    pixmap = page.get_pixmap(dpi=settings.WB_OCR_DPI)
    image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
    return _ocr_then_vision(image, page_number)


def _extract_pdf_pages(path: Path) -> list[PageResult]:
    try:
        doc = pymupdf.open(path)
    except Exception as exc:
        raise DocumentExtractionError("UNSUPPORTED_FILE", f"could not open PDF: {exc}") from exc

    pages: list[PageResult] = []
    try:
        for index in range(doc.page_count):
            page_number = index + 1
            try:
                pages.append(_extract_pdf_page(doc, index, page_number))
            except Exception as exc:  # a broken page must not stop the others
                pages.append(PageResult(page=page_number, text="", method="text_layer", error=str(exc)))
    finally:
        doc.close()
    return pages


def _extract_single_image_page(path: Path) -> PageResult:
    try:
        image = Image.open(path).convert("RGB")
    except Exception as exc:
        return PageResult(page=1, text="", method="ocr", error=f"could not open image: {exc}")
    return _ocr_then_vision(image, 1)


# ---------------------------------------------------------------- plain-text / office formats
def _extract_text_file(path: Path) -> PageResult:
    try:
        text = path.read_text(encoding="utf-8", errors="replace")
    except Exception as exc:
        return PageResult(page=1, text="", method="text_file", error=str(exc))
    return PageResult(page=1, text=text, method="text_file")


def _extract_docx(path: Path) -> PageResult:
    try:
        document = docx.Document(path)
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text for cell in row.cells]
                if any(cell.strip() for cell in cells):
                    parts.append(" | ".join(cells))
        text = "\n".join(parts)
    except Exception as exc:
        return PageResult(page=1, text="", method="docx", error=str(exc))
    return PageResult(page=1, text=text, method="docx")


def _extract_xlsx(path: Path) -> PageResult:
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            lines: list[str] = []
            for sheet_name in workbook.sheetnames:
                lines.append(f"[Sheet: {sheet_name}]")
                for row in workbook[sheet_name].iter_rows(values_only=True):
                    values = ["" if v is None else str(v) for v in row]
                    if any(v.strip() for v in values):
                        lines.append(" | ".join(values))
            text = "\n".join(lines)
        finally:
            workbook.close()
    except Exception as exc:
        return PageResult(page=1, text="", method="xlsx", error=str(exc))
    return PageResult(page=1, text=text, method="xlsx")


def _load_first_page_as_image(path: Path, suffix: str) -> Image.Image:
    if suffix == ".pdf":
        try:
            doc = pymupdf.open(path)
        except Exception as exc:
            raise DocumentExtractionError("UNSUPPORTED_FILE", f"could not open PDF: {exc}") from exc
        try:
            pixmap = doc[0].get_pixmap(dpi=settings.WB_OCR_DPI)
            return Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
        finally:
            doc.close()
    return Image.open(path).convert("RGB")


def _resolve_source(source: Union[str, Path]) -> Path:
    if isinstance(source, Path):
        return source
    from backend.file_store import file_store  # local import: avoids a module import cycle

    maybe_path = file_store.get_path(source)
    return maybe_path if maybe_path is not None else Path(source)


def extract(source: Union[str, Path], kind: str = "auto") -> ExtractResult:
    """
    Extracts per-page text from a file (a file_id already known to
    backend.file_store, or a plain path). `kind` is "auto" (default: text
    layer / OCR / vision, in that order) or "pid" (skip straight to tiled
    vision extraction for a P&ID drawing).
    """
    path = _resolve_source(source)
    if not path.exists():
        raise DocumentExtractionError("FILE_NOT_FOUND", f"no such file: {path}")

    file_bytes = path.read_bytes()
    cache_key = _cache_key(file_bytes, kind=kind)
    cached = _load_cache(cache_key)
    if cached is not None:
        return cached

    suffix = path.suffix.lower()

    if kind == "pid":
        image = _load_first_page_as_image(path, suffix)
        pages = _vision_extract_pid_tiles(image, purpose="document_extract_pid")
    elif suffix == ".pdf":
        pages = _extract_pdf_pages(path)
    elif suffix in (".png", ".jpg", ".jpeg"):
        pages = [_extract_single_image_page(path)]
    elif suffix in _TEXT_FILE_SUFFIXES:
        pages = [_extract_text_file(path)]
    elif suffix == ".docx":
        pages = [_extract_docx(path)]
    elif suffix == ".xlsx":
        pages = [_extract_xlsx(path)]
    else:
        raise DocumentExtractionError("UNSUPPORTED_FILE", f"cannot extract from {suffix!r}")

    result = ExtractResult(pages=pages)
    _save_cache(cache_key, result)
    return result
