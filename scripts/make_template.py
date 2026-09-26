"""
Build templates/approval_note.docx (WB_TEMPLATES_DIR) with python-docx.

    python scripts/make_template.py

The template holds the fixed layout; backend/tools/office.py fills it. Every
value is a placeholder paragraph or table cell such as {{subject}}; the
findings and SOP tables have ONE placeholder row that office.py copies once
per item. No logos or images.
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from backend.settings import settings  # noqa: E402
from backend.tools.office import DRAFT_FOOTER, TEMPLATE_NAME  # noqa: E402

_REPO_ROOT = Path(__file__).resolve().parent.parent

HEADER_FILL = "1F3864"      # dark blue header rows, white text
LABEL_FILL = "D9E2F3"       # light blue label cells
ACCENT = RGBColor(0x1F, 0x38, 0x64)

FINDINGS_HEADER = ("S.No", "Item", "Observation", "Severity", "Source page")
FINDINGS_ROW = ("{{finding.sno}}", "{{finding.item}}", "{{finding.observation}}",
                "{{finding.severity}}", "{{finding.source_page}}")
SOP_HEADER = ("S.No", "Document", "Page")
SOP_ROW = ("{{sop.sno}}", "{{sop.document}}", "{{sop.page}}")
APPROVAL_ROLES = ("Prepared by", "Reviewed by", "Approved by")
APPROVAL_LINES = ("Name", "Designation", "Signature", "Date")


def _shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def _cell_text(cell, text: str, *, bold: bool = False, white: bool = False, size: int = 10) -> None:
    para = cell.paragraphs[0]
    run = para.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if white:
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def _header_row(table, labels) -> None:
    row = table.rows[0]
    _repeat_header(row)
    for cell, label in zip(row.cells, labels):
        _cell_text(cell, label, bold=True, white=True)
        _shade(cell, HEADER_FILL)


def _field(paragraph, instr: str) -> None:
    """Insert a Word field (PAGE / NUMPAGES) that Word computes when the file opens."""
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        elif text == instr:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = f" {instr} "
        else:
            el = OxmlElement("w:t")
            el.text = text
        run._r.append(el)
    run.font.size = Pt(8)


def _footer(section) -> None:
    footer = section.footer
    para = footer.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(DRAFT_FOOTER)
    run.font.size = Pt(8)
    run.italic = True
    run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    pages = footer.add_paragraph()
    pages.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label = pages.add_run("Page ")
    label.font.size = Pt(8)
    _field(pages, "PAGE")
    of = pages.add_run(" of ")
    of.font.size = Pt(8)
    _field(pages, "NUMPAGES")


def _heading(doc, text: str) -> None:
    heading = doc.add_heading(text, level=1)
    for run in heading.runs:
        run.font.color.rgb = ACCENT


def build(path: Path) -> Path:
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)     # A4
    section.left_margin = section.right_margin = Cm(2.0)
    section.top_margin = section.bottom_margin = Cm(2.0)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    _footer(section)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("INSPECTION APPROVAL NOTE")
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = ACCENT

    meta = doc.add_table(rows=3, cols=2)
    meta.style = "Table Grid"
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row, (label, key) in zip(meta.rows, (("Ref no", "ref_no"), ("Date", "date"), ("Subject", "subject"))):
        _cell_text(row.cells[0], label, bold=True, size=11)
        _shade(row.cells[0], LABEL_FILL)
        row.cells[0].width = Cm(4)
        row.cells[1].paragraphs[0].add_run("{{" + key + "}}")
        row.cells[1].width = Cm(13)

    _heading(doc, "Background")
    doc.add_paragraph("{{background}}")

    _heading(doc, "Findings")
    findings = doc.add_table(rows=2, cols=len(FINDINGS_HEADER))
    findings.style = "Table Grid"
    _header_row(findings, FINDINGS_HEADER)
    for cell, placeholder in zip(findings.rows[1].cells, FINDINGS_ROW):
        cell.paragraphs[0].add_run(placeholder).font.size = Pt(10)
    for col, width in zip(findings.columns, (Cm(1.2), Cm(4.0), Cm(7.8), Cm(2.2), Cm(1.8))):
        for cell in col.cells:
            cell.width = width

    _heading(doc, "SOP references")
    sops = doc.add_table(rows=2, cols=len(SOP_HEADER))
    sops.style = "Table Grid"
    _header_row(sops, SOP_HEADER)
    for cell, placeholder in zip(sops.rows[1].cells, SOP_ROW):
        cell.paragraphs[0].add_run(placeholder).font.size = Pt(10)
    for col, width in zip(sops.columns, (Cm(1.2), Cm(13.0), Cm(2.8))):
        for cell in col.cells:
            cell.width = width
    sop_note = doc.add_paragraph().add_run("{{sop_note}}")   # removed by office.py unless refs were auto-added
    sop_note.italic = True
    sop_note.font.size = Pt(9)

    _heading(doc, "Recommendation")
    doc.add_paragraph("{{recommendation}}")

    _heading(doc, "Cost implication")
    doc.add_paragraph("{{cost_implication}}")

    _heading(doc, "Approval")
    approval = doc.add_table(rows=1 + len(APPROVAL_LINES), cols=1 + len(APPROVAL_ROLES))
    approval.style = "Table Grid"
    _header_row(approval, ("",) + APPROVAL_ROLES)
    for row, label in zip(approval.rows[1:], APPROVAL_LINES):
        _cell_text(row.cells[0], label, bold=True)
        _shade(row.cells[0], LABEL_FILL)
        if label == "Signature":
            row.height = Cm(1.4)
    approval.rows[1].cells[1].paragraphs[0].add_run("{{prepared_by}}").font.size = Pt(10)

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def template_path() -> Path:
    root = settings.WB_TEMPLATES_DIR
    return (root if root.is_absolute() else _REPO_ROOT / root) / TEMPLATE_NAME


if __name__ == "__main__":
    print(f"wrote {build(template_path())}")
