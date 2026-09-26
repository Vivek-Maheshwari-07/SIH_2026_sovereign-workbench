"""
B9 tests: scripts/e2e_run.py (pass/fail logic, Ctrl+C report, data from ui/scenarios.py and
demo/expected.md) against a mocked HTTP backend (httpx.MockTransport + the real ApiClient),
the PowerShell scripts parse, and the OLLAMA_NO_CLOUD doctor check.
"""
from __future__ import annotations

import json
import re
import shutil
import subprocess
import sys
from datetime import datetime, timezone
from pathlib import Path

import httpx
import pytest
from fake_client import docx_bytes, xlsx_bytes

from shared.contracts import CONTRACT_VERSION
from ui import scenarios as ui_scenarios
from ui.api_client import ApiClient

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO / "scripts"))
import e2e_run  # noqa: E402

NOW = datetime(2026, 9, 26, tzinfo=timezone.utc).isoformat()
HDR = {"X-Contract-Version": CONTRACT_VERSION}
EXPECTED = (REPO / "demo" / "expected.md").read_text(encoding="utf-8")
A_SECTION = e2e_run.expected_section(EXPECTED, "scenario_a_report_1.pdf")
C_SECTION = e2e_run.expected_section(EXPECTED, "scenario_c_pid_generated.png")


# ---------------------------------------------------------------- deliverables made in the test
def approval_note(severities=("High", "High", "Medium", "Medium", "Low"), cost="Rs 4,50,000",
                  sops=("osha_h2s_fact_sheet.pdf", "osha_lockout_tagout.pdf")) -> bytes:
    from docx import Document
    import io

    doc = Document()
    doc.add_paragraph("INSPECTION APPROVAL NOTE")
    table = doc.add_table(rows=1, cols=4)
    for cell, text in zip(table.rows[0].cells, ["S.No", "Item", "Observation", "Severity"]):
        cell.text = text
    items = ["Shell course 2, north side", "Bottom plate near sump", "Inlet nozzle N2 and mixer MX-104",
             "Confined space entry permit", "Roof handrail and stairway"]
    for i, (item, sev) in enumerate(zip(items, severities), start=1):
        row = table.add_row().cells
        row[0].text, row[1].text, row[2].text, row[3].text = str(i), item, "obs", sev
    doc.add_paragraph("SOP references")
    for sop in sops:
        doc.add_paragraph(f"{sop}, p.2")
    if not sops:
        doc.add_paragraph("Not applicable")
    doc.add_paragraph("Cost implication")
    doc.add_paragraph(cost)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


ALL_TAGS = ["T-201", "LT-201", "XV-201", "P-201A", "P-201B", "PT-202", "FT-201", "E-201", "TT-203", "V-201",
            "PSV-201", "LT-202"]


# Types as the backend writes them (it uses "Instrument" for all transmitters).
GOOD_TYPES = {"T-201": "Tank", "LT-201": "Instrument", "XV-201": "Valve", "P-201A": "Pump", "P-201B": "Pump",
              "PT-202": "Instrument", "FT-201": "Instrument", "E-201": "Heat exchanger", "TT-203": "Instrument",
              "V-201": "Vessel", "PSV-201": "Pressure safety valve", "LT-202": "Instrument"}


def tag_list(tags, types=None, with_type_column=True) -> bytes:
    import io

    import pandas as pd

    types = {**GOOD_TYPES, **(types or {})}
    data = {"Tag": list(tags)}
    if with_type_column:
        data["Equipment type"] = [types.get(t.replace(" ", "-"), "Instrument") for t in tags]
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        pd.DataFrame(data).to_excel(xw, sheet_name="Tags", index=False)
    return buf.getvalue()


# ---------------------------------------------------------------- mock backend
class MockBackend:
    """Just enough of the API for e2e_run: files, tasks, events, final state, artifacts, network."""

    def __init__(self, status="succeeded", answer="t = 10 * 200 / (2 * 138) = 7.246 mm",
                 files=None, net=(0, 0), polls_until_done=2):
        self.status, self.answer = status, answer
        self.files = files if files is not None else {"solution.py": ("py", b"print(7.246)\n")}
        self.net = list(net)
        self.polls_until_done = polls_until_done
        self.created: list[dict] = []
        self.cancelled: list[str] = []
        self.polls: dict[str, int] = {}

    def artifacts(self, task_id: str) -> list[dict]:
        return [{"artifact_id": f"a_{i:012x}", "task_id": task_id, "filename": name, "kind": kind,
                 "size_bytes": len(content), "created_at": NOW, "download_url": f"/api/artifacts/a_{i:012x}"}
                for i, (name, (kind, content)) in enumerate(self.files.items(), start=1)]

    def handler(self, req: httpx.Request) -> httpx.Response:
        path = req.url.path
        if path == "/api/health":
            return httpx.Response(200, headers=HDR, json={
                "status": "ok", "contract_version": CONTRACT_VERSION, "mock": True, "ollama_ok": True,
                "sandbox_ok": True, "tesseract_ok": True, "kb_chunks": 1, "models": [], "time": NOW})
        if path == "/api/network/status":
            value = self.net.pop(0) if len(self.net) > 1 else self.net[0]
            return httpx.Response(200, headers=HDR, json={
                "checked_at": NOW, "external_count": 0, "external_seen_since_start": value,
                "total_connections": 0, "connections": []})
        if path == "/api/files":
            return httpx.Response(200, headers=HDR, json={"file_id": "f_000000000001", "filename": "x",
                                                          "mime_type": "x", "size_bytes": 1, "is_image": False})
        if path == "/api/tasks" and req.method == "POST":
            body = json.loads(req.content)
            self.created.append(body)
            return httpx.Response(202, headers=HDR, json={"task_id": f"t_{len(self.created):012x}", "status": "queued"})
        m = re.fullmatch(r"/api/tasks/(t_\w+)(/events|/cancel)?", path)
        if m:
            task_id, tail = m.group(1), m.group(2)
            if tail == "/events":
                n = self.polls[task_id] = self.polls.get(task_id, 0) + 1
                done = n >= self.polls_until_done
                return httpx.Response(200, headers=HDR, json={"task_id": task_id, "events": [],
                                                              "next_seq": 0, "done": done})
            if tail == "/cancel":
                self.cancelled.append(task_id)
            body = self.created[int(task_id[2:], 16) - 1]
            error = None if self.status == "succeeded" else {"code": "AGENT_TIMEOUT", "message": "too slow"}
            return httpx.Response(200, headers=HDR, json={
                "task_id": task_id, "status": self.status, "mode": body["mode"], "scenario": body["scenario"],
                "message": body["message"], "final_answer": self.answer if self.status == "succeeded" else None,
                "artifacts": self.artifacts(task_id), "error": error, "created_at": NOW, "elapsed_s": 12.5})
        m = re.fullmatch(r"/api/artifacts/a_(\w+)", path)
        if m:
            name = list(self.files)[int(m.group(1), 16) - 1]
            return httpx.Response(200, headers={**HDR, "content-disposition": f'attachment; filename="{name}"'},
                                  content=self.files[name][1])
        return httpx.Response(404, headers=HDR, json={"error": {"code": "BAD_REQUEST", "message": path}})


def client_for(backend: MockBackend) -> ApiClient:
    return ApiClient(base_url="http://127.0.0.1:9", transport=httpx.MockTransport(backend.handler))


def run_main(backend: MockBackend, tmp_path: Path, *args: str, monkeypatch=None) -> int:
    if monkeypatch is not None:
        monkeypatch.setattr(e2e_run.time, "sleep", lambda s: None)
    return e2e_run.main([*args, "--out", str(tmp_path / "soak")], client=client_for(backend))


@pytest.fixture(autouse=True)
def _no_sleep(monkeypatch):
    monkeypatch.setattr(e2e_run.time, "sleep", lambda s: None)


# ---------------------------------------------------------------- data sources
def test_scenarios_come_from_ui_scenarios():
    assert e2e_run.SCENARIOS is ui_scenarios.SCENARIOS
    picked = e2e_run.select_scenarios("c, a")
    assert [s.work_order for s in picked] == ["WO-C", "WO-A"]
    assert picked[1].prompt == ui_scenarios.get_scenario(ui_scenarios.Scenario.INSPECTION_NOTE).prompt
    with pytest.raises(ValueError):
        e2e_run.select_scenarios("a,z")


def test_expected_values_read_from_expected_md():
    assert e2e_run.expected_findings(A_SECTION)[0] == ("Shell course 2, north side", "High")
    assert [s for _, s in e2e_run.expected_findings(A_SECTION)] == ["High", "High", "Medium", "Medium", "Low"]
    assert e2e_run.expected_cost(A_SECTION) == "Rs 4,50,000"
    assert "osha_lockout_tagout.pdf" in e2e_run.expected_sops(A_SECTION)
    assert e2e_run.expected_min_sops(A_SECTION) == 2          # "must cite at least 2 of these files"
    assert e2e_run.expected_tags(C_SECTION) == ALL_TAGS


def test_request_uses_exact_prompt_and_mode(tmp_path):
    backend = MockBackend()
    assert run_main(backend, tmp_path, "--scenarios", "b", "--mode", "agent") == 0
    scn = ui_scenarios.get_scenario(ui_scenarios.Scenario.CODE_CALC)
    assert backend.created[0] == {"message": scn.prompt, "file_ids": [], "mode": "agent", "scenario": "code_calc"}


# ---------------------------------------------------------------- pass / fail logic
def test_all_pass_exit_0_and_report(tmp_path):
    backend = MockBackend(net=(1, 1))
    assert run_main(backend, tmp_path, "--runs", "2", "--scenarios", "b") == 0
    report = (tmp_path / "soak" / "report.md").read_text(encoding="utf-8")
    raw = json.loads((tmp_path / "soak" / "raw.json").read_text(encoding="utf-8"))
    assert "ALL PASSED" in report and "2/2 passed" in report and "before 1, after 1" in report
    assert raw["all_passed"] is True and len(raw["results"]) == 2
    assert all(r["passed"] and r["artifacts"][0]["downloaded"] > 0 for r in raw["results"])
    assert (tmp_path / "soak" / "artifacts" / "run01_code_calc__solution.py").read_bytes() == b"print(7.246)\n"


def test_wrong_answer_fails(tmp_path):
    backend = MockBackend(answer="t = 8.1 mm")
    assert run_main(backend, tmp_path, "--scenarios", "b") == 1
    raw = json.loads((tmp_path / "soak" / "raw.json").read_text(encoding="utf-8"))
    assert raw["all_passed"] is False
    assert "not in the final answer" in (tmp_path / "soak" / "report.md").read_text(encoding="utf-8")


def test_failed_task_and_empty_artifact_fail(tmp_path):
    assert run_main(MockBackend(status="failed"), tmp_path, "--scenarios", "b") == 1
    assert "AGENT_TIMEOUT: too slow" in (tmp_path / "soak" / "report.md").read_text(encoding="utf-8")
    assert run_main(MockBackend(files={"solution.py": ("py", b"")}), tmp_path, "--scenarios", "b") == 1


def test_new_core_connection_is_flagged(tmp_path):
    run_main(MockBackend(net=(0, 2)), tmp_path, "--scenarios", "b")
    assert "+2 new core connection(s)" in (tmp_path / "soak" / "report.md").read_text(encoding="utf-8")


def test_scenario_a_checks(tmp_path):
    good = MockBackend(files={"note.docx": ("docx", approval_note())})
    assert run_main(good, tmp_path, "--scenarios", "a") == 0
    assert good.created[0]["file_ids"] == ["f_000000000001"]              # demo PDF uploaded first
    checks = e2e_run.check_inspection(A_SECTION, {"n.docx": approval_note(
        severities=("High", "Medium", "High", "Medium", "Low"), cost="To be filled", sops=())})
    by_name = {c.name: c for c in checks}
    assert not by_name["A: severities"].ok and "Bottom plate near sump: Medium (expected High)" in by_name["A: severities"].detail
    assert not by_name["A: cost"].ok and not by_name["A: SOPs cited"].ok


def test_scenario_a_needs_two_sops():
    def sop_check(sops):
        return {c.name: c for c in e2e_run.check_inspection(A_SECTION, {"n.docx": approval_note(sops=sops)})}["A: SOPs cited"]

    one = sop_check(("osha_h2s_fact_sheet.pdf",))
    assert not one.ok and one.detail.startswith("only 1 cited, need 2: osha_h2s_fact_sheet.pdf")
    two = sop_check(("osha_h2s_fact_sheet.pdf", "nsw_hot_work_petroleum.pdf"))
    assert two.ok and two.detail.startswith("2 cited (need 2)")
    assert not sop_check(("made_up_sop.pdf", "another_fake.pdf")).ok     # only files from expected.md count


def test_scenario_c_checks():
    exact = e2e_run.check_tags(C_SECTION, {"t.xlsx": tag_list(ALL_TAGS)})[0]
    assert exact.ok and exact.detail == "12/12 tags"
    invented = e2e_run.check_tags(C_SECTION, {"t.xlsx": tag_list(ALL_TAGS + ["FT"])})[0]
    assert not invented.ok and invented.detail == "12/12 tags; invented FT"   # invented tags fail
    drawing_no = e2e_run.check_tags(C_SECTION, {"t.xlsx": tag_list(ALL_TAGS + ["DEMO-C-001"])})[0]
    assert not drawing_no.ok and "invented DEMO-C-001" in drawing_no.detail
    spaced = e2e_run.check_tags(C_SECTION, {"t.xlsx": tag_list([t.replace("-", " ") for t in ALL_TAGS])})[0]
    assert spaced.ok                                                      # "-" and space count the same
    missing = e2e_run.check_tags(C_SECTION, {"t.xlsx": tag_list(ALL_TAGS[:-1])})[0]
    assert not missing.ok and "11/12" in missing.detail and "LT-202" in missing.detail
    assert not e2e_run.check_tags(C_SECTION, {})[0].ok


def c_checks(xlsx: bytes) -> dict:
    return {c.name: c for c in e2e_run.check_tags(C_SECTION, {"t.xlsx": xlsx})}


def test_expected_types_read_from_expected_md():
    types = e2e_run.expected_tag_types(C_SECTION)
    assert len(types) == 12 and types["V-201"] == "Vessel (separator)" and types["T-201"] == "Tank (crude storage)"


def test_equipment_types_all_correct():
    check = c_checks(tag_list(ALL_TAGS))["C: equipment types"]
    assert check.ok and check.detail == "12/12 types match"


def test_equipment_type_synonyms_pass():
    synonyms = {"T-201": "Storage tank", "XV-201": "On/off valve", "LT-201": "level transmitter",
                "V-201": "Separator", "E-201": "HEAT EXCHANGER", "PSV-201": "Valve", "P-201B": "Pump (standby)"}
    assert c_checks(tag_list(ALL_TAGS, synonyms))["C: equipment types"].ok


def test_wrong_equipment_type_fails_and_names_the_tag():
    check = c_checks(tag_list(ALL_TAGS, {"V-201": "Valve", "P-201A": "Tank"}))["C: equipment types"]
    assert not check.ok
    assert "V-201 expected Vessel (separator), found Valve" in check.detail   # a vessel is not a valve
    assert "P-201A expected Pump, found Tank" in check.detail


def test_missing_type_column_fails():
    check = c_checks(tag_list(ALL_TAGS, with_type_column=False))["C: equipment types"]
    assert not check.ok and "T-201 expected Tank (crude storage), found no type" in check.detail


def test_wrong_type_fails_the_whole_run(tmp_path):
    backend = MockBackend(files={"pid_tags.xlsx": ("xlsx", tag_list(ALL_TAGS, {"XV-201": "Pump"}))})
    assert run_main(backend, tmp_path, "--scenarios", "c") == 1
    report = (tmp_path / "soak" / "report.md").read_text(encoding="utf-8")
    assert "C: equipment types: wrong type: XV-201 expected Shutdown (on/off) valve, actuated, found Pump" in report
    good = MockBackend(files={"pid_tags.xlsx": ("xlsx", tag_list(ALL_TAGS))})
    assert run_main(good, tmp_path, "--scenarios", "c") == 0


def test_equipment_class_rule():
    cls = e2e_run.equipment_class
    assert cls("Tank (crude storage)") == cls("Storage tank") == "tank"
    assert cls("Pressure transmitter (pump discharge)") == "instrument"        # bracket text ignored
    assert cls("Pressure safety valve (on V-201, to flare)") == cls("Control valve") == "valve"
    assert cls("Vessel") != cls("Valve")
    assert cls("Agitator") == "agitator"                                      # unknown: own words only


def test_backend_down_exit_2(tmp_path):
    down = ApiClient(base_url="http://127.0.0.1:9", transport=httpx.MockTransport(
        lambda req: (_ for _ in ()).throw(httpx.ConnectError("refused", request=req))))
    assert e2e_run.main(["--scenarios", "b", "--out", str(tmp_path / "x")], client=down) == 2


# ---------------------------------------------------------------- Ctrl+C
def test_ctrl_c_cancels_task_and_writes_report(tmp_path, monkeypatch):
    backend = MockBackend(polls_until_done=99)
    calls = {"n": 0}

    def sleep(_s):
        calls["n"] += 1
        if calls["n"] == 2:
            raise KeyboardInterrupt

    monkeypatch.setattr(e2e_run.time, "sleep", sleep)
    code = e2e_run.main(["--runs", "3", "--scenarios", "b", "--out", str(tmp_path / "soak")],
                        client=client_for(backend))
    assert code == e2e_run.EXIT_INTERRUPTED
    assert backend.cancelled == ["t_000000000001"]
    report = (tmp_path / "soak" / "report.md").read_text(encoding="utf-8")
    raw = json.loads((tmp_path / "soak" / "raw.json").read_text(encoding="utf-8"))
    assert "interrupted with Ctrl+C" in report and raw["interrupted"] is True and raw["all_passed"] is False
    assert [r["status"] for r in raw["results"]] == ["interrupted"]      # the unfinished run is kept


def test_ctrl_c_between_runs_keeps_finished_results(tmp_path):
    backend = MockBackend()
    soak = e2e_run.Soak(client_for(backend), e2e_run.select_scenarios("b"), 3, e2e_run.TaskMode.GUIDED,
                        tmp_path / "soak", EXPECTED, sleep=lambda s: None, log=lambda m: None)
    real = soak.run_one

    def run_one(run, scn):
        if run == 2:
            raise KeyboardInterrupt
        return real(run, scn)

    soak.run_one = run_one
    soak.run_all()
    soak.write_report()
    assert soak.interrupted and len(soak.results) == 1 and soak.results[0].passed and not soak.all_passed


# ---------------------------------------------------------------- PowerShell scripts parse (no execution)
PS1 = [REPO / "scripts" / "start_demo.ps1", REPO / "scripts" / "stop_demo.ps1"]


@pytest.mark.parametrize("script", PS1, ids=lambda p: p.name)
def test_ps1_is_ascii_without_bom(script):
    data = script.read_bytes()
    assert not data.startswith(b"\xef\xbb\xbf") and max(data) < 128


@pytest.mark.skipif(shutil.which("powershell") is None, reason="Windows PowerShell not available")
@pytest.mark.parametrize("script", PS1, ids=lambda p: p.name)
def test_ps1_parses(script):
    command = ("$e = $null; $t = $null; "
               f"[void][System.Management.Automation.Language.Parser]::ParseFile('{script}', [ref]$t, [ref]$e); "
               "Write-Output $e.Count; $e | ForEach-Object { Write-Output $_.ToString() }")
    out = subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-Command", command],
                         capture_output=True, text=True, timeout=60)
    assert out.stdout.strip().splitlines()[0] == "0", out.stdout + out.stderr


def test_start_demo_covers_required_steps():
    text = PS1[0].read_text(encoding="ascii")
    for needle in (".venv\\Scripts\\python.exe", "OLLAMA_NO_CLOUD", "Ollama cloud disabled: (true|false)",
                   "OLLAMA_KEEP_ALIVE = \"30m\"", "OLLAMA_NUM_PARALLEL = \"1\"", "wb-sandbox.tar",
                   "/api/admin/prewarm", "streamlit", "firewall_outbound_blocked", "READY", "SkipPrewarm",
                   "NoBrowser", "demo_pids.json"):
        assert needle in text, needle


# ---------------------------------------------------------------- doctor: OLLAMA_NO_CLOUD
def test_doctor_no_cloud_check(monkeypatch, capsys):
    import doctor

    monkeypatch.setattr(doctor, "_results", [])
    monkeypatch.setattr(doctor, "user_env_value", lambda name: "1")
    doctor.check_ollama_no_cloud()
    monkeypatch.setattr(doctor, "user_env_value", lambda name: None)
    doctor.check_ollama_no_cloud()
    out = capsys.readouterr().out
    assert "[PASS] OLLAMA_NO_CLOUD is set to 1 (user environment)" in out
    assert "[FAIL] OLLAMA_NO_CLOUD not set" in out and "setx OLLAMA_NO_CLOUD 1" in out
    assert doctor._results == [True, False]
