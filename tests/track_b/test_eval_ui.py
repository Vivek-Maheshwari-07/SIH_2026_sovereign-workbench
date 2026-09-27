"""
Evaluation pass, UI group: Markdown answers, route caption, "Models used" on the router card,
the model registry in the sidebar, example requests, the Audit "Only this job" filter, the
PowerPoint preview, the off-machine backend guard and the extra demo inputs.
"""
from __future__ import annotations

import io
import sys
from pathlib import Path

import pytest
from fake_client import FakeClient, ev, health, ok
from pptx import Presentation
from streamlit.testing.v1 import AppTest

from shared.contracts import ArtifactKind, EventType, ModelInfo, RouteDecision, TaskMode, TaskType
from ui import api_client
from ui.components import artifacts, audit_page, router_badge, timeline
from ui.config import is_local_url
from ui.scenarios import EXAMPLES

REPO = Path(__file__).resolve().parents[2]
APP = str(REPO / "ui" / "app.py")
MODELS = [ModelInfo(id="general", ollama_name="qwen3.5:4b", tasks=[TaskType.DOCUMENT, TaskType.VISION], loaded=True),
          ModelInfo(id="coder", ollama_name="qwen2.5-coder:3b", tasks=[TaskType.CODING]),
          ModelInfo(id="embed", ollama_name="bge-m3")]


def run_app(monkeypatch, fake: FakeClient) -> AppTest:
    monkeypatch.setattr(api_client, "get_client", lambda: fake)
    at = AppTest.from_file(APP, default_timeout=30)
    at.run()
    assert not at.exception, [e.message for e in at.exception]
    return at


def page_html(at: AppTest) -> str:
    return "\n".join(m.value for m in at.markdown)


# ---------------------------------------------------------------- result as Markdown
def test_answer_markdown_keeps_formatting_and_escapes_html_and_dollars():
    text = timeline.answer_markdown("**Limit**: 20 ppm [1]\n- costs $5 <script>x</script>\n\nSources:\n[1] a.pdf")
    assert text.startswith("**Limit**: 20 ppm [1]  \n")          # bold kept, hard line break
    assert "\\$5" in text and "<script>" not in text and "&lt;script&gt;" in text
    assert "\n\nSources:  \n[1] a.pdf  " in text                  # blank line kept as a paragraph break


def test_route_bubble_names_the_task_type_not_the_model_id():
    decision = RouteDecision(task_type="document", model_id="general", ollama_name="qwen3.5:4b", reason="r",
                             layer="rule", confidence=1.0)
    stages = timeline.build_stages([ev(1, EventType.ROUTE, "Routed", {"decision": decision.model_dump(mode="json")})],
                                   [], None)
    assert stages[0].caption == "Route: document"


def test_new_tools_have_short_codes_and_labels():
    for tool in ("answer_question", "create_document", "analyze_table", "inspect_image"):
        assert len(timeline.tool_code(tool)) == 2 and tool in timeline.TOOL_LABELS


# ---------------------------------------------------------------- router card: models used
def test_models_used_counts_llm_calls_in_order_of_first_use():
    events = [ev(1, EventType.LLM_CALL, "p", {"model_id": "general"}),
              ev(2, EventType.LLM_CALL, "c", {"model_id": "coder"}),
              ev(3, EventType.TOOL_CALL, "t", {"tool": "x"}),
              ev(4, EventType.LLM_CALL, "s", {"model_id": "general"})]
    assert router_badge.models_used(events) == [("general", 2), ("coder", 1)]
    decision = RouteDecision(task_type="coding", model_id="coder", ollama_name="qwen2.5-coder:3b", reason="r",
                             layer="rule", confidence=1.0)
    html = router_badge.badge_html(decision, router_badge.models_used(events))
    assert "Models used" in html and "general ×2, coder ×1" in html
    assert "Models used" not in router_badge.badge_html(decision, [])


# ---------------------------------------------------------------- sidebar models + examples
def test_sidebar_lists_the_model_registry(monkeypatch):
    h = health()
    h.models = MODELS
    at = run_app(monkeypatch, FakeClient(ok(h)))
    text = page_html(at)
    assert "Models" in text and "qwen2.5-coder:3b" in text and "embeddings (search)" in text
    assert text.count('<i class="b-ok m">') == 1 and text.count('<i class="b-off m">') == 2
    assert "in memory" in text and "on disk" in text


def test_examples_show_when_idle_and_start_agent_jobs(monkeypatch):
    fake = FakeClient()
    at = run_app(monkeypatch, fake)
    keys = [b.key for b in at.button if b.key and b.key.startswith("ex_")]
    assert keys == [f"ex_{e.key}" for e in EXAMPLES]
    at.button(key="ex_table").click().run()
    assert not at.exception
    req = fake.created[0]
    assert req.mode == TaskMode.AGENT and req.scenario is None
    assert req.message == next(e for e in EXAMPLES if e.key == "table").prompt
    assert [u[0] for u in fake.uploads] == ["equipment_thickness.csv"]


def test_example_demo_files_exist():
    for example in EXAMPLES:
        assert example.demo_file is None or example.demo_file.is_file(), example.demo_file


def test_off_machine_backend_is_refused(monkeypatch):
    fake = FakeClient()
    fake.base_url = "http://10.20.30.40:8000"
    at = run_app(monkeypatch, fake)
    text = page_html(at)
    assert "Backend address is not on this machine" in text
    assert fake.created == [] and "Work orders" not in text


@pytest.mark.parametrize("url,local", [("http://127.0.0.1:8000", True), ("http://localhost:8001", True),
                                       ("http://[::1]:8000", True), ("http://10.0.0.2:8000", False),
                                       ("https://api.example.com", False), ("", False)])
def test_is_local_url(url, local):
    assert is_local_url(url) is local


# ---------------------------------------------------------------- audit filter
def test_effective_task_id():
    assert audit_page.effective_task_id("", "t_current", True) == "t_current"
    assert audit_page.effective_task_id("", "t_current", False) is None
    assert audit_page.effective_task_id(" t_typed ", "t_current", True) == "t_typed"
    assert audit_page.effective_task_id("", None, True) is None


# ---------------------------------------------------------------- PowerPoint preview
def test_pptx_outline_lists_slide_titles_and_bullets():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "H2S briefing"
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Hazards"
    body = slide.placeholders[1].text_frame
    body.text = "Toxic gas"
    for extra in ("Ceiling 20 ppm", "Wear a monitor", "Fourth bullet"):
        body.add_paragraph().text = extra
    buf = io.BytesIO()
    prs.save(buf)
    outline = artifacts.build_preview(ArtifactKind.PPTX, buf.getvalue())
    assert outline[0][0] == "H2S briefing"
    assert outline[1] == ("Hazards", ["Toxic gas", "Ceiling 20 ppm", "Wear a monitor"])


# ---------------------------------------------------------------- extra demo inputs
def test_make_extra_inputs(tmp_path):
    sys.path.insert(0, str(REPO / "demo" / "tools"))
    try:
        import make_extra_inputs as gen
    finally:
        sys.path.pop(0)
    image = gen.make_note(["Line one", "Line two"])
    assert image.mode == "RGB" and image.width > 1400 and image.height > 1000
    path = tmp_path / "e.csv"
    gen.write_csv(path)
    lines = path.read_text(encoding="utf-8").splitlines()
    assert lines[0] == ",".join(gen.CSV_HEADER) and len(lines) == len(gen.EQUIPMENT_ROWS) + 1
