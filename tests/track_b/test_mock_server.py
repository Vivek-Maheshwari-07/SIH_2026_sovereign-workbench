"""mock/mock_server.py in-process: scripted runs per scenario, cancel, artifacts, errors."""
from __future__ import annotations

import io
import time

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from mock import mock_server
from shared.contracts import CONTRACT_VERSION, EventsPage, TaskState, TaskStatus


@pytest.fixture
def client(monkeypatch):
    monkeypatch.setattr(mock_server.state, "release_s", 0.01)
    with TestClient(mock_server.app) as c:
        yield c


def run_to_end(client: TestClient, body: dict) -> tuple[TaskState, EventsPage]:
    task_id = client.post("/api/tasks", json=body).json()["task_id"]
    for _ in range(200):
        page = EventsPage.model_validate(client.get(f"/api/tasks/{task_id}/events", params={"after": 0}).json())
        if page.done:
            break
        time.sleep(0.02)
    return TaskState.model_validate(client.get(f"/api/tasks/{task_id}").json()), page


@pytest.mark.parametrize("scenario,kind", [("inspection_note", "docx"), ("code_calc", "py"), ("pid_tags", "xlsx")])
def test_each_scenario_runs_to_a_downloadable_artifact(client, scenario, kind):
    state, page = run_to_end(client, {"message": "demo", "mode": "guided", "scenario": scenario})
    assert state.status == TaskStatus.SUCCEEDED and state.final_answer
    assert [e.seq for e in page.events] == list(range(1, len(page.events) + 1))
    assert page.events[0].type.value == "route" and page.events[-1].type.value == "final"
    art = state.artifacts[0]
    assert art.kind.value == kind
    resp = client.get(art.download_url)
    assert resp.status_code == 200 and len(resp.content) == art.size_bytes
    assert resp.headers["X-Contract-Version"] == CONTRACT_VERSION
    if kind == "docx":
        assert "INSPECTION APPROVAL NOTE" in Document(io.BytesIO(resp.content)).paragraphs[0].text
    if kind == "xlsx":
        assert load_workbook(io.BytesIO(resp.content)).active["A2"].value == "P-201A"


def test_free_agent_request_answers_with_sources(client):
    state, _ = run_to_end(client, {"message": "What is the H2S limit?", "mode": "agent"})
    assert state.status == TaskStatus.SUCCEEDED and "Sources:" in state.final_answer and not state.artifacts


def test_cancel_stops_the_run(client, monkeypatch):
    monkeypatch.setattr(mock_server.state, "release_s", 60)
    task_id = client.post("/api/tasks", json={"message": "x", "mode": "guided", "scenario": "code_calc"}).json()["task_id"]
    state = TaskState.model_validate(client.post(f"/api/tasks/{task_id}/cancel").json())
    assert state.status == TaskStatus.CANCELLED and state.error.code == "CANCELLED"
    page = EventsPage.model_validate(client.get(f"/api/tasks/{task_id}/events").json())
    assert page.done and page.events[-1].type.value == "error"


def test_errors_use_the_contract_shape(client):
    bad = client.post("/api/files", files={"file": ("x.exe", b"MZ", "application/octet-stream")})
    assert bad.status_code == 415 and bad.json()["error"]["code"] == "UNSUPPORTED_FILE"
    assert client.get("/api/tasks/t_000000000000").json()["error"]["code"] == "TASK_NOT_FOUND"
    assert client.post("/api/tasks", json={"message": "x", "mode": "guided"}).json()["error"]["code"] == "BAD_REQUEST"
    assert client.get("/api/health").json()["mock"] is True


def test_probe_never_connects(client, monkeypatch):
    import socket

    monkeypatch.setattr(socket, "create_connection", lambda *a, **k: pytest.fail("mock must not connect"))
    result = client.post("/api/network/probe", json={"target": "https://www.google.com"}).json()
    assert result["reachable"] is False
