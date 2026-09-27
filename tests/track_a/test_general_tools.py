"""
Evaluation pass, group 2: the general agent tools (answer_question, create_document,
analyze_table, inspect_image), the Office writers behind them (make_report, make_table_excel)
and the sandbox's result-file collection. The LLM and (except one live test) the sandbox are faked.
"""
from __future__ import annotations

import io
import tarfile
from io import BytesIO
from types import SimpleNamespace

import pytest
from docx import Document
from openpyxl import load_workbook
from PIL import Image
from pptx import Presentation

from backend import agent_tools, llm_client
from backend.agent_tools import DocEntry, ToolContext, ToolError
from backend.file_store import file_store
from backend.llm_client import ChatResult, JsonResult
from backend.settings import settings
from backend.tools import office, sandbox
from backend.tools.documents import PageResult
from backend.tools.sandbox import SandboxResult, run_in_sandbox, sandbox_available
from shared.contracts import ArtifactKind, KBHit

H2S = KBHit(text="H2S is toxic. The OSHA ceiling limit is 20 ppm; evacuate at 100 ppm.",
            source="osha_h2s_fact_sheet.pdf", page=1, score=0.66)
WEAK = KBHit(text="Unrelated hot work text.", source="nsw_hot_work.pdf", page=9, score=0.40)


@pytest.fixture(autouse=True)
def isolated(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "WB_WORKSPACE_DIR", tmp_path / "workspace")
    monkeypatch.setattr(settings, "WB_CACHE_DIR", tmp_path / "cache")
    monkeypatch.setattr(agent_tools.knowledge, "search", lambda query, top_k=4: [H2S, WEAK])


def make_ctx() -> tuple[ToolContext, list]:
    events: list = []
    ctx = ToolContext(task_id="t_gen000000000", emit=lambda *a, **k: events.append(a), add_artifact=lambda a: None,
                      task_message="the user's request")
    return ctx, events


def fake_chat(monkeypatch, replies: list[str]) -> list[dict]:
    """llm_client.chat returns `replies` in order; every call's kwargs are recorded."""
    calls: list[dict] = []

    def chat(model, messages, tools=None, images=None, *, purpose="chat", timeout_s=None):
        calls.append({"model": model, "messages": messages, "images": images, "purpose": purpose})
        return ChatResult(text=replies[min(len(calls), len(replies)) - 1], tokens_out=5)

    monkeypatch.setattr(llm_client, "chat", chat)
    return calls


def png_bytes() -> bytes:
    buf = BytesIO()
    Image.new("RGB", (64, 48), "grey").save(buf, format="PNG")
    return buf.getvalue()


# ---------------------------------------------------------------- citations and ids
def test_cite_sources_lists_only_cited_sources_in_order():
    text = agent_tools.cite_sources("Limit is 20 ppm [2]. Evacuate [2][1].", ["a.pdf, p.1", "b.pdf, p.4", "c.pdf"])
    assert text.endswith("Sources:\n[1] a.pdf, p.1\n[2] b.pdf, p.4")


def test_cite_sources_without_citations_lists_everything_consulted():
    assert agent_tools.cite_sources("No citation.", ["a.pdf"]).endswith("Sources consulted: a.pdf")
    assert agent_tools.cite_sources("Out of range [7].", []) == "Out of range [7]."


def test_humanize_ids_replaces_known_ids_only():
    ctx, _ = make_ctx()
    ctx.scratchpad.kb_hits["kb_2"] = H2S
    ctx.scratchpad.docs["doc_1"] = DocEntry("doc_1", "f_x", "report.pdf", "auto", [])
    ctx.scratchpad.answers["ans_1"] = "The limit is 20 ppm."
    text = agent_tools.humanize_ids(ctx, "See kb_2 and doc_1. ans_1 Unknown kb_9.")
    assert text == "See osha_h2s_fact_sheet.pdf, p.1 and report.pdf. The limit is 20 ppm. Unknown kb_9."


def test_finish_answer_has_no_internal_ids():
    ctx, _ = make_ctx()
    ctx.scratchpad.kb_hits["kb_1"] = H2S
    outcome = agent_tools.execute_tool(ctx, "finish", {"answer": "Details are in kb_1."})
    assert outcome.finish_answer == "Details are in osha_h2s_fact_sheet.pdf, p.1."


# ---------------------------------------------------------------- answer_question
def test_answer_question_gives_the_model_full_passages_and_adds_sources(monkeypatch):
    calls = fake_chat(monkeypatch, ["The ceiling limit is 20 ppm [1]."])
    ctx, _ = make_ctx()
    outcome = agent_tools.execute_tool(ctx, "answer_question", {"question": "What is the H2S limit?"})
    assert outcome.ok
    prompt = calls[0]["messages"][-1]["content"]
    assert H2S.text in prompt and WEAK.text not in prompt           # full passage, weak hit dropped
    answer = ctx.scratchpad.answers["ans_1"]
    assert answer.endswith("Sources:\n[1] osha_h2s_fact_sheet.pdf, p.1")
    assert "ans_1" in outcome.summary


def test_answer_question_uses_documents_read_earlier(monkeypatch):
    calls = fake_chat(monkeypatch, ["Five findings [1]."])
    ctx, _ = make_ctx()
    ctx.scratchpad.docs["doc_1"] = DocEntry("doc_1", "f_x", "report.pdf", "auto",
                                            [PageResult(page=1, text="Findings table with five rows", method="ocr")])
    agent_tools.answer_question(ctx, "Summarise the report", use_knowledge=False)
    prompt = calls[0]["messages"][-1]["content"]
    assert "[1] report.pdf" in prompt and "Findings table with five rows" in prompt
    assert "osha" not in prompt


def test_answer_question_says_so_when_nothing_is_relevant(monkeypatch):
    monkeypatch.setattr(agent_tools.knowledge, "search", lambda query, top_k=4: [WEAK])
    fake_chat(monkeypatch, ["must not be called"])
    ctx, _ = make_ctx()
    outcome = agent_tools.answer_question(ctx, "What is the weather?")
    assert outcome.ok and "does not cover it" in outcome.summary and not ctx.scratchpad.answers


# ---------------------------------------------------------------- create_document
def fake_outline(monkeypatch) -> list:
    seen: list = []

    def chat_json_meta(model, messages, schema, **kwargs):
        seen.append(messages)
        value = schema.model_validate({"sections": [
            {"heading": "Hazards", "bullets": ["H2S is toxic.", "Ceiling limit 20 ppm."]},
            {"heading": "Actions", "bullets": ["Wear a personal monitor."]},
        ]})
        return JsonResult(value=value, tokens_out=40, duration_ms=10)

    monkeypatch.setattr(llm_client, "chat_json_meta", chat_json_meta)
    return seen


def test_create_document_pptx_has_content_and_sources_slide(monkeypatch):
    seen = fake_outline(monkeypatch)
    ctx, events = make_ctx()
    ctx.scratchpad.answers["ans_1"] = "The ceiling limit is 20 ppm [1]."
    outcome = agent_tools.execute_tool(ctx, "create_document", {"format": "pptx", "title": "H2S briefing"})
    assert outcome.ok, outcome.summary
    art = ctx.artifacts[0]
    assert art.kind == ArtifactKind.PPTX and art.filename.endswith(".pptx")
    path, _ = office.artifact_path(art.artifact_id)
    titles = [s.shapes.title.text for s in Presentation(str(path)).slides]
    assert titles == ["H2S briefing", "Hazards", "Actions", "Sources"]
    assert "The ceiling limit is 20 ppm" in seen[0][-1]["content"]


def test_create_document_docx_opens_in_word_format(monkeypatch):
    fake_outline(monkeypatch)
    ctx, _ = make_ctx()
    outcome = agent_tools.execute_tool(ctx, "create_document", {"format": "word", "title": "H2S note"})
    assert outcome.ok
    path, art = office.artifact_path(ctx.artifacts[0].artifact_id)
    text = [p.text for p in Document(str(path)).paragraphs]
    assert art.kind == ArtifactKind.DOCX
    assert text[0] == "H2S note" and "Hazards" in text and "Ceiling limit 20 ppm." in text
    assert "osha_h2s_fact_sheet.pdf, p.1" in text                  # sources section


def test_create_document_rejects_unknown_format():
    ctx, _ = make_ctx()
    outcome = agent_tools.execute_tool(ctx, "create_document", {"format": "pdf", "title": "x"})
    assert not outcome.ok and "format must be 'docx'" in outcome.summary


# ---------------------------------------------------------------- Office writers
def test_make_report_and_cell_values():
    art = office.make_report("Plan", [("Scope", ["One.", "Two."]), ("Empty", [])], task_id="t_x")
    path, _ = office.artifact_path(art.artifact_id)
    paras = [p.text for p in Document(str(path)).paragraphs]
    assert paras[:2] == ["Plan", paras[1]] and "One." in paras and office.NOT_APPLICABLE in paras
    assert "2 sections" in art.preview
    assert [office.cell_value(v) for v in ["12", "-3.5", "1e3", "V-101", " 7 ", "15.000000000000002"]] == \
        [12, -3.5, 1000.0, "V-101", 7, 15.0]                    # float noise from pandas is dropped


def test_make_table_excel_has_result_and_input_sheets():
    art = office.make_table_excel(["tag", "loss_pct"], [["V-101", "15.0"], ["E-201"]],
                                  input_header=["tag", "t"], input_rows=[["V-101", "10.2"]], task_id="t_x")
    path, _ = office.artifact_path(art.artifact_id)
    wb = load_workbook(path)
    assert wb.sheetnames == ["Result", "Input data"]
    assert [c.value for c in wb["Result"][2]] == ["V-101", 15.0]
    assert wb["Result"]["B3"].value in (None, "")                   # short row padded
    assert "2 result rows x 2 columns" in art.preview


# ---------------------------------------------------------------- analyze_table helpers
def test_extract_python_and_check_script():
    reply = "Here:\n```python\nimport pandas as pd\ndf = pd.read_csv('input.csv')\ndf.to_csv('output.csv')\n```\nDone"
    code = agent_tools.extract_python(reply)
    assert code.startswith("import pandas") and agent_tools.check_table_script(code) == []
    assert agent_tools.extract_python("print(1)") == "print(1)"
    assert "syntax error" in agent_tools.check_table_script("def (:")[0]
    assert len(agent_tools.check_table_script("print(1)")) == 2


def test_describe_table_and_csv_rows(tmp_path):
    path = tmp_path / "d.csv"
    path.write_text("tag,thk\nV-1,10.5\nV-2,9\n", encoding="utf-8")
    df = agent_tools.load_table(path)
    text = agent_tools.describe_table(df)
    assert "2 rows" in text and "- thk (float64)" in text and "V-1,10.5" in text
    assert agent_tools.csv_rows(b"a,b\n1,2\n") == (["a", "b"], [["1", "2"]])
    assert agent_tools.csv_rows(b"") == ([], [])


# ---------------------------------------------------------------- analyze_table
GOOD_SCRIPT = ("```python\nimport pandas as pd\ndf = pd.read_csv('input.csv')\n"
               "df['loss'] = 1\ndf.to_csv('output.csv', index=False)\nprint('done')\n```")


def table_file() -> str:
    return file_store.save("equipment.csv", b"tag,thk\nV-101,10.2\nE-201,6.8\n", "text/csv").file_id


def test_analyze_table_retries_then_saves_excel_and_script(monkeypatch):
    calls = fake_chat(monkeypatch, [GOOD_SCRIPT])
    runs = []

    def fake_run(files, command, *, timeout_s=None, collect=None):
        runs.append(files)
        if len(runs) == 1:
            return SandboxResult(exit_code=1, stdout="", stderr="KeyError: 'thickness'", timed_out=False, duration_ms=5)
        return SandboxResult(exit_code=0, stdout="2 items analysed", stderr="", timed_out=False, duration_ms=5,
                             files={"output.csv": b"tag,loss\nV-101,1\nE-201,1\n"})

    monkeypatch.setattr(agent_tools, "run_in_sandbox", fake_run)
    ctx, events = make_ctx()
    outcome = agent_tools.execute_tool(ctx, "analyze_table", {"file_id": table_file(), "request": "wall loss"})
    assert outcome.ok, outcome.summary
    assert runs[0]["input.csv"].startswith("tag,thk") and "analysis.py" in runs[0]
    assert "KeyError" in calls[1]["messages"][-1]["content"]              # the error went back to the model
    assert [a.kind for a in ctx.artifacts] == [ArtifactKind.XLSX, ArtifactKind.PY]
    assert "2 rows x 2 columns" in outcome.summary and "2 items analysed" in outcome.summary


def test_analyze_table_gives_up_after_the_attempt_limit(monkeypatch):
    fake_chat(monkeypatch, ["no code here"])
    monkeypatch.setattr(agent_tools, "run_in_sandbox", lambda *a, **k: pytest.fail("must not run bad code"))
    ctx, _ = make_ctx()
    outcome = agent_tools.execute_tool(ctx, "analyze_table", {"file_id": table_file(), "request": "x"})
    assert not outcome.ok and outcome.error_code == "BAD_MODEL_OUTPUT"


def test_analyze_table_refuses_non_table_files():
    ctx, _ = make_ctx()
    text_id = file_store.save("notes.txt", b"hello", "text/plain").file_id
    with pytest.raises(ToolError, match="CSV or Excel"):
        agent_tools.analyze_table(ctx, text_id, "x")


# ---------------------------------------------------------------- inspect_image
def test_inspect_image_sends_the_picture_to_the_vision_model(monkeypatch):
    calls = fake_chat(monkeypatch, ["- Flange shows rust staining"])
    ctx, _ = make_ctx()
    image_id = file_store.save("flange.png", png_bytes(), "image/png").file_id
    outcome = agent_tools.execute_tool(ctx, "inspect_image", {"file_id": image_id, "question": "Is it corroded?"})
    assert outcome.ok
    assert calls[0]["images"] and calls[0]["images"][0][:4] == b"\x89PNG"
    assert "Is it corroded?" in calls[0]["messages"][0]["content"]
    assert "flange.png" in ctx.scratchpad.answers["ans_1"]


def test_inspect_image_rejects_text_files():
    ctx, _ = make_ctx()
    text_id = file_store.save("notes.txt", b"hello", "text/plain").file_id
    with pytest.raises(ToolError, match="image or a PDF"):
        agent_tools.inspect_image(ctx, text_id, "what is this?")


# ---------------------------------------------------------------- sandbox result files
def _tar_of(name: str, data: bytes) -> list[bytes]:
    buf = io.BytesIO()
    with tarfile.open(fileobj=buf, mode="w") as tar:
        info = tarfile.TarInfo(name)
        info.size = len(data)
        tar.addfile(info, io.BytesIO(data))
    return [buf.getvalue()]


def test_read_file_returns_bytes_or_none():
    container = SimpleNamespace(get_archive=lambda path: (_tar_of("output.csv", b"a,b\n"), {"size": 4}))
    assert sandbox._read_file(container, "output.csv") == b"a,b\n"
    big = SimpleNamespace(get_archive=lambda path: (_tar_of("x", b"1"), {"size": sandbox.MAX_COLLECT_BYTES + 1}))
    assert sandbox._read_file(big, "x") is None

    def missing(path):
        raise sandbox.NotFound("no such file")

    assert sandbox._read_file(SimpleNamespace(get_archive=missing), "x") is None


@pytest.mark.skipif(not sandbox_available(), reason="Docker or sandbox image not available")
def test_live_sandbox_collects_written_files():
    code = "open('output.csv', 'w').write('a,b\\n1,2\\n')\nprint('ok')\n"
    result = run_in_sandbox({"s.py": code}, ["python", "s.py"], collect=["output.csv", "missing.csv"])
    assert result.ok and result.files == {"output.csv": b"a,b\n1,2\n"}


# ---------------------------------------------------------------- final answer keeps the grounding
GROUNDED = "The IDLH level is 100 ppm [1].\n\nSources:\n[1] osha_h2s_fact_sheet.pdf, p.1"


def test_uncited_retelling_is_replaced_by_the_grounded_answer():
    ctx, _ = make_ctx()
    ctx.scratchpad.answers["ans_1"] = GROUNDED
    assert agent_tools.final_answer(ctx, "OSHA says 100 ppm is dangerous.") == GROUNDED
    assert agent_tools.final_answer(ctx, "IDLH is 100 ppm [1].") == "IDLH is 100 ppm [1]."   # cited: kept


def test_with_files_the_model_text_stays_and_gets_the_sources():
    ctx, _ = make_ctx()
    ctx.scratchpad.answers["ans_1"] = GROUNDED
    ctx.artifacts.append(office.save_text_artifact("x.txt", "x", ArtifactKind.TXT))
    text = agent_tools.final_answer(ctx, "Saved the deck as h2s.pptx.")
    assert text == "Saved the deck as h2s.pptx.\n\nSources:\n[1] osha_h2s_fact_sheet.pdf, p.1"


def test_without_answers_the_text_is_unchanged():
    ctx, _ = make_ctx()
    assert agent_tools.final_answer(ctx, "Done.") == "Done."


def test_image_answers_keep_the_model_wording():
    ctx, _ = make_ctx()
    ctx.scratchpad.answers["ans_1"] = "- Flange weeping\n\n(Seen in note.jpg by the local vision model.)"
    assert agent_tools.final_answer(ctx, "Problems: flange weeping.") == "Problems: flange weeping."
