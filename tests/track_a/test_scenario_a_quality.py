"""
Scenario A quality fix: per-finding KB queries, severities copied from the report,
cost amounts with OCR spaces. Fast tests use OCR-like text (psm 4 layout); the slow
test runs guided Scenario A on both demo reports, 3 times each, through the real API.
"""
from __future__ import annotations

import io
import time
from pathlib import Path
from types import SimpleNamespace

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient

from backend import agent_tools, llm_client
from backend.flows import guided
from backend.settings import settings
from backend.tools import findings, knowledge, office
from shared.contracts import API_PREFIX, ApprovalNote, EventType, Finding, KBHit, Severity, TaskState, TaskStatus

_REPO_ROOT = Path(__file__).resolve().parent.parent.parent

# Report 1 as OCR (psm 4) reads it: garbled row numbers, rows wrapped over several lines.
REPORT_1_TEXT = """--- Page 1 ---
. SAHYADRI DEMO REFINERY (fictional)
Inspection & Corrosion Control Section Report No. SDR/INSP/2026/0412
STORAGE TANK T-104 ANNUAL INSPECTION REPORT
Date of inspection 14-09-2026 Inspector M. P. Joshi, Level Il UT
1. Background
Tank T-104 was taken out of service for its annual inspection. General corrosion within limits.
3. Findings
No. | Item Observation nt Severity
+14" | Shell course 2, north Wall thinning to 6.1 mm against 8.0 mm nominal High
side (24% loss). Below 80% of nominal thickness. Plate
replacement by welding needed.
2 Bottom plate near Pitting up to 1.2 mm deep. About 150 mm sludge High ©
sump layer. H2S of 18 ppm measured at manway M1
+13 Inlet nozzle N2 and No spade blind fitted on N2 during entry. Mixer Medium
mixer MX-104 motor breaker was switched off but not locked or .
4 Confined space entry No gas test entries on the permit after 11:00. Medium
permit _ Attendant left the manway for about 10 minutes. ° :
a Roof handrail and + Surface rust and coating breakdown. No loss of Low
stairway section.
Fictional demo document - Sovereign Al Workbench - no real plant data Page 1 of 2

--- Page 2 ---
SAHYADRI DEMO REFINERY (fictional) Report No. SDR/INSP/2026/0412
4. Recommendation
(a) Replace the two thinned plates of shell course 2, north side, before the tank is returned to service.
Welding must be done under a hot work permit after the tank is cleaned, gas freed and gas tested,
with a fire watch in place. (b) Remove all sludge before any further entry.
Estimated cost: Rs 4,50,000 (plate replacement Rs 3,20,000; sludge removal Rs 90, 000: blinds and
coating Rs 40 ,000).
"""
HEADER_WORDS = ("report no", "refinery", "date of", "inspector", "background", "page 1 of", "fictional")


# ---------------------------------------------------------------- findings table + queries
def test_parse_finding_rows_pairs_each_row_with_its_severity():
    rows = findings.parse_finding_rows(REPORT_1_TEXT)
    assert [r.severity for r in rows] == [Severity.HIGH, Severity.HIGH, Severity.MEDIUM, Severity.MEDIUM,
                                          Severity.LOW]
    assert rows[0].text.startswith("Shell course 2, north") and "welding needed" in rows[0].text
    assert rows[1].text.startswith("Bottom plate near") and "H2S of 18 ppm" in rows[1].text
    assert rows[4].text.startswith("Roof handrail")                  # garbled row number "a" dropped


def test_kb_queries_one_per_finding_then_recommendations_skipping_header_lines():
    queries = findings.kb_queries(REPORT_1_TEXT)
    assert len(findings.finding_queries(REPORT_1_TEXT)) == 5
    assert queries[1].startswith("Bottom plate near Pitting")
    assert any("hot work permit" in q for q in queries[5:])           # recommendation sentences follow
    assert not any("Estimated cost" in q or "Rs 4,50,000" in q for q in queries)
    for query in queries:
        assert not any(word in query.lower() for word in HEADER_WORDS), query


def test_fallback_sentence_queries_skip_header_lines():
    text = ("SAHYADRI DEMO REFINERY Report No. X/1 corrosion section\nDate of inspection: pitting check\n"
            "Shell plate shows pitting corrosion near the sump. The paint is fine.\nPage 1 of 1")
    assert findings.finding_queries(text) == ["Shell plate shows pitting corrosion near the sump."]


def test_match_row_is_fuzzy_on_item_words():
    rows = findings.parse_finding_rows(REPORT_1_TEXT)
    assert findings.match_row("Bottom plate near sump", rows) is rows[1]
    assert findings.match_row("Confined space entry permit", rows) is rows[3]
    assert findings.match_row("Heat exchanger bundle", rows) is None


# ---------------------------------------------------------------- severity check
def _ctx(events: list):
    return agent_tools.ToolContext(task_id="t", emit=lambda *a, **k: events.append(a), add_artifact=lambda a: None)


def _warnings(events: list) -> list[str]:
    return [e[2]["text"] for e in events if e[0] == EventType.LOG and e[2]["level"] == "warn"]


def test_severity_override_uses_report_value_and_warns():
    note = ApprovalNote(ref_no="", date="", subject="s", background="b", recommendation="r", findings=[
        Finding(item="Shell course 2, north side", observation="6.1 mm", severity=Severity.HIGH),
        Finding(item="Bottom plate near sump", observation="pitting", severity=Severity.MEDIUM),
        Finding(item="Inlet nozzle N2 and mixer MX-104", observation="no blind", severity=Severity.LOW),
        Finding(item="Confined space entry permit", observation="no gas test", severity=Severity.MEDIUM),
        Finding(item="Roof handrail and stairway", observation="rust", severity=Severity.LOW)])
    events: list = []
    agent_tools.check_against_report(_ctx(events), note, REPORT_1_TEXT)
    assert [f.severity for f in note.findings] == [Severity.HIGH, Severity.HIGH, Severity.MEDIUM, Severity.MEDIUM,
                                                   Severity.LOW]
    warnings = _warnings(events)
    assert len(warnings) == 2
    assert "'Bottom plate near sump': model said medium, the report says high" in warnings[0]


def test_severity_check_warns_about_a_missing_report_finding():
    note = ApprovalNote(ref_no="", date="", subject="s", background="b", recommendation="r", findings=[
        Finding(item="Shell course 2, north side", observation="6.1 mm", severity=Severity.HIGH)])
    events: list = []
    agent_tools.check_against_report(_ctx(events), note, REPORT_1_TEXT)
    missing = [w for w in _warnings(events) if w.startswith("Report finding not in the note")]
    assert len(missing) == 4 and "(high)" in missing[0]


def test_severity_check_does_nothing_without_a_findings_table():
    note = ApprovalNote(ref_no="", date="", subject="s", background="b", recommendation="r", findings=[
        Finding(item="Anything", observation="x", severity=Severity.LOW)])
    events: list = []
    agent_tools.check_against_report(_ctx(events), note, "Plain report text with no table.")
    assert note.findings[0].severity == Severity.LOW and events == []


def test_item_text_comes_from_the_report_row_and_observation_is_kept():
    observation = "No spade blind fitted on N2 during entry. Mixer motor breaker was switched off but not locked or tagged."
    note = ApprovalNote(ref_no="", date="", subject="s", background="b", recommendation="r", findings=[
        Finding(item="Inlet nozzle N2 and No spade blind fitted on N2 during entry. Mixer mixer MX-104 motor",
                observation=observation, severity=Severity.MEDIUM),
        Finding(item="Confined space entry permit", severity=Severity.MEDIUM,
                observation="No gas test entries on the permit after 11:00. Attendant left the manway for about "
                            "10 minutes."),
        Finding(item="Shell course 2, north side", severity=Severity.HIGH,
                observation="Wall thinning to 6.1 mm against 8.0 mm nominal (24% loss). Below 80% of nominal "
                            "thickness. Plate replacement by welding needed.")])
    agent_tools.check_against_report(_ctx([]), note, REPORT_1_TEXT)
    assert [f.item for f in note.findings] == ["Inlet nozzle N2 and mixer MX-104", "Confined space entry permit",
                                               "Shell course 2, north side"]
    assert note.findings[0].observation == observation


def test_row_item_cleans_ocr_noise():
    rows = findings.parse_finding_rows(REPORT_1_TEXT)
    pitting = "Pitting up to 1.2 mm deep. About 150 mm sludge layer. H2S of 18 ppm measured at manway M1."
    assert findings.row_item(rows[1], pitting) == "Bottom plate near sump"
    assert findings.row_item(rows[4], "Surface rust and coating breakdown. No loss of section.") == \
        "Roof handrail and stairway"                                     # "+" and row number "a" dropped
    text = ("3. Findings\n2 Flange FL-3 near Weep at gasket. Personal H2S monitor alarmed at | High\n"
            "CML-04 - 12 ppm within 1 m of the flange..Area barricaded.\n4. Recommendation")
    row = findings.parse_finding_rows(text)[0]
    assert findings.row_item(row, "Weep at gasket. Personal H2S monitor alarmed at 12 ppm.") == "Flange FL-3 near CML-04"


def test_row_item_does_not_depend_on_the_model_observation():
    rows = findings.parse_finding_rows(REPORT_1_TEXT)
    for observation in ("", "Pitting.", "Something the model wrote differently."):
        assert findings.row_item(rows[0], observation) == "Shell course 2, north side"
        assert findings.row_item(rows[2], observation) == "Inlet nozzle N2 and mixer MX-104"


def test_model_item_is_kept_when_the_row_has_no_clear_item_column():
    text = "3. Findings\nshell plate thinning found near the sump low\n4. Recommendation"
    assert findings.row_item(findings.parse_finding_rows(text)[0], "thinning") is None
    note = ApprovalNote(ref_no="", date="", subject="s", background="b", recommendation="r", findings=[
        Finding(item="Shell plate near sump", observation="thinning found", severity=Severity.LOW)])
    agent_tools.check_against_report(_ctx([]), note, text)
    assert note.findings[0].item == "Shell plate near sump"


def test_garbled_two_letter_row_number_is_dropped():
    text = "3. Findings\noS Insulation cladding About 2 m of damaged cladding near PS-12. Low\n4. Recommendation"
    assert findings.parse_finding_rows(text)[0].text == "Insulation cladding About 2 m of damaged cladding near PS-12."


def test_draft_uses_temperature_0_fixed_seed_and_long_timeout(monkeypatch):
    seen = {}

    def fake_chat_json_meta(model, messages, schema, *, purpose, options=None, timeout_s=None):
        seen["options"], seen["system"], seen["timeout_s"] = options, messages[0]["content"], timeout_s
        raise llm_client.LLMError("MODEL_TIMEOUT", "stop here")

    monkeypatch.setattr(llm_client, "chat_json_meta", fake_chat_json_meta)
    ctx = _ctx([])
    ctx.scratchpad.docs["doc_1"] = agent_tools.DocEntry("doc_1", "f", "r.pdf", "auto", [])
    with pytest.raises(llm_client.LLMError):
        agent_tools.draft_approval_note(ctx, "doc_1")
    assert seen["options"] == {"temperature": 0.0, "seed": agent_tools.NOTE_SEED}
    assert "copy it EXACTLY" in seen["system"]
    assert seen["timeout_s"] == max(agent_tools.NOTE_TIMEOUT_S, settings.WB_LLM_TIMEOUT_S) > 180


def test_chat_json_meta_timeout_applies_to_that_call_only(monkeypatch):
    made = []

    class FakeClient:
        def chat(self, **kwargs):
            return SimpleNamespace(message=SimpleNamespace(content='{"steps": 1}'), eval_count=1)

    def fake_client(timeout_s=None):
        made.append(timeout_s)
        return FakeClient()

    from pydantic import BaseModel

    class Out(BaseModel):
        steps: int

    monkeypatch.setattr(llm_client, "_client", fake_client)
    llm_client.chat_json_meta("m", [{"role": "user", "content": "x"}], Out, timeout_s=420)
    llm_client.chat_json_meta("m", [{"role": "user", "content": "x"}], Out)
    assert made == [420, None]


def test_chat_json_meta_options_override_only_that_call(monkeypatch):
    calls = []

    class FakeClient:
        def chat(self, **kwargs):
            calls.append(kwargs["options"])
            return SimpleNamespace(message=SimpleNamespace(content='{"steps": 1}'), eval_count=1)

    from pydantic import BaseModel

    class Out(BaseModel):
        steps: int

    monkeypatch.setattr(llm_client, "_client", lambda: FakeClient())
    llm_client.chat_json_meta("m", [{"role": "user", "content": "x"}], Out, options={"temperature": 0, "seed": 7})
    llm_client.chat_json_meta("m", [{"role": "user", "content": "x"}], Out)
    assert calls[0]["temperature"] == 0 and calls[0]["seed"] == 7 and calls[0]["num_ctx"] == settings.WB_NUM_CTX
    assert calls[1]["temperature"] == settings.WB_TEMPERATURE and "seed" not in calls[1]


# ---------------------------------------------------------------- cost check
def test_cost_accepts_ocr_spaces_inside_numbers():
    assert office.cost_text("Rs 4,50,000", "Estimated cost: Rs 4,50, 000 (plate)") == "Rs 4,50,000"
    assert office.cost_text("Rs 90,000 and Rs 40,000", "sludge Rs 90, 000: coating Rs 40 ,000") == \
        "Rs 90,000 and Rs 40,000"
    assert office.cost_text("Rs 4,50, 000", "Estimated cost: Rs 4,50,000") == "Rs 4,50,000"


def test_cost_still_rejects_invented_amounts():
    assert office.cost_text("Rs 12,00,000", "Estimated cost: Rs 4,50, 000") == office.COST_PLACEHOLDER
    assert office.cost_text("Rs 4,50,000 plus Rs 5,000 extra", REPORT_1_TEXT) == office.COST_PLACEHOLDER


def test_join_number_spaces_needs_a_digit_on_both_sides():
    assert office.join_number_spaces("Rs 90, 000; 6.1 mm, 8.0 mm; items 1 , 2") == "Rs 90,000; 6.1 mm, 8.0 mm; items 1,2"
    assert office.join_number_spaces("sump, 150 mm") == "sump, 150 mm"


def test_rejected_cost_is_logged_with_the_model_text():
    note = ApprovalNote(ref_no="", date="", subject="s", background="b", recommendation="r",
                        cost_implication="Rs 9,99,999", findings=[Finding(item="x", observation="y",
                                                                          severity=Severity.LOW)])
    events: list = []
    agent_tools._check_cost(_ctx(events), note, REPORT_1_TEXT)
    assert "'Rs 9,99,999' has an amount that is not in the report" in _warnings(events)[0]


# ---------------------------------------------------------------- multi-query KB search
def _hit(source: str, page: int, score: float) -> KBHit:
    return KBHit(text=f"{source} page {page}", source=source, page=page, score=score)


def test_search_knowledge_merges_unique_hits_above_min_score(monkeypatch):
    results = {"q1": [_hit("a.pdf", 1, 0.60), _hit("b.pdf", 2, 0.50)],
               "q2": [_hit("a.pdf", 1, 0.65), _hit("c.pdf", 3, 0.58)],
               "q3": [_hit("d.pdf", 4, 0.40)]}
    monkeypatch.setattr(knowledge, "search", lambda query, top_k=4: results[query])
    events: list = []
    ctx = _ctx(events)
    outcome = agent_tools.search_knowledge(ctx, "q1", queries=["q2", "q3", "q1"])
    hits = list(ctx.scratchpad.kb_hits.values())
    assert [(h.source, h.score) for h in hits] == [("a.pdf", 0.65), ("c.pdf", 0.58)]   # unique, best score kept
    assert outcome.summary.startswith("2 relevant passage(s) from 3 queries")
    log = next(e[2]["text"] for e in events if e[1] == "Knowledge base queries")
    assert log.splitlines()[0].startswith("q1 best 0.600 a.pdf p.1 (1 >= 0.57)")
    assert log.splitlines()[2].startswith("q3 best 0.400 d.pdf p.4 (0 >= 0.57)")


def test_single_query_search_is_unchanged(monkeypatch):
    monkeypatch.setattr(knowledge, "search", lambda query, top_k=4: [_hit("a.pdf", 1, 0.60)])
    events: list = []
    outcome = agent_tools.search_knowledge(_ctx(events), "only one", queries=["only one", ""])
    assert outcome.summary.startswith("kb_1: a.pdf p.1 (score 0.60)") and events == []


def test_pick_kb_refs_prefers_different_sop_files():
    ctx = _ctx([])
    for kb_id, hit in {"kb_1": _hit("hot.pdf", 73, 0.73), "kb_2": _hit("hot.pdf", 27, 0.72),
                       "kb_3": _hit("h2s.pdf", 2, 0.61), "kb_4": _hit("hot.pdf", 101, 0.58)}.items():
        ctx.scratchpad.kb_hits[kb_id] = hit
    assert guided.pick_kb_refs(ctx, ["kb_1", "kb_2", "kb_3", "kb_4"]) == ["kb_1", "kb_3", "kb_2"]
    assert guided.pick_kb_refs(ctx, ["kb_1", "kb_2"], limit=1) == ["kb_1"]


# ---------------------------------------------------------------- live: 3 runs per report
EXPECTED = {
    1: {"file": "scenario_a_report_1.pdf", "cost": "Rs 4,50,000",
        "findings": [("shell course 2", "High"), ("bottom plate", "High"), ("nozzle n2", "Medium"),
                     ("confined space", "Medium"), ("roof handrail", "Low")],
        "sops": {"nsw_hot_work_petroleum.pdf", "osha_h2s_fact_sheet.pdf", "osha_h2s_quickcard.pdf",
                 "osha_lockout_tagout.pdf", "osha_confined_space.pdf"},
        "never": set()},
    2: {"file": "scenario_a_report_2.pdf", "cost": "Rs 2,85,000",
        "findings": [("cml-03", "High"), ("fl-3", "High"), ("dead leg", "Medium"), ("cml-05", "Medium"),
                     ("cladding", "Low")],
        "sops": {"nsw_hot_work_petroleum.pdf", "osha_lockout_tagout.pdf", "osha_h2s_fact_sheet.pdf",
                 "osha_h2s_quickcard.pdf"},
        "never": {"osha_confined_space.pdf"}},
}
RUNS_PER_REPORT = 3
RUN_LIMIT_S = 300
LIVE_RESULTS: list[str] = []


def _live_ready() -> bool:
    try:
        ok = httpx.get(f"{settings.OLLAMA_HOST}/api/tags", timeout=3).status_code == 200
    except httpx.HTTPError:
        return False
    return ok and all((_REPO_ROOT / "demo" / "inputs" / e["file"]).exists() for e in EXPECTED.values())


def _note_parts(docx_bytes: bytes) -> tuple[list[tuple[str, str]], str]:
    doc = Document(io.BytesIO(docx_bytes))
    rows = [[c.text.strip() for c in row.cells] for table in doc.tables for row in table.rows]
    finding_rows = [(r[1], r[3]) for r in rows if len(r) == 5 and r[3] in ("Low", "Medium", "High", "Critical")]
    text = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(" | ".join(r) for r in rows)
    return finding_rows, text


def _run_report(client: TestClient, report: int) -> tuple[TaskState, bytes]:
    path = _REPO_ROOT / "demo" / "inputs" / EXPECTED[report]["file"]
    up = client.post(f"{API_PREFIX}/files", files={"file": (path.name, path.read_bytes(), "application/pdf")})
    body = {"message": "Draft an approval note for this inspection report, citing the relevant SOPs.",
            "file_ids": [up.json()["file_id"]], "mode": "guided", "scenario": "inspection_note"}
    task_id = client.post(f"{API_PREFIX}/tasks", json=body).json()["task_id"]
    started = time.monotonic()
    while time.monotonic() - started < RUN_LIMIT_S + 120:
        state = TaskState.model_validate(client.get(f"{API_PREFIX}/tasks/{task_id}").json())
        if state.status in (TaskStatus.SUCCEEDED, TaskStatus.FAILED, TaskStatus.CANCELLED):
            break
        time.sleep(2)
    docx_bytes = client.get(state.artifacts[0].download_url).content if state.artifacts else b""
    return state, docx_bytes


@pytest.fixture(scope="module")
def live_client(tmp_path_factory):
    root = tmp_path_factory.mktemp("scenario_a")
    with pytest.MonkeyPatch.context() as mp:
        for key in ("WB_WORKSPACE_DIR", "WB_CACHE_DIR", "WB_LOG_DIR"):
            mp.setattr(settings, key, root / key.lower())     # real KB (data/kb, data/chroma), read only
        knowledge.reset_client()
        from backend.main import app

        with TestClient(app) as client:
            yield client
        knowledge.reset_client()
    print("\n\n" + "\n".join(LIVE_RESULTS))


@pytest.mark.slow
@pytest.mark.skipif(not _live_ready(), reason="Ollama or demo reports not available")
@pytest.mark.parametrize("report,run", [(r, n) for r in EXPECTED for n in range(1, RUNS_PER_REPORT + 1)])
def test_live_scenario_a_is_correct_and_stable(live_client, network_guard, report, run):
    expected = EXPECTED[report]
    state, docx_bytes = _run_report(live_client, report)
    assert state.status == TaskStatus.SUCCEEDED, state.error
    rows, text = _note_parts(docx_bytes)
    cited = {f for f in expected["sops"] | expected["never"] if f in text}
    lines = [f"Report {report}, run {run}: {state.elapsed_s:.0f} s, cost {'OK' if expected['cost'] in text else 'MISSING'}"
             f", SOPs cited: {', '.join(sorted(cited)) or 'none'}",
             f"  {'expected item':<16} {'report':<7} {'note':<7} note item"]
    problems = []
    for key, severity in expected["findings"]:
        match = next(((item, sev) for item, sev in rows if key in item.lower()), None)
        lines.append(f"  {key:<16} {severity:<7} {match[1] if match else '-':<7} {match[0] if match else 'MISSING'}")
        if match is None or match[1] != severity:
            problems.append(f"{key}: {match}")
    LIVE_RESULTS.append("\n".join(lines))
    print("\n" + "\n".join(lines))
    assert not problems, problems
    assert len(rows) == len(expected["findings"]), rows
    assert cited & expected["sops"], "no expected SOP cited"
    assert not cited & expected["never"], f"must not cite {cited & expected['never']}"
    assert expected["cost"] in text
    assert state.elapsed_s < RUN_LIMIT_S
