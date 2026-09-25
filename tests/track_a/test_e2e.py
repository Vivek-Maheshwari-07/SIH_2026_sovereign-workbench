"""
Live end-to-end runs through the REAL API, exactly like the UI (ticket A8b):
upload -> POST /api/tasks -> poll events -> GET task -> download artifact.

Scenarios A (inspection note), B (calculation code), C (P&ID tags), each in
guided and agent mode. Needs Ollama (all registry models) and Docker; marked
slow. Uses a TEMPORARY workspace, KB and Chroma dir (never data/kb or
data/chroma), and the network guard: 0 external connection attempts allowed.
"""
from __future__ import annotations

import io
import time
from pathlib import Path

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import e2e_inputs
from backend.flows.code_flow import PIPE_THICKNESS_DEMO_REQUEST
from backend.registry import registry
from backend.settings import settings
from backend.tools import knowledge
from backend.tools.sandbox import sandbox_available
from shared.contracts import API_PREFIX, EventType, TaskState, TaskStatus

_REPO_ROOT = Path(__file__).resolve().parent.parent.parent
RUN_TIMEOUT_S = 900
EVENT_KEYS = {
    "route": {"decision"}, "plan": {"steps"}, "step_start": {"index", "title"},
    "llm_call": {"model_id", "purpose", "duration_ms", "tokens_out"}, "tool_call": {"tool", "args"},
    "tool_result": {"tool", "ok", "summary", "duration_ms"}, "artifact": {"artifact"},
    "log": {"level", "text"}, "final": {"answer"}, "error": {"error"},
}
MESSAGES = {
    "A": "Draft an approval note for this inspection report, citing the relevant SOPs.",
    "B": PIPE_THICKNESS_DEMO_REQUEST,
    "C": "Extract all equipment and instrument tags from this P&ID drawing into an Excel tag list.",
}
SCENARIO_NAMES = {"A": "inspection_note", "B": "code_calc", "C": "pid_tags"}
RESULTS: list[dict] = []
TIMELINES: dict[str, list[str]] = {}


def _live_ready() -> bool:
    try:
        tags = httpx.get(f"{settings.OLLAMA_HOST}/api/tags", timeout=3).json()
    except (httpx.HTTPError, ValueError):
        return False
    pulled = {m["name"] for m in tags.get("models", [])} | {m["name"].split(":")[0] for m in tags.get("models", [])}
    needed = {m.ollama_name for m in registry.all_models()}
    return all(n in pulled or f"{n}:latest" in pulled for n in needed) and sandbox_available()


pytestmark = [
    pytest.mark.slow,
    pytest.mark.skipif(not _live_ready(), reason="Ollama models or Docker sandbox not available"),
]


@pytest.fixture(scope="module")
def live_env(tmp_path_factory):
    root = tmp_path_factory.mktemp("e2e")
    with pytest.MonkeyPatch.context() as mp:
        for key, sub in (("WB_WORKSPACE_DIR", "workspace"), ("WB_CACHE_DIR", "cache"), ("WB_LOG_DIR", "logs"),
                         ("WB_CHROMA_DIR", "chroma"), ("WB_KB_DIR", "kb")):
            mp.setattr(settings, key, root / sub)
        knowledge.reset_client()
        kb_files = e2e_inputs.build_sops(root / "kb")
        report = knowledge.ingest_folder(progress=lambda _m: None)
        assert report.chunks > 0 and not report.failed_files
        inputs = {
            "A": e2e_inputs.demo_input("inspection", _REPO_ROOT) or e2e_inputs.build_report(root / "inputs"),
            "C": e2e_inputs.demo_input("pid", _REPO_ROOT) or e2e_inputs.build_pid_image(root / "inputs"),
        }
        from backend.main import app

        with TestClient(app) as client:
            yield {"client": client, "kb_files": kb_files, "inputs": inputs, "root": root}
        knowledge.reset_client()
    _print_report()


def _print_report() -> None:
    if not RESULTS:
        return
    print("\n\nScenario x mode results")
    print(f"{'scenario':<18} {'mode':<7} {'status':<10} {'steps':>5} {'time s':>7}  fallback")
    for r in RESULTS:
        print(f"{r['scenario']:<18} {r['mode']:<7} {r['status']:<10} {r['steps']:>5} {r['seconds']:>7.0f}  "
              f"{'yes' if r['fallback'] else 'no'}")
    for key, lines in TIMELINES.items():
        print(f"\nEvent timeline: {key}")
        for line in lines:
            print("  " + line)


def _short(event: dict) -> str:
    data, kind = event["data"], event["type"]
    if kind == "llm_call":
        text = f"{data['model_id']} {data['purpose']} {data['duration_ms'] / 1000:.1f}s tokens={data['tokens_out']}"
    elif kind == "tool_call":
        text = f"{data['tool']}({', '.join(f'{k}={str(v)[:40]!r}' for k, v in data['args'].items())})"
    elif kind == "tool_result":
        text = f"{data['tool']} ok={data['ok']} {data['duration_ms'] / 1000:.1f}s: {' '.join(data['summary'].split())[:90]}"
    elif kind == "log":
        text = f"[{data['level']}] {data['text'][:100]}"
    elif kind == "final":
        text = " ".join(data["answer"].split())[:110]
    elif kind == "error":
        text = f"{data['error']['code']}: {data['error']['message'][:90]}"
    else:
        text = event["title"][:100]
    return f"{event['seq']:>3} {kind:<11} {text}"


def _upload(client: TestClient, path: Path) -> str:
    mime = "application/pdf" if path.suffix == ".pdf" else "image/png"
    resp = client.post(f"{API_PREFIX}/files", files={"file": (path.name, path.read_bytes(), mime)})
    assert resp.status_code == 200, resp.text
    return resp.json()["file_id"]


def _run(client: TestClient, message: str, file_ids: list[str], mode: str, scenario: str | None):
    body = {"message": message, "file_ids": file_ids, "mode": mode}
    if scenario:
        body["scenario"] = scenario
    resp = client.post(f"{API_PREFIX}/tasks", json=body)
    assert resp.status_code == 202, resp.text
    task_id = resp.json()["task_id"]
    events, after, started = [], 0, time.monotonic()
    while time.monotonic() - started < RUN_TIMEOUT_S:
        page = client.get(f"{API_PREFIX}/tasks/{task_id}/events", params={"after": after}).json()
        events.extend(page["events"])
        after = page["next_seq"]
        if page["done"]:
            break
        time.sleep(1.0)
    else:
        raise AssertionError(f"task {task_id} did not finish in {RUN_TIMEOUT_S}s")
    state = TaskState.model_validate(client.get(f"{API_PREFIX}/tasks/{task_id}").json())
    return state, events, time.monotonic() - started


def _download(client: TestClient, artifact) -> bytes:
    resp = client.get(artifact.download_url)
    assert resp.status_code == 200
    assert artifact.filename in resp.headers["content-disposition"]
    assert len(resp.content) == artifact.size_bytes
    return resp.content


@pytest.mark.parametrize("mode", ["guided", "agent"])
@pytest.mark.parametrize("scenario", ["A", "B", "C"])
def test_live_scenario(live_env, network_guard, scenario, mode):
    client = live_env["client"]
    file_ids = [_upload(client, live_env["inputs"][scenario])] if scenario in live_env["inputs"] else []
    state, events, seconds = _run(client, MESSAGES[scenario], file_ids, mode,
                                  SCENARIO_NAMES[scenario] if mode == "guided" else None)
    fallback = any(e["type"] == "log" and "switched to guided flow" in e["data"]["text"] for e in events)
    RESULTS.append({"scenario": f"{scenario} {SCENARIO_NAMES[scenario]}", "mode": mode, "status": state.status.value,
                    "steps": sum(e["type"] == "step_start" for e in events), "seconds": seconds,
                    "fallback": fallback})
    if scenario == "A" and mode == "agent":
        TIMELINES["agent-mode Scenario A"] = [_short(e) for e in events]
    print(f"\n{scenario}/{mode}: {state.status.value} in {seconds:.0f}s, fallback={fallback}")
    for event in events:
        print("   " + _short(event))

    assert state.status == TaskStatus.SUCCEEDED, state.error
    for event in events:
        assert set(event["data"]) == EVENT_KEYS[event["type"]], event
    assert [e["seq"] for e in events] == list(range(1, len(events) + 1))
    assert events[-1]["type"] == EventType.FINAL.value

    if scenario == "A":
        docx = [a for a in state.artifacts if a.kind == "docx"]
        assert docx, "no Word artifact"
        doc = Document(io.BytesIO(_download(client, docx[-1])))
        findings = next(t for t in doc.tables if t.rows[0].cells[1].text == "Item")
        assert len(findings.rows) - 1 >= 2
        sop = next(t for t in doc.tables if [c.text for c in t.rows[0].cells] == ["S.No", "Document", "Page"])
        cited = [r.cells[1].text for r in sop.rows[1:] if r.cells[1].text != "Not applicable"]
        assert cited and all(name in live_env["kb_files"] for name in cited), cited
    elif scenario == "B":
        py = [a for a in state.artifacts if a.kind == "py"]
        assert len(py) >= 2
        assert all(b"def " in _download(client, a) for a in py)
        sandbox = [e for e in events if e["type"] == "tool_result" and e["data"]["tool"] == "sandbox"
                   and "passed" in e["data"]["summary"]]
        assert sandbox and sandbox[-1]["data"]["summary"].endswith(", 0 failed")
        texts = [state.final_answer or ""] + [e["data"]["summary"] for e in events if e["type"] == "tool_result"]
        assert any("7.24" in t or "7.25" in t for t in texts)
    else:
        xlsx = [a for a in state.artifacts if a.kind == "xlsx"]
        assert xlsx, "no Excel artifact"
        wb = load_workbook(io.BytesIO(_download(client, xlsx[-1])))
        found = {str(r[0].value).strip().upper() for r in wb["Tags"].iter_rows(min_row=2) if r[0].value}
        hits = [t for t in e2e_inputs.EXPECTED_TAGS if t in found]
        print(f"   tags found {len(hits)}/10: {sorted(found)}")
        assert len(hits) >= 6, sorted(found)

    assert network_guard == [], network_guard
