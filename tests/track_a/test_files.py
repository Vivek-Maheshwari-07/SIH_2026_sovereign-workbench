"""
Tests for backend.tools.files: safe_path, upload validation (size + magic
bytes), filename sanitization, and the read_file/write_file agent tools.
"""
from __future__ import annotations

from pathlib import Path

import pytest

from backend.settings import settings
from backend.tools.files import (
    ALLOWED_SUFFIXES,
    FileSafetyError,
    read_file,
    safe_path,
    sanitize_filename,
    validate_upload,
    write_file,
)

_REPO_ROOT = Path(__file__).resolve().parent.parent.parent


def _workspace_root() -> Path:
    root = settings.WB_WORKSPACE_DIR
    return (root if root.is_absolute() else _REPO_ROOT / root).resolve()


# ---------------------------------------------------------------- safe_path
@pytest.mark.parametrize(
    "bad_path",
    [
        "..\\..\\Windows",
        "../../Windows",
        r"C:\Windows\win.ini",
        r"\\server\share\x",
        "/etc/passwd",
    ],
)
def test_safe_path_rejects_escapes(bad_path):
    with pytest.raises(FileSafetyError) as exc_info:
        safe_path(bad_path)
    assert exc_info.value.code == "BAD_REQUEST"


def test_safe_path_accepts_a_normal_relative_name():
    resolved = safe_path("reports/note.txt")
    assert resolved.is_relative_to(_workspace_root())
    assert resolved.name == "note.txt"


def test_safe_path_rejects_empty_string():
    with pytest.raises(FileSafetyError):
        safe_path("")


# ---------------------------------------------------------------- sanitize_filename
def test_sanitize_filename_strips_path_parts():
    assert sanitize_filename("../../etc/passwd") == "passwd"
    assert sanitize_filename(r"C:\Windows\win.ini") == "win.ini"


def test_sanitize_filename_replaces_odd_characters():
    assert sanitize_filename("weird name!.txt") == "weird_name_.txt"


def test_sanitize_filename_never_returns_empty():
    assert sanitize_filename("") == "upload"
    assert sanitize_filename("...") == "upload"


# ---------------------------------------------------------------- validate_upload
def test_allowed_suffixes_matches_contract_exactly():
    # Must match shared/contracts.py's ERROR_CODES["UNSUPPORTED_FILE"] comment exactly.
    assert ALLOWED_SUFFIXES == {".pdf", ".png", ".jpg", ".jpeg", ".txt", ".md", ".py", ".csv", ".xlsx", ".docx"}


def test_validate_upload_accepts_a_real_pdf():
    content = b"%PDF-1.4\n1 0 obj<<>>endobj\n%%EOF"
    name, suffix = validate_upload("report.pdf", content)
    assert name == "report.pdf"
    assert suffix == ".pdf"


def test_validate_upload_accepts_a_real_png():
    png_header = b"\x89PNG\r\n\x1a\n" + b"0" * 32
    name, suffix = validate_upload("photo.png", png_header)
    assert suffix == ".png"


def test_validate_upload_rejects_too_large_file():
    max_bytes = settings.WB_MAX_UPLOAD_MB * 1024 * 1024
    content = b"%PDF-1.4" + b"0" * (max_bytes + 1)
    with pytest.raises(FileSafetyError) as exc_info:
        validate_upload("big.pdf", content)
    assert exc_info.value.code == "FILE_TOO_LARGE"


def test_validate_upload_rejects_fake_pdf_with_exe_bytes():
    exe_bytes = b"MZ\x90\x00\x03\x00\x00\x00" + b"\x00" * 32  # real Windows PE header
    with pytest.raises(FileSafetyError) as exc_info:
        validate_upload("virus.pdf", exe_bytes)
    assert exc_info.value.code == "UNSUPPORTED_FILE"


def test_validate_upload_rejects_unlisted_extension():
    with pytest.raises(FileSafetyError) as exc_info:
        validate_upload("program.exe", b"MZ\x90\x00")
    assert exc_info.value.code == "UNSUPPORTED_FILE"


def test_validate_upload_rejects_tif_upload():
    # tif/tiff were removed from the allowed list; see docs/contract_change_requests.md.
    tiff_header = b"II*\x00" + b"0" * 32
    with pytest.raises(FileSafetyError) as exc_info:
        validate_upload("scan.tif", tiff_header)
    assert exc_info.value.code == "UNSUPPORTED_FILE"


def test_validate_upload_rejects_txt_with_binary_content():
    with pytest.raises(FileSafetyError) as exc_info:
        validate_upload("notes.txt", b"\x00\x01\x02binary-not-text\xff")
    assert exc_info.value.code == "UNSUPPORTED_FILE"


def test_validate_upload_accepts_md_py_csv_as_plain_text():
    for filename in ("readme.md", "script.py", "data.csv"):
        name, suffix = validate_upload(filename, b"hello, world\n")
        assert suffix == Path(filename).suffix.lower()


def _make_docx_bytes(tmp_path: Path) -> bytes:
    import docx

    doc_path = tmp_path / "sample.docx"
    document = docx.Document()
    document.add_paragraph("hello from a real docx")
    document.save(doc_path)
    return doc_path.read_bytes()


def _make_xlsx_bytes(tmp_path: Path) -> bytes:
    import openpyxl

    xlsx_path = tmp_path / "sample.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "hello"
    workbook.save(xlsx_path)
    return xlsx_path.read_bytes()


def test_validate_upload_accepts_a_real_docx(tmp_path):
    content = _make_docx_bytes(tmp_path)
    name, suffix = validate_upload("report.docx", content)
    assert suffix == ".docx"


def test_validate_upload_accepts_a_real_xlsx(tmp_path):
    content = _make_xlsx_bytes(tmp_path)
    name, suffix = validate_upload("sheet.xlsx", content)
    assert suffix == ".xlsx"


def test_validate_upload_rejects_random_zip_renamed_to_docx(tmp_path):
    import zipfile

    zip_path = tmp_path / "random.zip"
    with zipfile.ZipFile(zip_path, "w") as zf:
        zf.writestr("not_a_word_doc.txt", "just some random zip contents")

    with pytest.raises(FileSafetyError) as exc_info:
        validate_upload("fake.docx", zip_path.read_bytes())
    assert exc_info.value.code == "UNSUPPORTED_FILE"


# ---------------------------------------------------------------- read_file / write_file
def test_write_file_then_read_file_roundtrip():
    write_file("scratch/hello.txt", "hello world")
    assert read_file("scratch/hello.txt") == b"hello world"


def test_write_file_accepts_bytes_too():
    write_file("scratch/bytes.bin", b"\x01\x02\x03")
    assert read_file("scratch/bytes.bin") == b"\x01\x02\x03"


def test_read_file_missing_raises_file_not_found():
    with pytest.raises(FileSafetyError) as exc_info:
        read_file("scratch/does_not_exist_ever.txt")
    assert exc_info.value.code == "FILE_NOT_FOUND"


def test_write_file_rejects_escaping_path():
    with pytest.raises(FileSafetyError):
        write_file("../outside.txt", "nope")
