"""Tests for backend/tools/office.py and GET /api/artifacts/{id} (ticket A7)."""
from __future__ import annotations

import importlib
import importlib.util
import json
import re
import shutil
import subprocess
import threading
import zipfile
from datetime import datetime
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pptx import Presentation

from backend.settings import settings
from backend.tools import office
from backend.tools.office import (
    COST_PLACEHOLDER,
    DRAFT_FOOTER,
    NOT_APPLICABLE,
    OfficeError,
    clean_text,
    cost_text,
    make_excel,
    make_ppt,
    make_word,
    split_sop_reference,
)
from shared.contracts import API_PREFIX, ApprovalNote, ArtifactKind, Finding, PidTag, PidTagList, Severity

_REPO_ROOT = Path(__file__).resolve().parent.parent.parent
FIXTURE = json.loads((_REPO_ROOT / "shared" / "fixtures" / "inspection_note.json").read_text(encoding="utf-8"))
NASTY = "bad\x00ctrl\x0bchars\x1f 😀 हिंदी पाठ ₹5,00,000"


def fresh_office_module():
    """A second, independent copy of office.py: simulates a backend restart (no shared memory)."""
    spec = importlib.util.spec_from_file_location("office_after_restart", office.__file__)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def sample_note(**overrides) -> ApprovalNote:
    """Built from the inspection_note fixture's story (Tank T-104, 3 findings, 1 high)."""
    subject = FIXTURE["artifact"]["preview"].split(". Findings")[0].removeprefix("Subject: ")
    data = dict(
        ref_no="INSP-XYZ",
        date="01-01-1999",
        subject=subject,
        background=f"Inspection report {FIXTURE['file']['filename']} ({FIXTURE['file']['pages']} pages, scanned) "
                   "for storage tank T-104 was reviewed during the annual shutdown.",
        findings=[
            Finding(item="Shell course 2, north side", observation="Wall thinning to 6.1 mm vs 8.0 mm nominal",
                    severity=Severity.HIGH, source_page=1),
            Finding(item="Bottom plate near sump", observation="Pitting up to 1.2 mm deep, isolated",
                    severity=Severity.MEDIUM, source_page=2),
            Finding(item="Roof handrail", observation="Surface rust, coating breakdown",
                    severity=Severity.LOW, source_page=None),
        ],
        sop_references=["SOP-INSP-012, p.4", "SOP-MNT-003 page 7", "API 653 general guidance"],
        recommendation="Repair shell course 2 before next run.",
        cost_implication=None,
    )
    data.update(overrides)
    return ApprovalNote(**data)


@pytest.fixture(autouse=True)
def workspace(tmp_path, monkeypatch):
    ws = tmp_path / "workspace"
    ws.mkdir()
    monkeypatch.setattr(settings, "WB_WORKSPACE_DIR", ws)
    return ws


def _doc_blocks(doc) -> list[str]:
    """Body text in document order: paragraphs, and every cell of every table."""
    blocks: list[str] = []
    for child in doc.element.body.iterchildren():
        if child.tag.endswith("}p"):
            blocks.append("".join(t.text or "" for t in child.iter() if t.tag.endswith("}t")))
        elif child.tag.endswith("}tbl"):
            for tc in child.iter():
                if tc.tag.endswith("}tc"):
                    blocks.append("".join(t.text or "" for t in tc.iter() if t.tag.endswith("}t")))
    return blocks


def _path_of(artifact) -> Path:
    found = office.artifact_path(artifact.artifact_id)
    assert found is not None
    return found[0]


def _assert_valid_zip(path: Path) -> None:
    with zipfile.ZipFile(path) as zf:
        assert zf.testzip() is None


# ---------------------------------------------------------------- Word
def test_word_sections_in_order_and_footer():
    artifact = make_word(sample_note())
    path = _path_of(artifact)
    doc = Document(str(path))
    blocks = _doc_blocks(doc)
    order = ["Ref no", "Date", "Subject", "Background", "Findings", "SOP references",
             "Recommendation", "Cost implication", "Approval", "Prepared by", "Reviewed by", "Approved by"]
    positions = [blocks.index(label) for label in order]
    assert positions == sorted(positions)
    for line in ("Name", "Designation", "Signature"):
        assert line in blocks
    footer_text = "\n".join(p.text for p in doc.sections[0].footer.paragraphs)
    assert DRAFT_FOOTER in footer_text
    assert "PAGE" in doc.sections[0].footer._element.xml            # page number field
    assert abs(doc.sections[0].page_width.cm - 21.0) < 0.1           # A4
    assert not any("{{" in b for b in blocks)
    _assert_valid_zip(path)


def test_word_findings_table_and_severity_colours():
    note = sample_note()
    doc = Document(str(_path_of(make_word(note))))
    table = next(t for t in doc.tables if t.rows[0].cells[0].text == "S.No" and t.rows[0].cells[1].text == "Item")
    assert len(table.rows) == 1 + len(note.findings)
    assert [c.text for c in table.rows[0].cells] == ["S.No", "Item", "Observation", "Severity", "Source page"]
    assert [c.text for c in table.rows[1].cells] == ["1", "Shell course 2, north side",
                                                     "Wall thinning to 6.1 mm vs 8.0 mm nominal", "High", "1"]
    assert table.rows[3].cells[4].text == "-"
    fills = [re.search(r'w:fill="([0-9A-F]{6})"', r.cells[3]._tc.xml).group(1) for r in table.rows[1:]]
    assert fills == [office.SEVERITY_FILLS[Severity.HIGH], office.SEVERITY_FILLS[Severity.MEDIUM],
                     office.SEVERITY_FILLS[Severity.LOW]]


def test_word_overwrites_ref_no_and_date():
    artifact = make_word(sample_note(ref_no="INSP-XYZ", date="01-01-1999"))
    doc = Document(str(_path_of(artifact)))
    meta = {row.cells[0].text: row.cells[1].text for row in doc.tables[0].rows}
    year = datetime.now().year
    assert meta["Ref no"] == f"INSP-{year}-001"
    assert meta["Date"] == datetime.now().strftime("%d-%m-%Y")
    text = "\n".join(_doc_blocks(doc))
    assert "INSP-XYZ" not in text and "01-01-1999" not in text
    assert artifact.filename == f"INSP-{year}-001_approval_note.docx"


def test_word_sop_references_and_empty_sections():
    doc = Document(str(_path_of(make_word(sample_note()))))
    sop = next(t for t in doc.tables if [c.text for c in t.rows[0].cells] == ["S.No", "Document", "Page"])
    assert [[c.text for c in r.cells] for r in sop.rows[1:]] == [
        ["1", "SOP-INSP-012", "4"], ["2", "SOP-MNT-003", "7"], ["3", "API 653 general guidance", "-"]]

    empty = sample_note(sop_references=[], background="   ", recommendation="")
    doc = Document(str(_path_of(make_word(empty))))
    blocks = _doc_blocks(doc)
    assert blocks[blocks.index("Background") + 1] == NOT_APPLICABLE
    assert blocks[blocks.index("Recommendation") + 1] == NOT_APPLICABLE
    sop = next(t for t in doc.tables if [c.text for c in t.rows[0].cells] == ["S.No", "Document", "Page"])
    assert sop.rows[1].cells[1].text == NOT_APPLICABLE


def test_word_cost_never_invented():
    blocks = _doc_blocks(Document(str(_path_of(make_word(sample_note(cost_implication="Rs 4,50,000 approx"))))))
    assert blocks[blocks.index("Cost implication") + 1] == COST_PLACEHOLDER

    source = "Contractor quote attached: Rs 4,50,000 for plate replacement."
    kept = make_word(sample_note(cost_implication="Rs 4,50,000 as per quote"), source_text=source)
    blocks = _doc_blocks(Document(str(_path_of(kept))))
    assert blocks[blocks.index("Cost implication") + 1] == "Rs 4,50,000 as per quote"


def test_cost_text_rules():
    assert cost_text(None, None) == COST_PLACEHOLDER
    assert cost_text("", "anything") == COST_PLACEHOLDER
    assert cost_text("₹ 2 lakh", None) == COST_PLACEHOLDER
    assert cost_text("₹ 2 lakh", "budget ₹ 3 lakh") == COST_PLACEHOLDER
    assert cost_text("Within approved maintenance budget", None) == "Within approved maintenance budget"
    assert cost_text("USD 1,200", "quote USD 1200") == "USD 1,200"
    assert cost_text("Rs 4,50,000.00", "Rs 4,50,000") == "Rs 4,50,000.00"
    assert cost_text("Rs 4.5 lakh", "Rs 4,50,000") == COST_PLACEHOLDER      # lakh/crore words not converted


def test_split_sop_reference():
    assert split_sop_reference("SOP-INSP-012, p.4") == ("SOP-INSP-012", "4")
    assert split_sop_reference("manual.pdf page 12") == ("manual.pdf", "12")
    assert split_sop_reference("General note") == ("General note", "-")


def test_word_preview_is_short_and_useful():
    artifact = make_word(sample_note())
    assert artifact.kind == ArtifactKind.DOCX and artifact.preview.startswith("Subject: Approval for repair")
    assert "HIGH Shell course 2" in artifact.preview and len(artifact.preview) <= 500
    assert artifact.download_url == f"{API_PREFIX}/artifacts/{artifact.artifact_id}"
    assert artifact.size_bytes == _path_of(artifact).stat().st_size


# ---------------------------------------------------------------- ref counter
def test_ref_counter_sequence_and_restart():
    year = datetime.now().year
    assert [office.next_ref_no() for _ in range(3)] == [f"INSP-{year}-{n:03d}" for n in (1, 2, 3)]
    fresh = fresh_office_module()                       # "restart": new module state, same counter file
    assert fresh.next_ref_no() == f"INSP-{year}-004"
    assert fresh.next_ref_no(datetime(year + 1, 1, 1)) == f"INSP-{year + 1}-001"


def test_ref_counter_threads_never_repeat():
    results: list[str] = []
    lock = threading.Lock()

    def worker():
        for _ in range(10):
            ref = office.next_ref_no()
            with lock:
                results.append(ref)

    threads = [threading.Thread(target=worker) for _ in range(8)]
    for t in threads:
        t.start()
    for t in threads:
        t.join()
    assert len(results) == 80 and len(set(results)) == 80


# ---------------------------------------------------------------- illegal characters
def test_clean_text():
    cleaned = clean_text(NASTY)
    assert "\x00" not in cleaned and "\x0b" not in cleaned and "\x1f" not in cleaned
    assert "😀" in cleaned and "हिंदी" in cleaned and "₹" in cleaned
    assert clean_text("a\tb\nc") == "a\tb\nc"
    assert len(clean_text("x" * 10_000, 100)) == 100


def test_nasty_text_in_all_three_files():
    note = sample_note(subject=NASTY, background=NASTY, recommendation=NASTY,
                       findings=[Finding(item=NASTY, observation=NASTY, severity=Severity.CRITICAL, source_page=3)],
                       sop_references=[NASTY])
    word = _path_of(make_word(note))
    tags = PidTagList(drawing_title=NASTY, tags=[PidTag(tag="P-1" + NASTY, equipment_type=NASTY, description=NASTY, tile=1)])
    excel = _path_of(make_excel(tags))
    findings_xlsx = _path_of(make_excel(note.findings))
    ppt = _path_of(make_ppt(NASTY, [(NASTY, [NASTY] * 3)]))

    for path in (word, excel, findings_xlsx, ppt):
        _assert_valid_zip(path)
    assert "हिंदी पाठ ₹5,00,000" in "\n".join(_doc_blocks(Document(str(word))))
    assert "😀" in load_workbook(excel)["Tags"]["B2"].value
    assert "₹" in Presentation(str(ppt)).slides[1].shapes.title.text


def test_very_long_text_is_trimmed():
    doc = Document(str(_path_of(make_word(sample_note(recommendation="word " * 5000)))))
    blocks = _doc_blocks(doc)
    assert len(blocks[blocks.index("Recommendation") + 1]) <= office.MAX_PARAGRAPH_CHARS


# ---------------------------------------------------------------- Excel
def _tag_list() -> PidTagList:
    return PidTagList(drawing_title="P&ID Unit 100", tags=[
        PidTag(tag="P-101A", equipment_type="Pump", description="Feed pump", tile=1),
        PidTag(tag=" p-101a ", equipment_type="Pump", description="duplicate from overlap", tile=2),
        PidTag(tag="V-201", equipment_type="Valve", tile=1),
        PidTag(tag="V-202", equipment_type="valve", tile=3),
        PidTag(tag="P-102", equipment_type="Pump", tile=4),
        PidTag(tag="P-102", equipment_type="Pump", tile=4),
        PidTag(tag="T-301", equipment_type="Vessel", tile=None),
    ])


def test_excel_tags_dedupe_and_summary():
    artifact = make_excel(_tag_list())
    path = _path_of(artifact)
    assert re.fullmatch(r"pid_tags_[0-9a-f]{8}\.xlsx", artifact.filename)
    wb = load_workbook(path)
    tags = wb["Tags"]
    rows = [[c.value for c in r] for r in tags.iter_rows(min_row=2)]
    assert [c.value for c in tags[1]] == ["Tag", "Equipment type", "Description", "Tile"]
    assert [r[0] for r in rows] == ["P-101A", "V-201", "V-202", "P-102", "T-301"]
    assert rows[0][3] == "top-left, top-right" and rows[0][2] == "Feed pump"
    assert rows[3][3] == "bottom-right"
    assert tags.freeze_panes == "A2" and tags.auto_filter.ref == "A1:D6"
    assert tags["A1"].font.bold

    summary = {r[0].value: r[1].value for r in wb["Summary"].iter_rows(min_row=2) if r[0].value}
    assert summary["Pump"] == 2 and summary["Valve"] == 2 and summary["Vessel"] == 1 and summary["Total"] == 5
    assert wb["Summary"].freeze_panes == "A2"
    assert artifact.preview.startswith("5 unique tags")
    _assert_valid_zip(path)


def test_excel_findings_sheet():
    note = sample_note()
    artifact = make_excel(note.findings)
    ws = load_workbook(_path_of(artifact))["Findings"]
    assert [c.value for c in ws[1]] == ["S.No", "Item", "Observation", "Severity", "Source page"]
    assert ws.max_row == 1 + len(note.findings)
    assert ws["D2"].value == "High" and ws["D2"].fill.fgColor.rgb.endswith(office.SEVERITY_FILLS[Severity.HIGH])
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref
    assert ws.column_dimensions["C"].width > ws.column_dimensions["A"].width


def test_excel_bad_input_is_error_code():
    with pytest.raises(OfficeError) as exc:
        make_excel({"tags": []})
    assert exc.value.code == "BAD_REQUEST"


# ---------------------------------------------------------------- PowerPoint
def test_ppt_splits_bullets_six_per_slide():
    bullets = [f"Point {i}" for i in range(1, 15)]
    artifact = make_ppt("Tank T-104 inspection summary", [("Findings", bullets), ("Next steps", [])])
    prs = Presentation(str(_path_of(artifact)))
    slides = list(prs.slides)
    assert len(slides) == 1 + 3 + 1
    assert [s.shapes.title.text for s in slides[1:4]] == ["Findings", "Findings (cont.)", "Findings (cont.)"]
    counts = [len(s.placeholders[1].text_frame.paragraphs) for s in slides[1:4]]
    assert counts == [6, 6, 2]
    assert slides[4].placeholders[1].text_frame.text == NOT_APPLICABLE
    for slide in slides:
        assert any(DRAFT_FOOTER in sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)
    _assert_valid_zip(_path_of(artifact))


# ---------------------------------------------------------------- artifact endpoint
def test_artifact_download_endpoint():
    from backend.main import app

    artifact = make_word(sample_note())
    path = _path_of(artifact)
    with TestClient(app) as client:
        resp = client.get(artifact.download_url)
        assert resp.status_code == 200
        assert resp.content == path.read_bytes()
        assert resp.headers["content-type"] == office.MEDIA_TYPES[ArtifactKind.DOCX]
        disposition = resp.headers["content-disposition"]
        assert disposition.startswith("attachment") and artifact.filename in disposition

        missing = client.get(f"{API_PREFIX}/artifacts/a_000000000000")
        assert missing.status_code == 404 and missing.json()["error"]["code"] == "FILE_NOT_FOUND"
        garbage = client.get(f"{API_PREFIX}/artifacts/not-an-id")
        assert garbage.status_code == 404

        bad = client.get(f"{API_PREFIX}/artifacts/..%5Cx")
        assert bad.status_code == 400 and bad.json()["error"]["code"] == "BAD_REQUEST"


def test_artifact_registry_survives_restart():
    artifact = make_excel(_tag_list())
    fresh = fresh_office_module()
    path, stored = fresh.artifact_path(artifact.artifact_id)
    assert path.name == artifact.filename and stored.kind == ArtifactKind.XLSX


# ---------------------------------------------------------------- optional LibreOffice check
_SOFFICE = shutil.which("soffice") or shutil.which("soffice.exe")


@pytest.mark.slow
@pytest.mark.skipif(_SOFFICE is None, reason="LibreOffice (soffice) not installed")
def test_libreoffice_converts_all_files_to_pdf(tmp_path):
    paths = [
        _path_of(make_word(sample_note())),
        _path_of(make_excel(_tag_list())),
        _path_of(make_ppt("Summary", [("Findings", ["a", "b"])])),
    ]
    out = tmp_path / "pdf"
    for path in paths:
        subprocess.run([_SOFFICE, "--headless", "--convert-to", "pdf", "--outdir", str(out), str(path)],
                       check=True, capture_output=True, timeout=120)
        assert (out / (path.stem + ".pdf")).stat().st_size > 0


def test_make_template_script_builds_template(tmp_path):
    spec = importlib.util.spec_from_file_location("make_template", _REPO_ROOT / "scripts" / "make_template.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    path = module.build(tmp_path / "approval_note.docx")
    doc = Document(str(path))
    text = "\n".join(_doc_blocks(doc))
    for placeholder in ("{{ref_no}}", "{{subject}}", "{{finding.severity}}", "{{sop.page}}", "{{sop_note}}",
                        "{{cost_implication}}"):
        assert placeholder in text
    assert not doc.inline_shapes                                   # no logos or images
    assert DRAFT_FOOTER in doc.sections[0].footer.paragraphs[0].text


def test_word_sop_auto_note():
    auto = Document(str(_path_of(make_word(sample_note(), sop_auto=True))))
    blocks = _doc_blocks(auto)
    assert office.SOP_AUTO_NOTE in blocks
    assert blocks.index(office.SOP_AUTO_NOTE) < blocks.index("Recommendation")   # right under the SOP table

    plain = _doc_blocks(Document(str(_path_of(make_word(sample_note())))))
    assert office.SOP_AUTO_NOTE not in plain and not any("{{" in b for b in plain)

    empty = _doc_blocks(Document(str(_path_of(make_word(sample_note(sop_references=[]), sop_auto=True)))))
    assert office.SOP_AUTO_NOTE not in empty                                       # nothing auto-added to label
