"""
Tests for backend.tools.documents: text-layer/OCR/vision extraction order,
P&ID tiling, resizing, and the on-disk cache. All vision (Ollama) calls are
mocked except one live smoke test, which is skipped automatically if Ollama
is unreachable and marked slow.
"""
from __future__ import annotations

import io
import json
import re
import time
from pathlib import Path

import httpx
import pymupdf
import pytest
from PIL import Image, ImageDraw, ImageFont

import backend.tools.documents as documents
from backend.llm_client import ChatResult
from backend.settings import settings
from backend.tools.documents import PID_TILE_OVERLAP, extract


def _ollama_available() -> bool:
    try:
        resp = httpx.get(settings.OLLAMA_HOST, timeout=3.0)
        return resp.status_code < 500
    except Exception:
        return False


OLLAMA_UP = _ollama_available()
skip_if_no_ollama = pytest.mark.skipif(not OLLAMA_UP, reason="Ollama is not reachable at settings.OLLAMA_HOST")


@pytest.fixture(autouse=True)
def _isolate_cache(tmp_path, monkeypatch):
    """Gives every test its own cache dir so tests can't see each other's cached results."""
    monkeypatch.setattr(settings, "WB_CACHE_DIR", tmp_path / "cache")
    yield


def _font(size: int = 28):
    try:
        return ImageFont.truetype("arial.ttf", size)
    except Exception:
        return ImageFont.load_default()


def _text_image(lines: list[str], size=(900, 300)) -> Image.Image:
    image = Image.new("RGB", size, color="white")
    draw = ImageDraw.Draw(image)
    font = _font(28)
    y = 20
    for line in lines:
        draw.text((20, y), line, fill="black", font=font)
        y += 45
    return image


def _make_text_layer_pdf(path: Path, text: str) -> None:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), text, fontsize=14)
    doc.save(path)
    doc.close()


def _make_scanned_pdf(path: Path, pages_lines: list[list[str]]) -> None:
    """Builds an image-only PDF: each page is a rendered image with no text layer."""
    doc = pymupdf.open()
    for lines in pages_lines:
        image = _text_image(lines)
        buf = io.BytesIO()
        image.save(buf, format="PNG")
        page = doc.new_page(width=image.width, height=image.height)
        page.insert_image(pymupdf.Rect(0, 0, image.width, image.height), stream=buf.getvalue())
    doc.save(path)
    doc.close()


def _make_pid_image() -> Image.Image:
    image = Image.new("RGB", (800, 600), color="white")
    draw = ImageDraw.Draw(image)
    font = _font(26)
    draw.rectangle((50, 50, 200, 150), outline="black", width=3)
    draw.text((60, 90), "P-101", fill="black", font=font)
    draw.rectangle((600, 50, 750, 150), outline="black", width=3)
    draw.text((610, 90), "V-201", fill="black", font=font)
    draw.rectangle((50, 450, 200, 550), outline="black", width=3)
    draw.text((60, 490), "FT-301", fill="black", font=font)
    draw.line((0, 300, 800, 300), fill="black", width=2)
    return image


SCANNED_TEXTS = [
    ["Corrosion found on line 6-P-101,", "wall loss 2.1 mm."],
    ["Valve V-204 leaking at flange,", "recommend replacement gasket."],
    ["Support bracket cracked near", "shell course 2, north side."],
]


# ---------------------------------------------------------------- text layer
def test_text_layer_pdf_uses_text_layer_method(tmp_path):
    pdf_path = tmp_path / "layer.pdf"
    _make_text_layer_pdf(pdf_path, "This PDF has a real embedded text layer for testing.")

    result = extract(pdf_path)

    assert len(result.pages) == 1
    assert result.pages[0].method == "text_layer"
    assert "text layer" in result.pages[0].text.lower()
    assert result.pages[0].error is None


# ---------------------------------------------------------------- OCR
@pytest.mark.parametrize("lines", SCANNED_TEXTS)
def test_scanned_pdf_uses_ocr_method_and_extracts_key_words(tmp_path, lines):
    pdf_path = tmp_path / "scanned.pdf"
    _make_scanned_pdf(pdf_path, [lines])

    result = extract(pdf_path)

    assert len(result.pages) == 1
    page = result.pages[0]
    assert page.method == "ocr"
    assert page.ocr_confidence is not None

    key_words = [w.strip(",.") for w in " ".join(lines).lower().split() if len(w) > 3]
    lowered = page.text.lower()
    hits = sum(1 for w in key_words if w in lowered)
    assert hits >= max(1, len(key_words) // 2), f"too few OCR word matches in {page.text!r}"


def test_two_page_scanned_pdf_has_correct_page_numbers(tmp_path):
    pdf_path = tmp_path / "two_page.pdf"
    _make_scanned_pdf(pdf_path, [SCANNED_TEXTS[0], SCANNED_TEXTS[1]])

    result = extract(pdf_path)

    assert [p.page for p in result.pages] == [1, 2]
    assert all(p.method == "ocr" for p in result.pages)


def test_joined_text_has_page_markers(tmp_path):
    pdf_path = tmp_path / "two_page.pdf"
    _make_scanned_pdf(pdf_path, [SCANNED_TEXTS[0], SCANNED_TEXTS[1]])

    result = extract(pdf_path)
    joined = result.joined_text()

    assert "--- Page 1 ---" in joined
    assert "--- Page 2 ---" in joined


# ---------------------------------------------------------------- vision fallback (mocked)
def test_poor_ocr_triggers_vision_fallback(tmp_path, monkeypatch):
    # A pure noise image: Tesseract should find ~nothing readable on it.
    noisy = Image.effect_noise((400, 200), 60).convert("RGB")
    buf = io.BytesIO()
    noisy.save(buf, format="PNG")

    doc = pymupdf.open()
    page = doc.new_page(width=noisy.width, height=noisy.height)
    page.insert_image(pymupdf.Rect(0, 0, noisy.width, noisy.height), stream=buf.getvalue())
    pdf_path = tmp_path / "noisy.pdf"
    doc.save(pdf_path)
    doc.close()

    calls: list[str] = []

    def fake_chat(model_id, messages, tools=None, images=None, *, purpose="chat"):
        calls.append(purpose)
        return ChatResult(text="vision fallback text")

    monkeypatch.setattr(documents, "chat", fake_chat)

    result = extract(pdf_path)

    assert result.pages[0].method == "vision"
    assert result.pages[0].text == "vision fallback text"
    assert calls, "the vision model was never called"


# ---------------------------------------------------------------- P&ID tiling
def test_make_pid_tiles_returns_four_tiles_with_overlap():
    image = _make_pid_image()
    tiles = documents._make_pid_tiles(image)

    assert [name for name, _ in tiles] == ["top-left", "top-right", "bottom-left", "bottom-right"]

    by_name = dict(tiles)
    horizontal_overlap = by_name["top-left"].width + by_name["top-right"].width - image.width
    vertical_overlap = by_name["top-left"].height + by_name["bottom-left"].height - image.height
    assert horizontal_overlap > 0
    assert vertical_overlap > 0
    # each tile extends past the midline by (half-width * overlap), so the
    # shared strip between two neighbouring tiles is twice that
    expected_overlap = 2 * round(image.width / 2 * PID_TILE_OVERLAP)
    assert abs(horizontal_overlap - expected_overlap) <= 1


def test_pid_extraction_produces_four_tiles_each_within_max_px(tmp_path, monkeypatch):
    image = _make_pid_image()
    image_path = tmp_path / "pid.png"
    image.save(image_path)

    seen_sizes: list[tuple[int, int]] = []

    def fake_chat(model_id, messages, tools=None, images=None, *, purpose="chat"):
        tile_image = Image.open(io.BytesIO(images[0]))
        seen_sizes.append(tile_image.size)
        return ChatResult(text="P-101\nV-201")

    monkeypatch.setattr(documents, "chat", fake_chat)
    monkeypatch.setattr(documents, "_ocr_pid_tile", lambda tile: "")   # OCR finds nothing -> 4-tile vision path

    result = extract(image_path, kind="pid")

    assert len(result.pages) == 4
    assert {p.tile for p in result.pages} == {"top-left", "top-right", "bottom-left", "bottom-right"}
    assert all(p.method == "vision_tiles" for p in result.pages)
    assert len(seen_sizes) == 4
    for w, h in seen_sizes:
        assert max(w, h) <= settings.WB_VISION_MAX_PX


# ---------------------------------------------------------------- resize
def test_resize_for_vision_keeps_aspect_ratio_and_max_px():
    image = Image.new("RGB", (2000, 1000), color="white")
    resized = documents._resize_for_vision(image)

    assert max(resized.size) <= settings.WB_VISION_MAX_PX
    assert abs((2000 / 1000) - (resized.width / resized.height)) < 0.01


def test_resize_for_vision_leaves_small_images_untouched():
    image = Image.new("RGB", (100, 50), color="white")
    resized = documents._resize_for_vision(image)
    assert resized.size == (100, 50)


# ---------------------------------------------------------------- cache
def test_cache_second_call_skips_tesseract_and_vision(tmp_path, monkeypatch):
    pdf_path = tmp_path / "scanned.pdf"
    _make_scanned_pdf(pdf_path, [SCANNED_TEXTS[0]])

    real_ocr_page = documents._ocr_page
    ocr_calls: list[int] = []

    def counting_ocr_page(image):
        ocr_calls.append(1)
        return real_ocr_page(image)

    vision_calls: list[int] = []

    def fake_chat(model_id, messages, tools=None, images=None, *, purpose="chat"):
        vision_calls.append(1)
        return ChatResult(text="should not be called")

    monkeypatch.setattr(documents, "_ocr_page", counting_ocr_page)
    monkeypatch.setattr(documents, "chat", fake_chat)

    first = extract(pdf_path)
    assert first.from_cache is False
    assert len(ocr_calls) == 1

    second = extract(pdf_path)
    assert second.from_cache is True
    assert len(ocr_calls) == 1  # no new OCR call
    assert len(vision_calls) == 0  # OCR was good enough both times
    assert second.pages[0].text == first.pages[0].text
    assert second.pages[0].method == first.pages[0].method


def test_corrupted_cache_file_is_ignored_and_redone(tmp_path):
    pdf_path = tmp_path / "scanned.pdf"
    _make_scanned_pdf(pdf_path, [SCANNED_TEXTS[0]])

    first = extract(pdf_path)
    assert first.from_cache is False

    cache_dir = documents._cache_dir()
    cache_files = list(cache_dir.glob("*.json"))
    assert len(cache_files) == 1
    cache_files[0].write_text("{not valid json", encoding="utf-8")

    second = extract(pdf_path)
    assert second.from_cache is False  # corrupted cache -> ignored, redone
    assert second.pages[0].text == first.pages[0].text


# ---------------------------------------------------------------- plain text / office formats
@pytest.mark.parametrize("suffix", [".txt", ".md", ".py", ".csv"])
def test_text_like_files_use_text_file_method(tmp_path, suffix):
    path = tmp_path / f"note{suffix}"
    path.write_text("Corrosion found on line 6-P-101, wall loss 2.1 mm.", encoding="utf-8")

    result = extract(path)

    assert len(result.pages) == 1
    assert result.pages[0].method == "text_file"
    assert "corrosion" in result.pages[0].text.lower()
    assert result.pages[0].error is None


def test_docx_extracts_paragraphs_and_table_cells(tmp_path):
    import docx

    doc_path = tmp_path / "note.docx"
    document = docx.Document()
    document.add_paragraph("Inspection summary for line 6-P-101.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Wall loss"
    table.cell(0, 1).text = "2.1 mm"
    document.save(doc_path)

    result = extract(doc_path)

    assert len(result.pages) == 1
    page = result.pages[0]
    assert page.method == "docx"
    assert "inspection summary" in page.text.lower()
    assert "wall loss" in page.text.lower()
    assert "2.1 mm" in page.text.lower()


def test_xlsx_extracts_sheet_names_and_cell_values(tmp_path):
    import openpyxl

    xlsx_path = tmp_path / "note.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Findings"
    sheet["A1"] = "Tag"
    sheet["B1"] = "Wall loss mm"
    sheet["A2"] = "6-P-101"
    sheet["B2"] = 2.1
    workbook.save(xlsx_path)

    result = extract(xlsx_path)

    assert len(result.pages) == 1
    page = result.pages[0]
    assert page.method == "xlsx"
    assert "findings" in page.text.lower()
    assert "6-p-101" in page.text.lower()
    assert "2.1" in page.text


# ---------------------------------------------------------------- live vision (needs Ollama)
@skip_if_no_ollama
@pytest.mark.slow
def test_live_vision_call_on_one_pid_tile_returns_text():
    image = _make_pid_image()
    tile_name, tile_image = documents._make_pid_tiles(image)[0]

    start = time.monotonic()
    text = documents._vision_extract_page(tile_image, purpose="test_live_vision")
    duration_s = time.monotonic() - start

    assert text.strip() != ""
    print(f"\nLive vision call on tile {tile_name!r} took {duration_s:.2f}s, returned {len(text)} chars")


# ---------------------------------------------------------------- deskew
def _document_page() -> Image.Image:
    """A page-like image: 22 lines of body text on white, like a report page."""
    image = Image.new("RGB", (1240, 1754), color="white")
    draw = ImageDraw.Draw(image)
    font = _font(26)
    for row in range(22):
        draw.text((90, 120 + row * 64), f"{row + 1}. Shell course {row % 4 + 1} north side wall 6.{row % 9} mm "
                                        f"against 8.0 mm nominal, severity High", fill="black", font=font)
    return image


@pytest.mark.parametrize("true_angle", [1.5, -1.0])
def test_detect_skew_finds_rotation_within_tolerance(true_angle):
    skewed = _document_page().rotate(true_angle, resample=Image.BICUBIC, expand=True, fillcolor="white")
    assert documents.detect_skew(skewed) == pytest.approx(true_angle, abs=0.3)


def test_deskew_leaves_straight_page_unrotated():
    page = _document_page()
    straight, angle = documents._deskew(page)
    assert abs(angle) < documents.DESKEW_MIN_ANGLE
    assert straight is page  # not rotated at all, not even by a tiny angle


def test_deskew_rotates_skewed_page_and_angle_is_recorded(tmp_path):
    skewed = _text_image(SCANNED_TEXTS[0]).rotate(2.0, resample=Image.BICUBIC, expand=True, fillcolor="white")
    straight, angle = documents._deskew(skewed)
    assert angle == pytest.approx(2.0, abs=0.3) and straight is not skewed
    assert documents.detect_skew(straight) == pytest.approx(0.0, abs=0.3)

    path = tmp_path / "skewed.png"
    skewed.save(path)
    page = extract(path).pages[0]
    assert page.method == "ocr" and page.deskew_angle == pytest.approx(2.0, abs=0.3)


def test_cache_key_includes_deskew_setting(tmp_path, monkeypatch):
    file_bytes = b"same bytes"
    with_deskew = documents._cache_key(file_bytes, kind="auto")
    monkeypatch.setattr(documents, "DESKEW_ENABLED", False)
    assert documents._cache_key(file_bytes, kind="auto") != with_deskew


def _pre_deskew_cache_key(file_bytes: bytes, kind: str) -> str:
    """The cache key exactly as it was built before deskew (cache version 1)."""
    import hashlib

    hasher = hashlib.sha256()
    hasher.update(file_bytes)
    for part in ("1", kind, str(settings.WB_OCR_DPI), str(settings.WB_VISION_MAX_PX), documents._vision_model_name()):
        hasher.update(b"\x00")
        hasher.update(part.encode("utf-8"))
    return hasher.hexdigest()


def test_old_cache_made_without_deskew_is_not_reused(tmp_path):
    pdf_path = tmp_path / "scanned.pdf"
    _make_scanned_pdf(pdf_path, [SCANNED_TEXTS[0]])
    old_key = _pre_deskew_cache_key(pdf_path.read_bytes(), "auto")
    stale = {"pages": [{"page": 1, "text": "STALE OCR WITHOUT DESKEW", "method": "ocr", "ocr_confidence": 90.0}]}
    (documents._cache_dir() / f"{old_key}.json").write_text(json.dumps(stale), encoding="utf-8")

    result = extract(pdf_path)

    assert result.from_cache is False
    assert "STALE" not in result.pages[0].text and result.pages[0].deskew_angle is not None


# ---------------------------------------------------------------- live: B6 scanned demo reports
_DEMO_INPUTS = Path(__file__).resolve().parents[2] / "demo" / "inputs"

# Key values per page, from demo/expected.md (findings, background, cost). Numbers that appear ONLY
# inside the ruled thickness tables are left out: Tesseract drops those cells (docs/known_issues.md).
# That includes 6.8 mm (report 1 bottom plate), which expected.md lists as a "should".
_KEY_VALUES = {
    "scenario_a_report_1": {
        1: ["6.1 mm", "8.0 mm", "24%", "1.2 mm", "18 ppm"],
        2: ["Rs 4,50,000", "3,20,000", "90,000", "40,000"],
    },
    "scenario_a_report_2": {
        1: ["5.20 mm", "7.11 mm", "27%", "12 ppm", "5.90 mm", "17%", "6.10 mm", "14%", "4.80 mm", "0.38 mm"],
        2: ["Rs 2,85,000", "1,95,000", "55,000", "35,000"],
    },
}


def _normalize(text: str) -> str:
    """expected.md matching rules: case-insensitive, extra spaces ignored (also OCR's 'Rs 90, 000')."""
    text = re.sub(r"(?<=\d)\s*([,.])\s*(?=\d)", r"\1", text.lower())  # only inside numbers
    text = re.sub(r"(?<=\d)\s+%", "%", text)
    return re.sub(r"\s+", " ", text)


def _has_value(text: str, value: str) -> bool:
    return re.search(rf"(?<![\d.,]){re.escape(_normalize(value))}(?!\d)", _normalize(text)) is not None


def _answer_key_pages(name: str) -> dict[int, str]:
    text = (_DEMO_INPUTS / "_source_text" / f"{name}.txt").read_text(encoding="utf-8")
    parts = re.split(r"^--- Page (\d+) ---$", text, flags=re.MULTILINE)
    return {int(parts[i]): parts[i + 1] for i in range(1, len(parts), 2)}


@pytest.mark.slow
@pytest.mark.parametrize("name", ["scenario_a_report_1", "scenario_a_report_2"])
def test_demo_report_ocr_keeps_every_severity_and_key_value(name):
    pdf_path = _DEMO_INPUTS / f"{name}.pdf"
    if not pdf_path.exists():
        pytest.skip(f"B6 demo data not present: {pdf_path}")
    key_pages = _answer_key_pages(name)

    result = extract(pdf_path)

    got = {p.page: p for p in result.pages}
    assert sorted(got) == sorted(key_pages)
    for number, key_text in key_pages.items():
        page = got[number]
        assert page.method == "ocr" and page.deskew_angle is not None
        for word in ("High", "Medium", "Low"):
            expected = len(re.findall(rf"\b{word}\b", key_text))
            found = len(re.findall(rf"\b{word}\b", page.text))
            assert found == expected, f"page {number}: {word!r} expected {expected}, found {found}"
        missing = [v for v in _KEY_VALUES[name][number] if not _has_value(page.text, v)]
        assert missing == [], f"page {number}: key values not found: {missing}"
