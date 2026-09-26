"""
Tests for backend/agent.py and backend/agent_tools.py (ticket A8, first half).

No Ollama: a scripted fake replaces the Ollama client (llm_client._client), so
the real chat()/chat_json_meta() code runs, including the <tool_call> text
fallback parser and JSON validation.
"""
from __future__ import annotations

import json
import time
from pathlib import Path
from types import SimpleNamespace
from typing import Any, Callable, Optional

import pymupdf
import pytest
from docx import Document

from backend import agent, agent_tools, llm_client
from backend.agent import PINNED_MESSAGES, SYSTEM_PROMPT, TRIM_NOTE, estimate_tokens, trim_conversation
from backend.file_store import file_store
from backend.registry import registry
from backend.settings import settings
from backend.task_store import TaskStore
from backend.tools import office
from backend.tools.office import COST_PLACEHOLDER, cost_text
from shared.contracts import (
    TERMINAL_STATUSES,
    ApprovalNote,
    EventType,
    Finding,
    KBHit,
    PlanStep,
    RouteDecision,
    Severity,
    TaskMode,
    TaskStatus,
    TaskType,
)

EVENT_KEYS = {
    EventType.ROUTE: {"decision"},
    EventType.PLAN: {"steps"},
    EventType.STEP_START: {"index", "title"},
    EventType.LLM_CALL: {"model_id", "purpose", "duration_ms", "tokens_out"},
    EventType.TOOL_CALL: {"tool", "args"},
    EventType.TOOL_RESULT: {"tool", "ok", "summary", "duration_ms"},
    EventType.ARTIFACT: {"artifact"},
    EventType.LOG: {"level", "text"},
    EventType.FINAL: {"answer"},
    EventType.ERROR: {"error"},
}

REPORT_PAGES = [
    "Inspection report for storage tank T-104. Page one. Shell course 2 north side measured wall "
    "thickness 6.1 mm against nominal 8.0 mm.",
    "Page two. Bottom plate pitting up to 1.2 mm. Contractor quote attached: Rs 4,50,000 for plate "
    "replacement of shell course 2.",
]
HOT_WORK = KBHit(text="Minimum wall thickness criteria for repair ...", source="SOP-INSP-012.pdf", page=4, score=0.81)
WEAK = KBHit(text="Canteen hygiene rules ...", source="SOP-ADM-001.pdf", page=2, score=0.31)


# ---------------------------------------------------------------- fake Ollama
def _response(content: str = "", tool_calls: Optional[list] = None, eval_count: int = 42):
    calls = [SimpleNamespace(function=SimpleNamespace(name=n, arguments=a)) for n, a in (tool_calls or [])]
    return SimpleNamespace(message=SimpleNamespace(content=content, tool_calls=calls or None), eval_count=eval_count)


class FakeOllama:
    """Scripted replies. `steps` are agent turns: ("tool", name, args) or ("text", str).
    `json` maps a schema title (e.g. "ApprovalNote", "_Plan") to a list of raw JSON strings."""

    def __init__(self, steps: list, json_replies: Optional[dict[str, list[str]]] = None,
                 repeat_last: bool = False, delay_s: float = 0.0,
                 on_step: Optional[Callable[[int], None]] = None, step_delay_s: float = 0.0) -> None:
        self.steps = list(steps)
        self.json = {k: list(v) for k, v in (json_replies or {}).items()}
        self.repeat_last = repeat_last
        self.delay_s = delay_s
        self.on_step = on_step
        self.step_delay_s = step_delay_s            # delay for agent turns only (not JSON calls)
        self.step_calls: list[list[dict]] = []

    def chat(self, **kwargs):
        if self.delay_s:
            time.sleep(self.delay_s)
        fmt = kwargs.get("format")
        if fmt:
            replies = self.json.get(fmt.get("title", ""), [])
            return _response(replies.pop(0) if replies else "{}")
        self.step_calls.append([dict(m) for m in kwargs["messages"]])
        if self.step_delay_s:
            time.sleep(self.step_delay_s)
        if self.on_step:
            self.on_step(len(self.step_calls))
        if not self.steps:
            step = ("text", "Nothing more to do.")
        elif self.repeat_last and len(self.steps) == 1:
            step = self.steps[0]
        else:
            step = self.steps.pop(0)
        if step[0] == "tool":
            return _response("", [(step[1], step[2])])
        return _response(step[1])


def plan_json(*tools: str) -> str:
    steps = [{"index": i, "title": t.replace("_", " "), "tool": t} for i, t in enumerate(tools, 1)]
    return json.dumps({"steps": steps})


def note_json(**overrides) -> str:
    data = dict(
        ref_no="INSP-FAKE", date="01-01-1999", subject="Approval for repair of Tank T-104 shell course 2",
        background="Annual inspection of tank T-104.",
        findings=[Finding(item="Shell course 2, north side", observation="Wall thinning to 6.1 mm vs 8.0 mm",
                          severity=Severity.HIGH, source_page=1)],
        sop_references=["SOP-INSP-012.pdf, p.4"], recommendation="Repair shell course 2 before next run.",
        cost_implication=None,
    )
    data.update(overrides)
    return ApprovalNote(**data).model_dump_json()


# ---------------------------------------------------------------- fixtures
@pytest.fixture(autouse=True)
def isolated(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "WB_WORKSPACE_DIR", tmp_path / "workspace")
    monkeypatch.setattr(settings, "WB_CACHE_DIR", tmp_path / "cache")
    monkeypatch.setattr(settings, "WB_LOG_DIR", tmp_path / "logs")
    model = registry.model_for_task(TaskType.DOCUMENT)
    decision = RouteDecision(task_type=TaskType.DOCUMENT, model_id=model.id, ollama_name=model.ollama_name,
                             reason="test", layer="forced", confidence=1.0)
    monkeypatch.setattr(agent, "route", lambda request: decision)
    monkeypatch.setattr(agent_tools.knowledge, "search", lambda query, top_k=4: [HOT_WORK, WEAK])


@pytest.fixture
def store():
    s = TaskStore()
    s.start(agent.run)
    yield s
    s.stop()


@pytest.fixture
def report_file_id(tmp_path) -> str:
    path = tmp_path / "report.pdf"
    doc = pymupdf.open()
    for text in REPORT_PAGES:
        doc.new_page().insert_textbox(pymupdf.Rect(60, 60, 540, 780), text, fontsize=12)
    doc.save(str(path))
    doc.close()
    return file_store.save("inspection_report.pdf", path.read_bytes(), "application/pdf").file_id


def use_fake(monkeypatch, fake: FakeOllama) -> FakeOllama:
    monkeypatch.setattr(llm_client, "_client", lambda: fake)
    return fake


def run_task(store: TaskStore, message: str, file_ids: Optional[list[str]] = None, timeout_s: float = 30):
    state = store.create(message, file_ids or [], TaskMode.AGENT, None)
    deadline = time.time() + timeout_s
    while time.time() < deadline:
        current = store.get(state.task_id)
        if current.status in TERMINAL_STATUSES:
            events, _, _ = store.events(state.task_id, 0)
            return current, events
        time.sleep(0.02)
    raise AssertionError("task did not finish")


def happy_steps(file_id: str, note: Optional[str] = None) -> tuple[list, dict]:
    steps = [
        ("tool", "read_document", {"file_id": file_id}),
        ("tool", "search_knowledge", {"query": "minimum wall thickness repair criteria"}),
        ("tool", "draft_approval_note", {"doc_id": "doc_1", "kb_ref_ids": ["kb_1"]}),
        ("tool", "finish", {"answer": "Approval note drafted."}),
    ]
    json_replies = {"_Plan": [plan_json("read_document", "search_knowledge", "draft_approval_note", "finish")],
                    "ApprovalNote": [note or note_json()]}
    return steps, json_replies


def _docx_of(state) -> Document:
    found = office.artifact_path(state.artifacts[0].artifact_id)
    return Document(str(found[0]))


def _warns(events) -> list[str]:
    return [e.data["text"] for e in events if e.type == EventType.LOG and e.data["level"] == "warn"]


# ---------------------------------------------------------------- happy path
def test_happy_path_events_order_and_keys(monkeypatch, store, report_file_id):
    steps, json_replies = happy_steps(report_file_id)
    use_fake(monkeypatch, FakeOllama(steps, json_replies))

    state, events = run_task(store, "Draft an approval note for this inspection report", [report_file_id])

    assert state.status == TaskStatus.SUCCEEDED, state.error
    assert state.final_answer == "Approval note drafted."
    types = [e.type for e in events]
    T = EventType
    assert types == [
        T.ROUTE, T.LLM_CALL, T.PLAN,
        T.LLM_CALL, T.STEP_START, T.TOOL_CALL, T.TOOL_RESULT,                            # read_document
        T.LLM_CALL, T.STEP_START, T.TOOL_CALL, T.TOOL_RESULT,                            # search_knowledge
        T.LLM_CALL, T.STEP_START, T.TOOL_CALL, T.LLM_CALL, T.ARTIFACT, T.TOOL_RESULT,    # draft note
        T.LLM_CALL, T.STEP_START, T.TOOL_CALL, T.TOOL_RESULT, T.FINAL,                   # finish
    ]
    for event in events:
        assert set(event.data) == EVENT_KEYS[event.type], event
    assert [e.seq for e in events] == list(range(1, len(events) + 1))
    llm = [e.data for e in events if e.type == T.LLM_CALL]
    assert all(d["tokens_out"] == 42 for d in llm)
    assert [d["purpose"] for d in llm] == ["plan", "agent_step", "agent_step", "agent_step",
                                            "draft_approval_note", "agent_step"]
    assert [p.tool for p in state.plan] == ["read_document", "search_knowledge", "draft_approval_note", "finish"]
    assert len(state.artifacts) == 1 and state.artifacts[0].filename.endswith("_approval_note.docx")
    read_result = next(e for e in events if e.type == T.TOOL_RESULT and e.data["tool"] == "read_document")
    assert read_result.data["summary"].startswith("doc_1: inspection_report.pdf, 2 pages")
    assert [e.step for e in events if e.type == T.STEP_START] == [1, 2, 3, 4]


def test_tool_results_sent_to_model_are_short(monkeypatch, store, report_file_id):
    steps, json_replies = happy_steps(report_file_id)
    fake = use_fake(monkeypatch, FakeOllama(steps, json_replies))
    run_task(store, "Draft an approval note", [report_file_id])
    tool_messages = [m for m in fake.step_calls[-1] if m.get("role") == "tool"]
    assert len(tool_messages) == 3
    assert all(len(m["content"]) <= agent_tools.TOOL_RESULT_MAX_CHARS for m in tool_messages)
    assert fake.step_calls[-1][0]["content"] == SYSTEM_PROMPT


def test_tool_call_written_as_text_still_works(monkeypatch, store):
    text_call = '<tool_call>{"name": "finish", "arguments": {"answer": "Parsed from text."}}</tool_call>'
    use_fake(monkeypatch, FakeOllama([("text", text_call)], {"_Plan": [plan_json("finish")]}))
    state, events = run_task(store, "Say hello")
    assert state.status == TaskStatus.SUCCEEDED and state.final_answer == "Parsed from text."
    assert any(e.type == EventType.TOOL_CALL and e.data["tool"] == "finish" for e in events)


def test_plain_text_reply_is_treated_as_finish(monkeypatch, store):
    use_fake(monkeypatch, FakeOllama([("text", "A relief valve protects equipment from overpressure.")],
                                     {"_Plan": [plan_json("finish")]}))
    state, _ = run_task(store, "What does a relief valve do?")
    assert state.status == TaskStatus.SUCCEEDED and state.final_answer.startswith("A relief valve")


def test_bad_plan_json_twice_uses_default_plan(monkeypatch, store):
    use_fake(monkeypatch, FakeOllama([("tool", "finish", {"answer": "ok"})], {"_Plan": ["not json", "{}"]}))
    state, events = run_task(store, "Explain LOTO")
    assert state.status == TaskStatus.SUCCEEDED
    assert [p.tool for p in state.plan] == ["search_knowledge", "finish"]
    assert any("default plan" in w for w in _warns(events))


# ---------------------------------------------------------------- errors inside the loop
def test_unknown_tool_error_is_fed_back_and_loop_continues(monkeypatch, store):
    fake = use_fake(monkeypatch, FakeOllama([("tool", "delete_everything", {}), ("tool", "finish", {"answer": "ok"})],
                                            {"_Plan": [plan_json("finish")]}))
    state, events = run_task(store, "Do something")
    assert state.status == TaskStatus.SUCCEEDED
    bad = next(e for e in events if e.type == EventType.TOOL_RESULT and e.data["tool"] == "delete_everything")
    assert bad.data["ok"] is False and "unknown tool" in bad.data["summary"]
    assert "unknown tool" in fake.step_calls[1][-1]["content"]           # the model saw the error
    assert [e.data["index"] for e in events if e.type == EventType.STEP_START] == [1, 2]


def test_bad_arguments_are_reported(monkeypatch, store):
    fake = use_fake(monkeypatch, FakeOllama([("tool", "search_knowledge", {"top_k": 3}),
                                             ("tool", "finish", {"answer": "ok"})], {"_Plan": [plan_json("finish")]}))
    state, _ = run_task(store, "Search")
    assert state.status == TaskStatus.SUCCEEDED
    assert "missing required argument(s): query" in fake.step_calls[1][-1]["content"]


def test_tool_exception_emits_error_and_worker_survives(monkeypatch, store):
    def broken_search(query, top_k=4):
        raise RuntimeError("chroma exploded")

    monkeypatch.setattr(agent_tools.knowledge, "search", broken_search)
    use_fake(monkeypatch, FakeOllama([("tool", "search_knowledge", {"query": "x"}),
                                      ("tool", "finish", {"answer": "recovered"})], {"_Plan": [plan_json("finish")]}))
    state, events = run_task(store, "Search the SOPs")
    errors = [e for e in events if e.type == EventType.ERROR]
    assert len(errors) == 1 and errors[0].data["error"]["code"] == "INTERNAL"
    assert "chroma exploded" in errors[0].data["error"]["message"]
    assert state.status == TaskStatus.SUCCEEDED and state.final_answer == "recovered"

    use_fake(monkeypatch, FakeOllama([("tool", "finish", {"answer": "second task ran"})],
                                     {"_Plan": [plan_json("finish")]}))
    second, _ = run_task(store, "Next task")
    assert second.status == TaskStatus.SUCCEEDED and second.final_answer == "second task ran"


def test_agent_crash_marks_failed_and_next_task_runs(monkeypatch, store):
    def crash(self):
        raise ValueError("bug in agent")

    original_run = agent.AgentRun.run
    monkeypatch.setattr(agent.AgentRun, "run", crash)
    state, events = run_task(store, "anything")
    assert state.status == TaskStatus.FAILED and state.error.code == "INTERNAL"
    assert events[-1].type == EventType.ERROR
    monkeypatch.setattr(agent.AgentRun, "run", original_run)
    use_fake(monkeypatch, FakeOllama([("tool", "finish", {"answer": "fine"})], {"_Plan": [plan_json("finish")]}))
    second, _ = run_task(store, "next")
    assert second.status == TaskStatus.SUCCEEDED


# ---------------------------------------------------------------- stop conditions
def test_step_limit(monkeypatch, store):
    monkeypatch.setattr(settings, "WB_AGENT_MAX_STEPS", 3)
    use_fake(monkeypatch, FakeOllama([("tool", "search_knowledge", {"query": "again"})],
                                     {"_Plan": [plan_json("search_knowledge", "finish")]}, repeat_last=True))
    state, events = run_task(store, "Loop forever")
    assert state.status == TaskStatus.FAILED and state.error.code == "AGENT_STEP_LIMIT"
    # the same call 3 times: runs once, the 2 repeats are skipped but still use up steps
    assert len([e for e in events if e.type == EventType.STEP_START]) == 1
    assert sum("Skipped repeated call" in w for w in _warns(events)) == 2
    assert events[-1].type == EventType.ERROR and events[-1].data["error"]["code"] == "AGENT_STEP_LIMIT"


def test_timeout(monkeypatch, store):
    monkeypatch.setattr(settings, "WB_AGENT_TIMEOUT_S", 0.3)
    use_fake(monkeypatch, FakeOllama([("tool", "search_knowledge", {"query": "slow"})],
                                     {"_Plan": [plan_json("finish")]}, repeat_last=True, delay_s=0.15))
    state, events = run_task(store, "Slow task")
    assert state.status == TaskStatus.FAILED and state.error.code == "AGENT_TIMEOUT"
    assert events[-1].data["error"]["code"] == "AGENT_TIMEOUT"
    assert state.elapsed_s < 2


def test_cancel_in_the_middle(monkeypatch, store):
    holder: dict[str, Any] = {}

    def cancel_on_second_turn(turn: int) -> None:
        if turn == 2:
            store.cancel(holder["task_id"])

    use_fake(monkeypatch, FakeOllama([("tool", "search_knowledge", {"query": "a"})],
                                     {"_Plan": [plan_json("finish")]}, repeat_last=True,
                                     on_step=cancel_on_second_turn))
    original_create = store.create

    def create_and_remember(*args, **kwargs):
        state = original_create(*args, **kwargs)
        holder["task_id"] = state.task_id
        return state

    monkeypatch.setattr(store, "create", create_and_remember)
    state, events = run_task(store, "Cancel me")
    assert state.status == TaskStatus.CANCELLED and state.error.code == "CANCELLED"
    steps = [e for e in events if e.type == EventType.STEP_START]
    assert len(steps) == 1                                   # the cancelled turn never ran its tool
    assert events[-1].type == EventType.ERROR and events[-1].data["error"]["code"] == "CANCELLED"


# ---------------------------------------------------------------- validation of the drafted note
def test_invented_sop_reference_is_dropped(monkeypatch, store, report_file_id):
    note = note_json(sop_references=["SOP-INSP-012.pdf, p.4", "SOP-FAKE-999, p.3"])
    steps, json_replies = happy_steps(report_file_id, note)
    use_fake(monkeypatch, FakeOllama(steps, json_replies))
    state, events = run_task(store, "Draft note", [report_file_id])
    assert state.status == TaskStatus.SUCCEEDED
    assert any("SOP-FAKE-999" in w and "Dropped" in w for w in _warns(events))
    doc = _docx_of(state)
    sop = next(t for t in doc.tables if [c.text for c in t.rows[0].cells] == ["S.No", "Document", "Page"])
    assert [[c.text for c in r.cells] for r in sop.rows[1:]] == [["1", "SOP-INSP-012.pdf", "4"]]


def test_sop_reference_to_unretrieved_page_is_corrected(monkeypatch, store, report_file_id):
    note = note_json(sop_references=["SOP-INSP-012, p.9"])
    steps, json_replies = happy_steps(report_file_id, note)
    use_fake(monkeypatch, FakeOllama(steps, json_replies))
    state, events = run_task(store, "Draft note", [report_file_id])
    assert any("page 9 was not retrieved" in w for w in _warns(events))
    sop = next(t for t in _docx_of(state).tables if t.rows[0].cells[1].text == "Document")
    assert sop.rows[1].cells[2].text == "4"


def test_invalid_source_page_is_fixed(monkeypatch, store, report_file_id):
    findings = [Finding(item="Shell course 2", observation="Thinning", severity=Severity.HIGH, source_page=7),
                Finding(item="Bottom plate", observation="Pitting", severity=Severity.MEDIUM, source_page=2)]
    steps, json_replies = happy_steps(report_file_id, note_json(findings=findings))
    use_fake(monkeypatch, FakeOllama(steps, json_replies))
    state, events = run_task(store, "Draft note", [report_file_id])
    assert any("page 7 does not exist" in w for w in _warns(events))
    table = next(t for t in _docx_of(state).tables if t.rows[0].cells[1].text == "Item")
    assert [r.cells[4].text for r in table.rows[1:]] == ["-", "2"]


def test_kb_hits_below_minimum_are_ignored(monkeypatch):
    ctx = agent_tools.ToolContext(task_id="t", emit=lambda *a, **k: None, add_artifact=lambda a: None)
    outcome = agent_tools.search_knowledge(ctx, "wall thickness")
    assert "kb_1: SOP-INSP-012.pdf p.4" in outcome.summary and "SOP-ADM-001" not in outcome.summary
    assert "1 weaker hit(s) below 0.50 ignored" in outcome.summary
    assert list(ctx.scratchpad.kb_hits) == ["kb_1"]

    monkeypatch.setattr(agent_tools.knowledge, "search", lambda query, top_k=4: [WEAK])
    outcome = agent_tools.search_knowledge(ctx, "canteen")
    assert outcome.summary.startswith("No relevant SOP passages") and "Do not cite" in outcome.summary


# ---------------------------------------------------------------- cost check through the agent
@pytest.mark.parametrize("cost, expected", [
    ("Rs 4,50,000 as per contractor quote", "Rs 4,50,000 as per contractor quote"),
    ("Rs 12,00,000 estimated", COST_PLACEHOLDER),
])
def test_cost_check_uses_full_document_text(monkeypatch, store, report_file_id, cost, expected):
    steps, json_replies = happy_steps(report_file_id, note_json(cost_implication=cost))
    use_fake(monkeypatch, FakeOllama(steps, json_replies))
    state, _ = run_task(store, "Draft note", [report_file_id])
    doc = _docx_of(state)
    paragraphs = [p.text for p in doc.paragraphs]
    assert paragraphs[paragraphs.index("Cost implication") + 1] == expected


def test_cost_number_format_behaviour():
    """Documents today's behaviour for reformatted amounts (reported, office.py unchanged)."""
    source = "Contractor quote attached: Rs 4,50,000 for plate replacement."
    assert cost_text("Rs 450000", source) == "Rs 450000"            # commas are ignored when comparing
    assert cost_text("Rs 450,000", source) == "Rs 450,000"          # western grouping also matches
    assert cost_text("Rs 4.5 lakh", source) == COST_PLACEHOLDER     # words/units are not converted
    assert cost_text("Rs 4,50,000.00", source) == "Rs 4,50,000.00"  # trailing .00 is the same whole number
    assert cost_text("Rs 4,50,000.50", source) == COST_PLACEHOLDER  # real decimals are a different number
    assert cost_text("Rs 4,50,000", "quote Rs 4,50,000.00") == "Rs 4,50,000"   # works both ways


# ---------------------------------------------------------------- context trimming
def test_context_trimming_keeps_pinned_messages():
    pinned = [{"role": "system", "content": SYSTEM_PROMPT},
              {"role": "user", "content": "Task: draft a note"},
              {"role": "assistant", "content": "Plan:\n1. Read [read_document]"}]
    exchanges = []
    for i in range(40):
        exchanges.append({"role": "assistant", "content": "", "tool_calls": [{"function": {"name": "search_knowledge",
                                                                                          "arguments": {"query": i}}}]})
        exchanges.append({"role": "tool", "tool_name": "search_knowledge", "content": f"result {i} " + "x" * 1400})
    messages = pinned + exchanges
    budget = agent.context_budget()
    assert estimate_tokens(messages) > budget

    trimmed = trim_conversation(messages)

    assert trimmed[:PINNED_MESSAGES] == pinned
    assert trimmed[PINNED_MESSAGES]["content"] == TRIM_NOTE
    assert estimate_tokens(trimmed) <= budget
    assert trimmed[-1]["content"].startswith("result 39")               # newest result kept
    assert trimmed[PINNED_MESSAGES + 1]["role"] == "assistant"            # no orphan tool result
    assert trim_conversation(trimmed) == trimmed                          # stable on re-trim
    short = pinned + exchanges[:2]
    assert trim_conversation(short) == short


def test_default_plan_per_task_type():
    assert [s.tool for s in agent.default_plan(TaskType.VISION, True)] == ["read_document", "extract_pid_tags", "finish"]
    assert [s.tool for s in agent.default_plan(TaskType.CODING, False)] == ["run_code_task", "finish"]
    assert all(isinstance(s, PlanStep) for s in agent.default_plan(TaskType.DOCUMENT, True))


def test_tool_definitions_format():
    defs = agent_tools.tool_definitions()
    assert [d["function"]["name"] for d in defs] == ["read_document", "search_knowledge", "draft_approval_note",
                                                      "extract_pid_tags", "run_code_task", "finish"]
    for d in defs:
        assert d["type"] == "function" and d["function"]["parameters"]["type"] == "object"
        assert set(d["function"]["parameters"]["required"]) <= set(d["function"]["parameters"]["properties"])


# ================================================================ guided flows (A8 second half)
from backend.flows import guided  # noqa: E402
from shared.contracts import CodeResult, PidTag, PidTagList, Scenario  # noqa: E402

FAKE_CODE = CodeResult(code="def pipe_wall_thickness(P, D, S):\n    return P * D / (2 * S)\n",
                       tests="def test_a():\n    assert 1 == 1\n", passed=3, failed=0, attempts=1,
                       stdout_tail="t = 10 * 200 / (2 * 138) = 7.246 mm\n")


def run_guided_task(store, message, file_ids=None, scenario=None, mode=TaskMode.GUIDED, timeout_s=30):
    state = store.create(message, file_ids or [], mode, scenario)
    deadline = time.time() + timeout_s
    while time.time() < deadline:
        current = store.get(state.task_id)
        if current.status in TERMINAL_STATUSES:
            return current, store.events(state.task_id, 0)[0]
        time.sleep(0.02)
    raise AssertionError("task did not finish")


def _plan_tools(events):
    plan = next(e for e in events if e.type == EventType.PLAN)
    return [s["tool"] for s in plan.data["steps"]]


def test_guided_scenario_a(monkeypatch, store, report_file_id):
    use_fake(monkeypatch, FakeOllama([], {"ApprovalNote": [note_json(cost_implication="Rs 4,50,000")]}))
    state, events = run_guided_task(store, "Approval note please", [report_file_id], Scenario.INSPECTION_NOTE)
    assert state.status == TaskStatus.SUCCEEDED, state.error
    assert _plan_tools(events) == ["read_document", "search_knowledge", "draft_approval_note", "finish"]
    assert [e.data["title"] for e in events if e.type == EventType.STEP_START] == \
        ["read_document", "search_knowledge", "draft_approval_note", "finish"]
    assert not any(e.type == EventType.LLM_CALL and e.data["purpose"] in ("plan", "agent_step") for e in events)
    for event in events:
        assert set(event.data) == EVENT_KEYS[event.type]
    call = next(e for e in events if e.type == EventType.TOOL_CALL and e.data["tool"] == "search_knowledge")
    assert "6.1 mm" in call.data["args"]["query"]                       # query built from the finding lines
    draft = next(e for e in events if e.type == EventType.TOOL_CALL and e.data["tool"] == "draft_approval_note")
    assert draft.data["args"] == {"doc_id": "doc_1", "kb_ref_ids": ["kb_1"]}
    assert "1 findings (1 high/critical)" in state.final_answer
    assert state.final_answer.endswith("_approval_note.docx")


def test_guided_scenario_b(monkeypatch, store):
    seen = {}

    def fake_code_task(request, emit=None, seed_code=None):
        seen["request"] = request
        return FAKE_CODE

    monkeypatch.setattr(agent_tools.code_flow, "run_code_task", fake_code_task)
    state, events = run_guided_task(store, "pipe wall thickness code", scenario=Scenario.CODE_CALC)
    assert state.status == TaskStatus.SUCCEEDED
    assert _plan_tools(events) == ["run_code_task", "finish"]
    assert "7.246 mm" in state.final_answer and "3 passed, 0 failed" in state.final_answer
    assert sorted(a.kind for a in state.artifacts) == ["py", "py"]
    assert seen["request"] == "pipe wall thickness code"               # same text: not appended twice


def test_guided_scenario_c(monkeypatch, store, tmp_path):
    from PIL import Image
    png = tmp_path / "pid.png"
    Image.new("RGB", (64, 64), "white").save(png)
    file_id = file_store.save("pid.png", png.read_bytes(), "image/png").file_id
    tiles = [agent_tools.PageResult(page=1, text=f"tile text {n}", method="vision_tiles", tile=n)
             for n in ("top-left", "top-right", "bottom-left", "bottom-right")]
    monkeypatch.setattr(agent_tools, "extract", lambda source, kind="auto": SimpleNamespace(pages=tiles))
    tags = PidTagList(tags=[PidTag(tag="P-101A", equipment_type="Pump", tile=1),
                            PidTag(tag="P-101A", equipment_type="Pump", tile=2),
                            PidTag(tag="V-201", equipment_type="Valve", tile=3),
                            PidTag(tag="FIC-101", equipment_type="Instrument", tile=9)])
    use_fake(monkeypatch, FakeOllama([], {"PidTagList": [tags.model_dump_json()]}))
    state, events = run_guided_task(store, "tags please", [file_id], Scenario.PID_TAGS)
    assert state.status == TaskStatus.SUCCEEDED, state.error
    assert _plan_tools(events) == ["read_document", "extract_pid_tags", "finish"]
    read = next(e for e in events if e.type == EventType.TOOL_CALL and e.data["tool"] == "read_document")
    assert read.data["args"]["kind"] == "pid"
    assert state.final_answer.startswith("Found 3 unique tags")
    assert "Pump: 1" in state.final_answer and state.artifacts[0].kind == "xlsx"
    assert any("tile 9 does not exist" in w for w in _warns(events))


def test_guided_mode_without_scenario_is_bad_request(store):
    state, _ = run_guided_task(store, "hello", scenario=None)
    assert state.status == TaskStatus.FAILED and state.error.code == "BAD_REQUEST"


def test_guided_step_failure_uses_tool_error_code(monkeypatch, store):
    def failing(request, emit=None, seed_code=None):
        raise agent_tools.code_flow.CodeFlowError("SANDBOX_UNAVAILABLE", "docker down")

    monkeypatch.setattr(agent_tools.code_flow, "run_code_task", failing)
    state, _ = run_guided_task(store, "code", scenario=Scenario.CODE_CALC)
    assert state.status == TaskStatus.FAILED and state.error.code == "SANDBOX_UNAVAILABLE"


def test_agent_falls_back_to_guided_on_step_limit(monkeypatch, store, report_file_id):
    monkeypatch.setattr(settings, "WB_AGENT_MAX_STEPS", 2)
    use_fake(monkeypatch, FakeOllama([("tool", "search_knowledge", {"query": "loop"})],
                                     {"_Plan": [plan_json("search_knowledge", "finish")],
                                      "ApprovalNote": [note_json()]}, repeat_last=True))
    state, events = run_guided_task(store, "Draft approval note", [report_file_id], mode=TaskMode.AGENT)
    assert state.status == TaskStatus.SUCCEEDED, state.error
    assert any(w.startswith(agent.FALLBACK_MESSAGE) and "AGENT_STEP_LIMIT" in w for w in _warns(events))
    plans = [e for e in events if e.type == EventType.PLAN]
    assert len(plans) == 2
    assert [p.tool for p in state.plan] == ["read_document", "search_knowledge", "draft_approval_note", "finish"]
    assert state.artifacts[0].kind == "docx"
    assert not any(e.type == EventType.ERROR for e in events)


def test_agent_without_matching_scenario_does_not_fall_back(monkeypatch, store):
    monkeypatch.setattr(settings, "WB_AGENT_MAX_STEPS", 2)
    use_fake(monkeypatch, FakeOllama([("tool", "search_knowledge", {"query": "loop"})],
                                     {"_Plan": [plan_json("finish")]}, repeat_last=True))
    state, events = run_guided_task(store, "General question", mode=TaskMode.AGENT)
    assert state.status == TaskStatus.FAILED and state.error.code == "AGENT_STEP_LIMIT"
    assert not any(agent.FALLBACK_MESSAGE in w for w in _warns(events))


def test_explicit_scenario_without_deliverable_falls_back(monkeypatch, store, report_file_id):
    use_fake(monkeypatch, FakeOllama([("tool", "finish", {"answer": "I read it, looks fine."})],
                                     {"_Plan": [plan_json("finish")], "ApprovalNote": [note_json()]}))
    state, events = run_guided_task(store, "Note", [report_file_id], Scenario.INSPECTION_NOTE, mode=TaskMode.AGENT)
    assert state.status == TaskStatus.SUCCEEDED
    assert any("no docx file" in w for w in _warns(events)) and state.artifacts[0].kind == "docx"


def test_pick_scenario_rule(tmp_path):
    from PIL import Image
    png = tmp_path / "x.png"
    Image.new("RGB", (8, 8)).save(png)
    image_id = file_store.save("x.png", png.read_bytes(), "image/png").file_id
    text_id = file_store.save("r.txt", b"report text", "text/plain").file_id

    def decision(task_type):
        return RouteDecision(task_type=task_type, model_id="general", ollama_name="x", reason="t",
                             layer="rule", confidence=1)

    assert guided.pick_scenario(Scenario.CODE_CALC, decision(TaskType.DOCUMENT), [text_id]) == Scenario.CODE_CALC
    assert guided.pick_scenario(None, decision(TaskType.GENERAL), [image_id]) == Scenario.PID_TAGS
    assert guided.pick_scenario(None, decision(TaskType.CODING), []) == Scenario.CODE_CALC
    assert guided.pick_scenario(None, decision(TaskType.DOCUMENT), [text_id]) == Scenario.INSPECTION_NOTE
    assert guided.pick_scenario(None, decision(TaskType.DOCUMENT), []) is None
    assert guided.pick_scenario(None, decision(TaskType.GENERAL), [text_id]) is None


def test_search_query_from_text():
    text = "INSPECTION REPORT\nTank T-104\nShell course 2: wall thinning to 6.1 mm\nRoof: surface corrosion\nSigned"
    assert guided.search_query_from_text(text) == "Shell course 2: wall thinning to 6.1 mm Roof: surface corrosion"
    assert guided.search_query_from_text("Just some words here") == "Just some words here"
    assert len(guided.search_query_from_text("thinning " * 200)) <= guided.SEARCH_QUERY_MAX_CHARS


def test_api_guided_mode_runs_guided_flow(monkeypatch):
    from fastapi.testclient import TestClient

    from backend.main import app
    from shared.contracts import API_PREFIX, TaskState

    monkeypatch.setattr(agent_tools.code_flow, "run_code_task", lambda request, emit=None, seed_code=None: FAKE_CODE)
    with TestClient(app) as client:
        resp = client.post(f"{API_PREFIX}/tasks", json={"message": "Pipe wall thickness code", "mode": "guided",
                                                         "scenario": "code_calc"})
        assert resp.status_code == 202
        task_id = resp.json()["task_id"]
        deadline = time.time() + 30
        while time.time() < deadline:
            page = client.get(f"{API_PREFIX}/tasks/{task_id}/events", params={"after": 0}).json()
            if page["done"]:
                break
            time.sleep(0.1)
        state = TaskState.model_validate(client.get(f"{API_PREFIX}/tasks/{task_id}").json())
    plan = next(e for e in page["events"] if e["type"] == "plan")
    assert [s["tool"] for s in plan["data"]["steps"]] == [t for _, t in guided.GUIDED_PLANS[Scenario.CODE_CALC]]
    assert plan["title"].startswith("Guided plan (code_calc)")
    assert state.status == TaskStatus.SUCCEEDED
    assert state.mode == TaskMode.GUIDED and state.scenario == Scenario.CODE_CALC


def test_empty_sop_references_are_filled_from_chosen_hits(monkeypatch, store, report_file_id):
    steps, json_replies = happy_steps(report_file_id, note_json(sop_references=[]))
    use_fake(monkeypatch, FakeOllama(steps, json_replies))
    state, events = run_task(store, "Draft note", [report_file_id])
    assert state.status == TaskStatus.SUCCEEDED
    infos = [e.data["text"] for e in events if e.type == EventType.LOG and e.data["level"] == "info"]
    assert any("cited the retrieved passages" in t for t in infos)
    sop = next(t for t in _docx_of(state).tables if t.rows[0].cells[1].text == "Document")
    assert [[c.text for c in r.cells] for r in sop.rows[1:]] == [["1", "SOP-INSP-012.pdf", "4"]]


def test_search_query_skips_labels_and_scope():
    text = ("1. Scope: ultrasonic thickness survey.\n2. Findings Finding 1 - Shell course 2: wall thickness 6.1 mm. "
            "Severity: HIGH.\nFinding 2 - Bottom plate pitting corrosion 1.2 mm.")
    assert guided.search_query_from_text(text) == \
        "Shell course 2: wall thickness 6.1 mm. Bottom plate pitting corrosion 1.2 mm."


def test_plan_is_sent_as_user_turn(monkeypatch, store):
    """After an assistant-role plan the model replied with nothing (seen live); the plan must be a user turn."""
    fake = use_fake(monkeypatch, FakeOllama([("tool", "finish", {"answer": "ok"})], {"_Plan": [plan_json("finish")]}))
    run_task(store, "Say ok")
    first = fake.step_calls[0]
    assert [m["role"] for m in first] == ["system", "user", "user"]
    assert first[2]["content"].startswith("Plan:") and first[2]["content"].endswith(agent.PLAN_FOLLOW_UP)


# ================================================================ A8b follow-ups
def _tool_calls(events, tool):
    return [e for e in events if e.type == EventType.TOOL_CALL and e.data["tool"] == tool]


def test_auto_finish_when_done_and_model_starts_other_deliverable(monkeypatch, store, report_file_id):
    steps, json_replies = happy_steps(report_file_id)
    steps[-1] = ("tool", "extract_pid_tags", {"doc_id": "doc_1"})          # instead of finish
    use_fake(monkeypatch, FakeOllama(steps, json_replies))
    state, events = run_task(store, "Draft an approval note", [report_file_id])
    assert state.status == TaskStatus.SUCCEEDED, state.error
    assert agent.AUTO_FINISH_MESSAGE in _warns(events)
    assert _tool_calls(events, "extract_pid_tags") == []                     # never ran
    assert state.final_answer.startswith("Finished automatically. Files produced: INSP-")
    assert state.final_answer.endswith("_approval_note.docx.")
    assert events[-1].type == EventType.FINAL


def test_auto_finish_when_done_and_model_repeats_the_deliverable_call(monkeypatch, store, report_file_id):
    steps, json_replies = happy_steps(report_file_id)
    steps[-1] = steps[2]                                                      # draft again, same args
    use_fake(monkeypatch, FakeOllama(steps, json_replies))
    state, events = run_task(store, "Draft an approval note", [report_file_id])
    assert state.status == TaskStatus.SUCCEEDED
    assert agent.AUTO_FINISH_MESSAGE in _warns(events)
    assert len(_tool_calls(events, "draft_approval_note")) == 1 and len(state.artifacts) == 1


def test_same_call_twice_is_not_run_again(monkeypatch, store):
    fake = use_fake(monkeypatch, FakeOllama([("tool", "search_knowledge", {"query": "loto"}),
                                             ("tool", "search_knowledge", {"query": "loto"}),
                                             ("tool", "search_knowledge", {"query": "lockout"}),
                                             ("tool", "finish", {"answer": "ok"})],
                                            {"_Plan": [plan_json("search_knowledge", "finish")]}))
    state, events = run_task(store, "What is LOTO?")
    assert state.status == TaskStatus.SUCCEEDED
    queries = [e.data["args"]["query"] for e in _tool_calls(events, "search_knowledge")]
    assert queries == ["loto", "lockout"]                                     # the repeat was skipped
    told = fake.step_calls[2][-1]["content"]
    assert told.startswith("You already called search_knowledge with these arguments")
    assert "kb_1: SOP-INSP-012.pdf" in told                                   # it got the earlier result back
    assert any("Skipped repeated call" in w for w in _warns(events))


def test_call_key_ignores_argument_order():
    assert agent.call_key("t", {"a": 1, "b": 2}) == agent.call_key("t", {"b": 2, "a": 1})
    assert agent.call_key("t", {"a": 1}) != agent.call_key("t", {"a": 2})


def test_time_budget_switches_to_guided(monkeypatch, store, report_file_id):
    monkeypatch.setattr(settings, "WB_AGENT_TIMEOUT_S", 1.0)
    use_fake(monkeypatch, FakeOllama([("tool", "search_knowledge", {"query": "slow"}),
                                      ("tool", "search_knowledge", {"query": "slower"}),
                                      ("tool", "search_knowledge", {"query": "slowest"})],
                                     {"_Plan": [plan_json("finish")], "ApprovalNote": [note_json()]},
                                     step_delay_s=0.35))
    state, events = run_task(store, "Draft approval note", [report_file_id])
    assert state.status == TaskStatus.SUCCEEDED, state.error
    assert agent.TIME_SWITCH_MESSAGE in _warns(events)
    assert state.artifacts[0].kind == "docx" and state.elapsed_s < 1.0
    assert [p.tool for p in state.plan] == ["read_document", "search_knowledge", "draft_approval_note", "finish"]


def test_time_budget_does_not_switch_once_deliverable_exists(monkeypatch):
    run = agent.AgentRun.__new__(agent.AgentRun)
    run.started, run.scenario = time.monotonic() - 1000, Scenario.INSPECTION_NOTE
    run.ctx = SimpleNamespace(artifacts=[SimpleNamespace(kind="docx")])
    run.check_time_budget()                                                   # no switch: file exists
    run.ctx = SimpleNamespace(artifacts=[])
    with pytest.raises(agent._SwitchToGuided):
        run.check_time_budget()


def test_tag_types_checked_against_prefix_table():
    warns: list[str] = []
    ctx = agent_tools.ToolContext(task_id="t", emit=lambda *a, **k: warns.append(a[2].get("text", "")),
                                  add_artifact=lambda a: None)
    tags = PidTagList(tags=[
        PidTag(tag="P-101A", equipment_type="Valve"),            # wrong -> Pump
        PidTag(tag="t-301", equipment_type="Instrument"),        # wrong -> Tank (case-insensitive prefix)
        PidTag(tag="FIC-101", equipment_type="Instrument"),      # generic instrument is fine
        PidTag(tag="PT-102", equipment_type="Pressure transmitter"),
        PidTag(tag="PSV-7", equipment_type="Relief valve"),
        PidTag(tag="XV-201", equipment_type="Valve"),            # unknown prefix: keep the model's answer
        PidTag(tag="E-401", equipment_type="Heat exchanger"),
        PidTag(tag="LIC-3", equipment_type="Pump"),              # wrong -> Level indicating controller
    ])
    agent_tools.check_tag_types(ctx, tags)
    assert [t.equipment_type for t in tags.tags] == [
        "Pump", "Tank", "Instrument", "Pressure transmitter", "Relief valve", "Valve", "Heat exchanger",
        "Level indicating controller"]
    assert len([w for w in warns if "prefix" in w]) == 3
    assert agent_tools.tag_prefix(" fic-101 ") == "FIC" and agent_tools.tag_prefix("101") == ""


def test_auto_added_sop_references_are_labelled_in_word(monkeypatch, store, report_file_id):
    steps, json_replies = happy_steps(report_file_id, note_json(sop_references=[]))
    use_fake(monkeypatch, FakeOllama(steps, json_replies))
    state, _ = run_task(store, "Draft note", [report_file_id])
    paragraphs = [p.text for p in _docx_of(state).paragraphs]
    assert office.SOP_AUTO_NOTE in paragraphs


def test_model_cited_sop_references_have_no_auto_label(monkeypatch, store, report_file_id):
    steps, json_replies = happy_steps(report_file_id)
    use_fake(monkeypatch, FakeOllama(steps, json_replies))
    state, _ = run_task(store, "Draft note", [report_file_id])
    paragraphs = [p.text for p in _docx_of(state).paragraphs]
    assert office.SOP_AUTO_NOTE not in paragraphs and not any("{{" in p for p in paragraphs)
