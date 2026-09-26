"""
Ticket A10 (performance): prewarm, the P&ID OCR fast path and its fallback.
Fast tests mock Ollama and OCR; the slow test runs Scenario C live on the demo P&ID.
"""
from __future__ import annotations

import io
import time
from pathlib import Path

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image

from backend import agent_tools, llm_client, prewarm, router
from backend.llm_client import ChatResult, LLMError
from backend.settings import settings
from backend.tools import documents, knowledge, office
from backend.tools.documents import PID_FAST_MIN_TAGS, PageResult, find_tag_candidates
from shared.contracts import API_PREFIX, EventType, PrewarmResult, TaskState, TaskStatus

_REPO_ROOT = Path(__file__).resolve().parent.parent.parent
DEMO_PID = _REPO_ROOT / "demo" / "inputs" / "scenario_c_pid_generated.png"
DEMO_TAGS = {"T-201", "XV-201", "P-201A", "P-201B", "E-201", "V-201", "FT-201", "PT-202", "LT-201", "LT-202",
             "TT-203", "PSV-201"}
SCENARIO_C_LIMIT_S = 300


@pytest.fixture(autouse=True)
def _temp_cache(tmp_path, monkeypatch):
    """Extraction results go to a temp cache, never data/cache (and never a cache hit between tests)."""
    monkeypatch.setattr(settings, "WB_CACHE_DIR", tmp_path / "cache")


# ---------------------------------------------------------------- prewarm
@pytest.fixture
def fake_prewarm(monkeypatch):
    """All prewarm items succeed without Ollama, Chroma or Docker; records the model load order."""
    loads: list[str] = []
    monkeypatch.setattr(llm_client, "load_model", lambda name, embedding=False: loads.append(name) or 5)
    monkeypatch.setattr(llm_client, "loaded_models", lambda: ["qwen3.5:4b", "bge-m3:latest"])
    monkeypatch.setattr(router, "_load_example_vectors", lambda: {})
    monkeypatch.setattr(knowledge, "search", lambda query, top_k=4: [])
    monkeypatch.setattr(prewarm, "sandbox_available", lambda: True)
    return loads


def test_prewarm_all_items_ok_general_model_last(fake_prewarm):
    result, items, loaded = prewarm.run_prewarm()
    PrewarmResult.model_validate(result.model_dump())
    assert result.failed == [] and len(result.warmed) == 6
    assert all(label.endswith(" ms)") for label in result.warmed)
    assert [i.name.split()[0] for i in items] == ["model", "model", "router", "kb", "sandbox", "model"]
    assert fake_prewarm == ["qwen2.5-coder:3b", "bge-m3", "qwen3.5:4b"]     # general last: survives eviction
    assert loaded == ["qwen3.5:4b", "bge-m3:latest"]


def test_prewarm_reports_failed_items_without_crashing(fake_prewarm, monkeypatch):
    def load(name, embedding=False):
        if name.startswith("qwen2.5-coder"):
            raise LLMError("MODEL_UNAVAILABLE", "coder not pulled")
        return 5

    monkeypatch.setattr(llm_client, "load_model", load)
    monkeypatch.setattr(prewarm, "sandbox_available", lambda: False)
    result, _, _ = prewarm.run_prewarm()
    assert len(result.warmed) == 4
    assert [f.split(":")[0] for f in result.failed] == ["model coder (qwen2.5-coder", "sandbox image"]
    assert "coder not pulled" in result.failed[0]


def test_prewarm_survives_a_broken_registry(monkeypatch):
    def broken():
        raise RuntimeError("models.yaml is broken")

    monkeypatch.setattr(prewarm, "prewarm_steps", broken)
    monkeypatch.setattr(prewarm, "_loaded", lambda: [])
    result, _, _ = prewarm.run_prewarm()
    assert result.warmed == [] and result.failed == ["model registry: RuntimeError: models.yaml is broken"]


def test_prewarm_endpoint_reports_failed_item(fake_prewarm, monkeypatch):
    def locked(query, top_k=4):
        raise OSError("chroma locked")

    monkeypatch.setattr(knowledge, "search", locked)
    from backend.main import app

    with TestClient(app) as client:
        resp = client.post(f"{API_PREFIX}/admin/prewarm")
    assert resp.status_code == 200
    result = PrewarmResult.model_validate(resp.json())
    assert result.failed == ["kb collection: OSError: chroma locked"]
    assert '"name": "kb collection"' in resp.headers["X-Prewarm-Items"]
    assert resp.headers["X-Prewarm-Loaded-Models"] == "qwen3.5:4b,bge-m3:latest"


def test_load_model_uses_real_call_options_and_keep_alive(monkeypatch):
    calls: list[dict] = []

    class FakeClient:
        def generate(self, **kwargs):
            calls.append(kwargs)

        def embed(self, **kwargs):
            calls.append(kwargs)

    monkeypatch.setattr(llm_client, "_client", lambda: FakeClient())
    llm_client.load_model("qwen3.5:4b")
    llm_client.load_model("bge-m3", embedding=True)
    assert calls[0]["prompt"] == "" and calls[0]["options"]["num_ctx"] == settings.WB_NUM_CTX
    assert all(c["keep_alive"] == llm_client.OLLAMA_KEEP_ALIVE for c in calls)
    assert calls[1]["model"] == "bge-m3" and calls[1]["input"]


# ---------------------------------------------------------------- tag regex
def test_tag_regex_finds_tag_shapes():
    text = "PT-202 PT FT-201 P-201A\nP-201B PSV-201 TT-203 V-201 LT-202 E-201 XV-201 T-201 FIC-1001 LT-201"
    assert find_tag_candidates(text) == ["PT-202", "FT-201", "P-201A", "P-201B", "PSV-201", "TT-203", "V-201",
                                         "LT-202", "E-201", "XV-201", "T-201", "FIC-1001", "LT-201"]


@pytest.mark.parametrize("text", ["DEMO-C-001", "DRG DEMO-C-001 REV 0", "UNIT 300", "TO UNIT 300", "CW",
                                  "SAHYADRI DEMO REFINERY", "P-1", "P-12345", "ABCDE-201", "FT", "6-P-101"])
def test_tag_regex_rejects_non_tags(text):
    assert find_tag_candidates(text) == []


def test_tag_regex_dedupes_and_uppercases():
    assert find_tag_candidates("p-101a and P-101A, v-201") == ["P-101A", "V-201"]


# ---------------------------------------------------------------- path selection
def _drawing(tmp_path) -> Path:
    path = tmp_path / "pid.png"
    Image.new("RGB", (400, 300), "white").save(path)
    return path


def _patch_ocr_and_vision(monkeypatch, tile_texts: list[str], vision_text: str = "") -> list[str]:
    """OCR returns tile_texts in tile order; every vision call is recorded by its purpose."""
    purposes: list[str] = []
    replies = iter(tile_texts)
    monkeypatch.setattr(documents, "_ocr_pid_tile", lambda tile: next(replies))
    monkeypatch.setattr(documents, "_deskew", lambda image: (image, 0.0))

    def fake_chat(model_id, messages, tools=None, images=None, *, purpose="chat"):
        purposes.append(purpose)
        return ChatResult(text=vision_text)

    monkeypatch.setattr(documents, "chat", fake_chat)
    return purposes


def test_enough_ocr_tags_uses_fast_path_with_one_vision_call(tmp_path, monkeypatch):
    purposes = _patch_ocr_and_vision(monkeypatch, ["PT-202 PT", "FT-201 E-201", "P-201A", "DRG DEMO-C-001"],
                                     "PT-202\nFT-201\nE-201\nP-201A")
    pages = documents.extract(_drawing(tmp_path), kind="pid").pages
    assert [p.method for p in pages] == ["ocr_tiles"] * 4 + ["vision_confirm"]
    assert [p.tile for p in pages][-1] == documents.PID_WHOLE_DRAWING
    assert pages[3].text == ""                                   # drawing number is not a tag
    assert purposes == ["document_extract_pid_confirm"]          # the vision model runs ONCE
    assert "fast path" in documents.pid_path_note(pages) and "4 confirm OCR" in documents.pid_path_note(pages)


def test_too_few_ocr_tags_falls_back_to_four_vision_tiles(tmp_path, monkeypatch):
    purposes = _patch_ocr_and_vision(monkeypatch, ["P-101", "", "", ""], "P-101")
    pages = documents.extract(_drawing(tmp_path), kind="pid").pages
    assert [p.method for p in pages] == ["vision_tiles"] * 4
    assert purposes == ["document_extract_pid"] * 4
    assert "4-tile vision path" in documents.pid_path_note(pages) and "only 1 tag" in pages[0].note


def test_use_fast_path_threshold():
    pages = [PageResult(page=1, text="\n".join(f"P-10{i}" for i in range(n)), method="ocr_tiles", tile="top-left")
             for n in (PID_FAST_MIN_TAGS - 1, PID_FAST_MIN_TAGS)]
    assert documents.use_pid_fast_path([pages[0]]) is False
    assert documents.use_pid_fast_path([pages[1]]) is True


# ---------------------------------------------------------------- merging OCR + vision
def _ctx(events: list):
    return agent_tools.ToolContext(task_id="t", emit=lambda *a, **k: events.append(a), add_artifact=lambda a: None)


def _fast_pages(vision: str, error: str | None = None) -> list[PageResult]:
    tiles = {"top-left": "PT-202", "top-right": "PT-202\nFT-201\nV-201", "bottom-left": "P-201A\nT-201",
             "bottom-right": ""}
    pages = [PageResult(page=1, text=text, method="ocr_tiles", tile=name) for name, text in tiles.items()]
    pages[0].note = "P&ID read by the fast path"
    return pages + [PageResult(page=1, text=vision, method="vision_confirm", tile="whole drawing", error=error)]


def test_vision_confirm_does_not_duplicate_ocr_tags():
    events: list = []
    tag_list = agent_tools.merge_pid_tags(_ctx(events), _fast_pages("PT-202\nFT-201\nV-201\nP-201A\nT-201"))
    rows = office.dedupe_tags(tag_list)
    assert [r["tag"] for r in rows] == ["PT-202", "FT-201", "V-201", "P-201A", "T-201"]
    assert next(r for r in rows if r["tag"] == "PT-202")["tile"] == "top-left, top-right"
    assert all(r["description"] == "OCR, confirmed by vision" for r in rows)
    assert {r["tag"]: r["equipment_type"] for r in rows}["V-201"] == "Vessel"
    assert events == []


def test_vision_adds_missed_tag_but_not_echoes_or_variants():
    events: list = []
    pages = _fast_pages("PT-202\nLT-202\nP-101\nP-201\nP-201B\nLT-202")
    rows = office.dedupe_tags(agent_tools.merge_pid_tags(_ctx(events), pages))
    added = [r for r in rows if r["description"].startswith("Vision only")]
    assert [r["tag"] for r in added] == ["LT-202", "P-201B"] and all(r["tile"] == "" for r in added)
    assert next(r for r in rows if r["tag"] == "FT-201")["description"] == "OCR only (vision did not read it)"
    warnings = " ".join(e[2]["text"] for e in events)
    assert "'P-101'" in warnings and "'P-201'" in warnings       # prompt echo and suffix-less variant ignored


def test_failed_vision_check_keeps_ocr_tags():
    rows = office.dedupe_tags(agent_tools.merge_pid_tags(_ctx([]), _fast_pages("", error="MODEL_TIMEOUT")))
    assert len(rows) == 5 and rows[0]["description"] == "OCR (vision check unavailable)"


def test_fast_path_tool_flow_logs_path_and_skips_the_llm(monkeypatch, tmp_path):
    from backend.file_store import file_store

    monkeypatch.setattr(settings, "WB_WORKSPACE_DIR", tmp_path / "ws")
    file_id = file_store.save("pid.png", _drawing(tmp_path).read_bytes(), "image/png").file_id
    monkeypatch.setattr(agent_tools, "extract", lambda source, kind="auto": type("R", (), {"pages": _fast_pages(
        "PT-202\nFT-201\nV-201\nP-201A\nT-201")})())

    def no_llm(*args, **kwargs):
        raise AssertionError("the fast path must not call the LLM")

    monkeypatch.setattr(llm_client, "chat_json_meta", no_llm)
    events: list = []
    ctx = _ctx(events)
    assert agent_tools.execute_tool(ctx, "read_document", {"file_id": file_id, "kind": "pid"}).ok
    outcome = agent_tools.execute_tool(ctx, "extract_pid_tags", {"doc_id": "doc_1"})
    assert outcome.ok, outcome.summary
    logs = [e[2]["text"] for e in events if e[0] == EventType.LOG]
    assert any(text.startswith("P&ID read by the fast path") for text in logs)
    assert ctx.artifacts[0].kind == "xlsx"


# ---------------------------------------------------------------- live Scenario C
def _ollama_ready() -> bool:
    try:
        return httpx.get(f"{settings.OLLAMA_HOST}/api/tags", timeout=3).status_code == 200
    except httpx.HTTPError:
        return False


@pytest.mark.slow
@pytest.mark.skipif(not DEMO_PID.exists() or not _ollama_ready(), reason="demo P&ID or Ollama not available")
def test_live_scenario_c_finds_12_of_12_tags_under_5_minutes(tmp_path, monkeypatch, network_guard):
    for key in ("WB_WORKSPACE_DIR", "WB_CACHE_DIR", "WB_LOG_DIR"):
        monkeypatch.setattr(settings, key, tmp_path / key.lower())        # no cached extraction, no real logs
    from backend.main import app

    with TestClient(app) as client:
        up = client.post(f"{API_PREFIX}/files", files={"file": (DEMO_PID.name, DEMO_PID.read_bytes(), "image/png")})
        body = {"message": "Extract all equipment and instrument tags from this P&ID drawing into an Excel tag list.",
                "file_ids": [up.json()["file_id"]], "mode": "guided", "scenario": "pid_tags"}
        started = time.monotonic()
        task_id = client.post(f"{API_PREFIX}/tasks", json=body).json()["task_id"]
        events, after = [], 0
        while time.monotonic() - started < SCENARIO_C_LIMIT_S + 60:
            page = client.get(f"{API_PREFIX}/tasks/{task_id}/events", params={"after": after}).json()
            events += page["events"]
            after = page["next_seq"]
            if page["done"]:
                break
            time.sleep(1)
        state = TaskState.model_validate(client.get(f"{API_PREFIX}/tasks/{task_id}").json())
        xlsx = client.get(state.artifacts[0].download_url).content if state.artifacts else b""
    assert state.status == TaskStatus.SUCCEEDED, state.error
    rows = list(load_workbook(io.BytesIO(xlsx))["Tags"].iter_rows(min_row=2, values_only=True))
    tags = [r[0] for r in rows]
    print(f"\nScenario C: {state.elapsed_s:.1f} s, {len(tags)} tags: {tags}")
    print("Path: " + next(e["data"]["text"] for e in events if e["title"] == "P&ID read path"))
    assert set(tags) == DEMO_TAGS and len(tags) == 12                 # 12/12, nothing invented, no duplicates
    assert not {"DEMO-C-001", "UNIT 300", "CW"} & set(tags)
    assert state.elapsed_s < SCENARIO_C_LIMIT_S
